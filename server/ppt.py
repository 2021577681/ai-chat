# ============================================================
# server/ppt.py - PPT 生成 mixin
# ============================================================
# 提供 PptMixin，给 Handler 用。MVP 能力：
#   - 接收结构化 deck/slides JSON
#   - 生成 .pptx 文件到 output/ 目录或用户指定的沙箱内路径
#   - 支持 cover / agenda / bullets / section / summary 基础版式
#   - 支持 three_cards / process / timeline / comparison / two_column / matrix
#     / pyramid / cycle / funnel / quote / architecture / method_pipeline
#     / table / chart / experiment_design / ablation 原生 PPT 矢量版式
# ============================================================

import os
import re

from . import config
from .sandbox import check_path_or_error


_SAFE_FILENAME_RE = re.compile(r'[\\/:*?"<>|]+')


def _as_text(value, default=''):
    if value is None:
        return default
    return str(value)


def _as_list(value):
    if isinstance(value, list):
        return value
    if value is None:
        return []
    return [value]


def _item_text(value):
    if isinstance(value, dict):
        return _as_text(value.get('text') or value.get('title') or value.get('name') or value.get('desc') or value.get('value'))
    return _as_text(value)


def _rows_from_records(records, headers):
    rows = []
    for rec in _as_list(records):
        rows.append([_as_text(rec.get(h, '') if isinstance(rec, dict) else rec) for h in headers])
    return rows


def _safe_filename(name):
    name = _as_text(name, 'generated.pptx').strip() or 'generated.pptx'
    name = _SAFE_FILENAME_RE.sub('_', name)
    if not name.lower().endswith('.pptx'):
        name += '.pptx'
    return name


class PptMixin:
    """Handler mixin：生成 PPTX 文件。"""

    def handle_generate_ppt(self, body):
        try:
            data = body.get('data') or {}
            if not isinstance(data, dict):
                return self._send_json(200, {'ok': False, 'error': 'data 必须是对象'})

            try:
                from pptx import Presentation
                from pptx.util import Inches, Pt
            except ImportError:
                return self._send_json(200, {
                    'ok': False,
                    'error': '缺少依赖 python-pptx，请先运行：pip install -r requirements.txt'
                })

            filename = _safe_filename(data.get('filename') or data.get('file_name'))
            output_path = data.get('path') or data.get('output_path') or os.path.join('output', filename)
            if str(output_path).lower().endswith(os.sep) or str(output_path).endswith('/'):
                output_path = os.path.join(output_path, filename)
            if not str(output_path).lower().endswith('.pptx'):
                output_path = os.path.join(str(output_path), filename)

            out_path, err = check_path_or_error(output_path, must_exist=False)
            if err:
                return self._send_json(200, {'ok': False, 'error': err})

            slides = _as_list(data.get('slides'))
            if not slides:
                slides = [{
                    'type': 'cover',
                    'title': data.get('title') or '未命名 PPT',
                    'subtitle': data.get('subtitle') or ''
                }]

            slides = self._auto_visualize_slides(slides)

            prs = Presentation()
            # python-pptx 默认是 10x7.5（4:3），而矢量版式按 16:9 设计。
            # 不显式设置会导致所有 12.x 英寸坐标横向越界。
            prs.slide_width = Inches(13.333333)
            prs.slide_height = Inches(7.5)
            for slide_data in slides:
                if not isinstance(slide_data, dict):
                    slide_data = {'type': 'bullets', 'title': _as_text(slide_data)}
                slide_type = _as_text(slide_data.get('type') or slide_data.get('layout'), 'bullets').lower()
                if slide_type == 'cover':
                    self._ppt_add_cover(prs, slide_data)
                elif slide_type in ('agenda', 'toc'):
                    self._ppt_add_bullets(prs, {
                        'title': slide_data.get('title') or '目录',
                        'bullets': slide_data.get('items') or slide_data.get('bullets') or []
                    }, Pt)
                elif slide_type in ('section', 'divider'):
                    self._ppt_add_section(prs, slide_data)
                elif slide_type == 'summary':
                    self._ppt_add_bullets(prs, {
                        'title': slide_data.get('title') or '总结',
                        'bullets': slide_data.get('bullets') or slide_data.get('items') or []
                    }, Pt)
                elif slide_type in ('three_cards', 'cards'):
                    self._ppt_add_three_cards(prs, slide_data)
                elif slide_type in ('process', 'flow', 'workflow'):
                    self._ppt_add_process(prs, slide_data)
                elif slide_type == 'timeline':
                    self._ppt_add_timeline(prs, slide_data)
                elif slide_type in ('comparison', 'compare', 'vs'):
                    self._ppt_add_comparison(prs, slide_data)
                elif slide_type in ('two_column', 'two_columns', 'columns'):
                    self._ppt_add_two_column(prs, slide_data)
                elif slide_type in ('matrix', 'quadrant', 'four_quadrant'):
                    self._ppt_add_matrix(prs, slide_data)
                elif slide_type == 'pyramid':
                    self._ppt_add_pyramid(prs, slide_data)
                elif slide_type in ('cycle', 'loop'):
                    self._ppt_add_cycle(prs, slide_data)
                elif slide_type == 'funnel':
                    self._ppt_add_funnel(prs, slide_data)
                elif slide_type in ('quote', 'quotation'):
                    self._ppt_add_quote(prs, slide_data)
                elif slide_type in ('architecture', 'system_architecture', 'tech_architecture'):
                    self._ppt_add_architecture(prs, slide_data)
                elif slide_type in ('method_pipeline', 'research_pipeline', 'pipeline'):
                    self._ppt_add_method_pipeline(prs, slide_data)
                elif slide_type in ('table', 'data_table'):
                    self._ppt_add_table(prs, slide_data)
                elif slide_type in ('chart', 'bar_chart', 'line_chart', 'pie_chart'):
                    self._ppt_add_chart(prs, slide_data)
                elif slide_type in ('experiment_design', 'experiment', 'experimental_design'):
                    self._ppt_add_experiment_design(prs, slide_data)
                elif slide_type in ('ablation', 'ablation_study'):
                    self._ppt_add_ablation(prs, slide_data)
                else:
                    self._ppt_add_bullets(prs, slide_data, Pt)

            parent = os.path.dirname(out_path)
            if parent:
                os.makedirs(parent, exist_ok=True)
            prs.save(out_path)

            rel_path = os.path.relpath(out_path, config.WORKSPACE_ROOT).replace(os.sep, '/')
            self._send_json(200, {
                'ok': True,
                'path': rel_path,
                'slides': len(slides),
                'message': f'PPT 已生成：{rel_path}'
            })
        except Exception as e:
            self._send_json(200, {'ok': False, 'error': str(e)})

    def _ppt_add_cover(self, prs, data):
        slide = prs.slides.add_slide(prs.slide_layouts[0])
        slide.shapes.title.text = _as_text(data.get('title'), '未命名 PPT')
        if len(slide.placeholders) > 1:
            slide.placeholders[1].text = _as_text(data.get('subtitle') or data.get('desc'))

    def _ppt_add_section(self, prs, data):
        layout = prs.slide_layouts[2] if len(prs.slide_layouts) > 2 else prs.slide_layouts[0]
        slide = prs.slides.add_slide(layout)
        slide.shapes.title.text = _as_text(data.get('title'), '章节')
        if len(slide.placeholders) > 1:
            slide.placeholders[1].text = _as_text(data.get('subtitle') or data.get('desc'))

    def _ppt_add_bullets(self, prs, data, Pt):
        slide = prs.slides.add_slide(prs.slide_layouts[1])
        slide.shapes.title.text = _as_text(data.get('title'), '未命名页面')

        body = slide.placeholders[1]
        tf = body.text_frame
        tf.clear()

        bullets = data.get('bullets') or data.get('items') or []
        bullets = _as_list(bullets)
        if not bullets and data.get('content'):
            bullets = _as_list(data.get('content'))

        if not bullets:
            bullets = ['']

        for i, item in enumerate(bullets):
            if isinstance(item, dict):
                text = item.get('text') or item.get('title') or item.get('desc') or ''
                level = int(item.get('level') or 0)
            else:
                text = _as_text(item)
                level = 0
            p = tf.paragraphs[0] if i == 0 else tf.add_paragraph()
            p.text = text
            p.level = max(0, min(level, 4))
            p.font.size = Pt(20 if p.level == 0 else 16)

    # ============ 原生 PPT 矢量版式 ============

    def _theme(self):
        from pptx.dml.color import RGBColor
        return {
            'primary': RGBColor(37, 99, 235),
            'primary_dark': RGBColor(30, 64, 175),
            'accent': RGBColor(249, 115, 22),
            'bg': RGBColor(248, 250, 252),
            'card': RGBColor(255, 255, 255),
            'line': RGBColor(203, 213, 225),
            'text': RGBColor(15, 23, 42),
            'muted': RGBColor(71, 85, 105),
            'soft_blue': RGBColor(219, 234, 254),
            'soft_orange': RGBColor(255, 237, 213),
            'soft_green': RGBColor(220, 252, 231),
            'soft_purple': RGBColor(237, 233, 254),
            'green': RGBColor(22, 163, 74),
            'purple': RGBColor(124, 58, 237),
        }

    def _blank_slide(self, prs, title):
        from pptx.util import Inches, Pt
        slide = prs.slides.add_slide(prs.slide_layouts[6])
        self._add_background(slide, prs)
        box = slide.shapes.add_textbox(Inches(0.55), Inches(0.3), Inches(12.2), Inches(0.55))
        p = box.text_frame.paragraphs[0]
        p.text = _as_text(title, '未命名页面')
        p.font.size = Pt(28)
        p.font.bold = True
        p.font.color.rgb = self._theme()['text']
        return slide

    def _add_background(self, slide, prs):
        from pptx.enum.shapes import MSO_SHAPE
        from pptx.util import Inches
        theme = self._theme()
        bg = slide.shapes.add_shape(MSO_SHAPE.RECTANGLE, 0, 0, prs.slide_width, prs.slide_height)
        bg.fill.solid()
        bg.fill.fore_color.rgb = theme['bg']
        bg.line.fill.background()
        bar = slide.shapes.add_shape(MSO_SHAPE.RECTANGLE, 0, 0, Inches(0.11), prs.slide_height)
        bar.fill.solid()
        bar.fill.fore_color.rgb = theme['primary']
        bar.line.fill.background()

    def _set_shape_fill(self, shape, color, line_color=None):
        shape.fill.solid()
        shape.fill.fore_color.rgb = color
        if line_color is None:
            shape.line.fill.background()
        else:
            shape.line.color.rgb = line_color

    def _add_textbox(self, slide, text, left, top, width, height, size=16, bold=False, color=None, align=None):
        from pptx.util import Pt
        theme = self._theme()
        box = slide.shapes.add_textbox(left, top, width, height)
        tf = box.text_frame
        tf.word_wrap = True
        tf.margin_left = Pt(2)
        tf.margin_right = Pt(2)
        tf.margin_top = Pt(1)
        tf.margin_bottom = Pt(1)
        p = tf.paragraphs[0]
        p.text = _as_text(text)
        p.font.size = Pt(size)
        p.font.bold = bold
        p.font.color.rgb = color or theme['text']
        if align:
            p.alignment = align
        return box

    def _auto_visualize_slides(self, slides):
        """把模型偷懒生成的纯 bullets，按标题关键词自动提升成矢量版式。

        这不是替代模型设计，而是兜底：用户明明要卡片/流程/时间轴/对比时，
        模型经常仍只给 bullets。这里在不破坏 cover/agenda/summary 的前提下，
        将常见页面转换为更像 PPT 的原生形状页面。
        """
        out = []
        for idx, slide in enumerate(_as_list(slides)):
            if not isinstance(slide, dict):
                out.append(slide)
                continue
            s = dict(slide)
            slide_type = _as_text(s.get('type') or s.get('layout'), 'bullets').lower()
            if slide_type not in ('bullets', 'content', 'text'):
                out.append(s)
                continue

            title = _as_text(s.get('title')).lower()
            bullets = _as_list(s.get('bullets') or s.get('items') or [])
            if len(bullets) < 2:
                out.append(s)
                continue

            if any(k in title for k in ('对比', '比较', '区别', '差异', ' vs ', 'vs')) and len(bullets) >= 2:
                mid = max(1, len(bullets) // 2)
                s['type'] = 'comparison'
                s.setdefault('left', {'title': '传统方式', 'items': bullets[:mid]})
                s.setdefault('right', {'title': 'Agent 方式', 'items': bullets[mid:] or bullets[:mid]})
            elif any(k in title for k in ('流程', '步骤', '链路', '过程', '工作流', '方法')):
                s['type'] = 'process'
                s['steps'] = [self._bullet_to_item(x, i + 1) for i, x in enumerate(bullets[:5])]
            elif any(k in title for k in ('路线', '时间轴', '阶段', '计划', '排期', '里程碑', '演进')):
                s['type'] = 'timeline'
                s['events'] = [self._bullet_to_item(x, i + 1, with_time=True) for i, x in enumerate(bullets[:5])]
            elif any(k in title for k in ('矩阵', '四象限', '象限')):
                s['type'] = 'matrix'
                s['quadrants'] = [self._bullet_to_item(x, i + 1) for i, x in enumerate(bullets[:4])]
            elif any(k in title for k in ('金字塔', '层级', '分层')):
                s['type'] = 'pyramid'
                s['levels'] = bullets[:5]
            elif any(k in title for k in ('闭环', '循环', '飞轮')):
                s['type'] = 'cycle'
                s['steps'] = [self._bullet_to_item(x, i + 1) for i, x in enumerate(bullets[:5])]
            elif any(k in title for k in ('能力', '优势', '价值', '场景', '模块', '特点', '特性')) or (idx > 0 and 3 <= len(bullets) <= 4):
                s['type'] = 'three_cards'
                s['cards'] = [self._bullet_to_item(x, i + 1) for i, x in enumerate(bullets[:4])]
            out.append(s)
        return out

    def _bullet_to_item(self, item, idx=1, with_time=False):
        if isinstance(item, dict):
            return dict(item)
        text = _as_text(item)
        parts = re.split(r'[：:，,。；;\-—]', text, maxsplit=1)
        result = {'title': parts[0].strip() or f'项目 {idx}'}
        if len(parts) > 1 and parts[1].strip():
            result['desc'] = parts[1].strip()
        if with_time:
            result['time'] = f'阶段 {idx}'
        return result

    def _item_title_desc(self, item, idx=None):
        if isinstance(item, dict):
            title = item.get('title') or item.get('text') or item.get('name') or (f'步骤 {idx}' if idx else '')
            desc = item.get('desc') or item.get('description') or item.get('content') or ''
            icon = item.get('icon') or ''
            time = item.get('time') or item.get('date') or item.get('phase') or ''
        else:
            title, desc, icon, time = _as_text(item), '', '', ''
        return _as_text(title), _as_text(desc), _as_text(icon), _as_text(time)

    def _ppt_add_three_cards(self, prs, data):
        from pptx.enum.shapes import MSO_SHAPE
        from pptx.enum.text import PP_ALIGN
        from pptx.util import Inches
        theme = self._theme()
        slide = self._blank_slide(prs, data.get('title'))
        cards = _as_list(data.get('cards') or data.get('items') or data.get('bullets'))[:4]
        if not cards:
            cards = [{'title': '能力', 'desc': '在此补充说明'}]
        n = len(cards)
        gap = Inches(0.28)
        left0 = Inches(0.85)
        top = Inches(1.55)
        total_w = Inches(11.6)
        card_w = int((total_w - gap * (n - 1)) / n)
        card_h = Inches(4.55)
        colors = [theme['soft_blue'], theme['soft_orange'], theme['card'], theme['soft_blue']]
        icons = ['①', '②', '③', '④']
        for i, item in enumerate(cards):
            title, desc, icon, _ = self._item_title_desc(item, i + 1)
            icon = icon or icons[i]
            left = left0 + i * (card_w + gap)
            card = slide.shapes.add_shape(MSO_SHAPE.ROUNDED_RECTANGLE, left, top, card_w, card_h)
            self._set_shape_fill(card, colors[i % len(colors)], theme['line'])
            self._add_textbox(slide, icon, left + Inches(0.2), top + Inches(0.28), card_w - Inches(0.4), Inches(0.62), 30, True, theme['primary'], PP_ALIGN.CENTER)
            self._add_textbox(slide, title, left + Inches(0.22), top + Inches(1.15), card_w - Inches(0.44), Inches(0.55), 20, True, theme['text'], PP_ALIGN.CENTER)
            self._add_textbox(slide, desc, left + Inches(0.28), top + Inches(2.05), card_w - Inches(0.56), Inches(1.65), 14, False, theme['muted'], PP_ALIGN.CENTER)

    def _ppt_add_process(self, prs, data):
        from pptx.enum.shapes import MSO_SHAPE
        from pptx.enum.text import PP_ALIGN
        from pptx.util import Inches
        theme = self._theme()
        slide = self._blank_slide(prs, data.get('title'))
        steps = _as_list(data.get('steps') or data.get('items') or data.get('bullets'))[:5]
        if not steps:
            steps = ['理解目标', '规划步骤', '调用工具', '检查结果']
        n = len(steps)
        left0 = Inches(0.85)
        top = Inches(2.0)
        area_w = Inches(11.45)
        node = Inches(1.18)
        step_gap = int((area_w - node * n) / max(1, n - 1)) if n > 1 else 0
        for i, item in enumerate(steps):
            title, desc, _, _ = self._item_title_desc(item, i + 1)
            x = left0 + i * (node + step_gap)
            if i > 0:
                x1 = left0 + (i - 1) * (node + step_gap) + node
                y = top + node / 2
                arrow = slide.shapes.add_shape(MSO_SHAPE.RIGHT_ARROW, x1 + Inches(0.14), y - Inches(0.08), max(Inches(0.2), x - x1 - Inches(0.28)), Inches(0.16))
                arrow.fill.solid(); arrow.fill.fore_color.rgb = theme['line']
                arrow.line.fill.background()
            circle = slide.shapes.add_shape(MSO_SHAPE.OVAL, x, top, node, node)
            self._set_shape_fill(circle, theme['primary'], theme['primary_dark'])
            self._add_textbox(slide, str(i + 1), x, top + Inches(0.28), node, Inches(0.42), 22, True, theme['card'], PP_ALIGN.CENTER)
            self._add_textbox(slide, title, x - Inches(0.28), top + Inches(1.45), node + Inches(0.56), Inches(0.48), 17, True, theme['text'], PP_ALIGN.CENTER)
            self._add_textbox(slide, desc, x - Inches(0.42), top + Inches(2.02), node + Inches(0.84), Inches(0.9), 12, False, theme['muted'], PP_ALIGN.CENTER)

    def _ppt_add_timeline(self, prs, data):
        from pptx.enum.shapes import MSO_SHAPE
        from pptx.enum.text import PP_ALIGN
        from pptx.util import Inches
        theme = self._theme()
        slide = self._blank_slide(prs, data.get('title'))
        events = _as_list(data.get('events') or data.get('steps') or data.get('items'))[:5]
        if not events:
            events = [{'time': '阶段一', 'title': '启动'}, {'time': '阶段二', 'title': '落地'}, {'time': '阶段三', 'title': '优化'}]
        n = len(events)
        left0 = Inches(1.05)
        right = Inches(12.1)
        y = Inches(3.05)
        line = slide.shapes.add_shape(MSO_SHAPE.RECTANGLE, left0, y - Inches(0.015), right - left0, Inches(0.03))
        line.fill.solid(); line.fill.fore_color.rgb = theme['line']
        line.line.fill.background()
        span = int((right - left0) / max(1, n - 1)) if n > 1 else 0
        for i, item in enumerate(events):
            title, desc, _, time = self._item_title_desc(item, i + 1)
            x = left0 + i * span
            dot = slide.shapes.add_shape(MSO_SHAPE.OVAL, x - Inches(0.16), y - Inches(0.16), Inches(0.32), Inches(0.32))
            self._set_shape_fill(dot, theme['accent'] if i % 2 else theme['primary'], theme['card'])
            top = Inches(1.45) if i % 2 == 0 else Inches(3.45)
            self._add_textbox(slide, time or f'阶段 {i + 1}', x - Inches(0.75), top, Inches(1.5), Inches(0.35), 13, True, theme['primary'], PP_ALIGN.CENTER)
            self._add_textbox(slide, title, x - Inches(0.95), top + Inches(0.42), Inches(1.9), Inches(0.48), 15, True, theme['text'], PP_ALIGN.CENTER)
            self._add_textbox(slide, desc, x - Inches(1.05), top + Inches(0.92), Inches(2.1), Inches(0.68), 11, False, theme['muted'], PP_ALIGN.CENTER)

    def _ppt_add_comparison(self, prs, data):
        from pptx.enum.shapes import MSO_SHAPE
        from pptx.enum.text import PP_ALIGN
        from pptx.util import Inches, Pt
        theme = self._theme()
        slide = self._blank_slide(prs, data.get('title'))
        left_data = data.get('left') if isinstance(data.get('left'), dict) else {'title': '方案 A', 'items': data.get('left') or []}
        right_data = data.get('right') if isinstance(data.get('right'), dict) else {'title': '方案 B', 'items': data.get('right') or []}
        panels = [(left_data, Inches(0.95), theme['soft_blue'], theme['primary']), (right_data, Inches(7.0), theme['soft_orange'], theme['accent'])]
        for panel, left, fill, color in panels:
            box = slide.shapes.add_shape(MSO_SHAPE.ROUNDED_RECTANGLE, left, Inches(1.45), Inches(5.25), Inches(4.75))
            self._set_shape_fill(box, fill, theme['line'])
            self._add_textbox(slide, panel.get('title') or '对比项', left + Inches(0.28), Inches(1.75), Inches(4.69), Inches(0.55), 22, True, color, PP_ALIGN.CENTER)
            items = _as_list(panel.get('items') or panel.get('bullets') or panel.get('points'))[:6]
            tf_box = slide.shapes.add_textbox(left + Inches(0.52), Inches(2.65), Inches(4.25), Inches(2.9))
            tf = tf_box.text_frame
            tf.clear()
            for i, item in enumerate(items or ['']):
                p = tf.paragraphs[0] if i == 0 else tf.add_paragraph()
                p.text = _as_text(item.get('text') if isinstance(item, dict) else item)
                p.font.size = Pt(15)
                p.font.color.rgb = theme['text']
                p.level = 0
        vs = slide.shapes.add_shape(MSO_SHAPE.OVAL, Inches(6.05), Inches(3.08), Inches(0.82), Inches(0.82))
        self._set_shape_fill(vs, theme['primary'], theme['card'])
        self._add_textbox(slide, 'VS', Inches(6.05), Inches(3.29), Inches(0.82), Inches(0.3), 18, True, theme['card'], PP_ALIGN.CENTER)

    def _ppt_add_two_column(self, prs, data):
        from pptx.enum.shapes import MSO_SHAPE
        from pptx.enum.text import PP_ALIGN
        from pptx.util import Inches, Pt
        theme = self._theme()
        slide = self._blank_slide(prs, data.get('title'))

        left_data = data.get('left') if isinstance(data.get('left'), dict) else {'title': data.get('left_title') or '左侧', 'items': data.get('left') or []}
        right_data = data.get('right') if isinstance(data.get('right'), dict) else {'title': data.get('right_title') or '右侧', 'items': data.get('right') or []}
        panels = [
            (left_data, Inches(0.9), theme['soft_blue'], theme['primary']),
            (right_data, Inches(6.85), theme['soft_green'], theme['green']),
        ]
        for panel, left, fill, color in panels:
            box = slide.shapes.add_shape(MSO_SHAPE.ROUNDED_RECTANGLE, left, Inches(1.45), Inches(5.45), Inches(4.85))
            self._set_shape_fill(box, fill, theme['line'])
            self._add_textbox(slide, panel.get('title') or '栏目', left + Inches(0.35), Inches(1.78), Inches(4.75), Inches(0.5), 22, True, color, PP_ALIGN.CENTER)
            items = _as_list(panel.get('items') or panel.get('bullets') or panel.get('points') or panel.get('content'))[:7]
            tf_box = slide.shapes.add_textbox(left + Inches(0.55), Inches(2.65), Inches(4.5), Inches(3.15))
            tf = tf_box.text_frame
            tf.clear()
            for i, item in enumerate(items or ['']) :
                p = tf.paragraphs[0] if i == 0 else tf.add_paragraph()
                p.text = '• ' + _as_text(item.get('text') if isinstance(item, dict) else item)
                p.font.size = Pt(15)
                p.font.color.rgb = theme['text']

    def _ppt_add_matrix(self, prs, data):
        from pptx.enum.shapes import MSO_SHAPE
        from pptx.enum.text import PP_ALIGN
        from pptx.util import Inches
        theme = self._theme()
        slide = self._blank_slide(prs, data.get('title'))
        quadrants = _as_list(data.get('quadrants') or data.get('items') or data.get('cards') or data.get('bullets'))[:4]
        if not quadrants:
            quadrants = ['高价值高频', '高价值低频', '低价值高频', '低价值低频']
        while len(quadrants) < 4:
            quadrants.append({'title': f'象限 {len(quadrants) + 1}', 'desc': ''})

        left0, top0 = Inches(1.45), Inches(1.55)
        cell_w, cell_h = Inches(4.95), Inches(2.15)
        gap = Inches(0.12)
        colors = [theme['soft_blue'], theme['soft_orange'], theme['soft_green'], theme['soft_purple']]
        accent = [theme['primary'], theme['accent'], theme['green'], theme['purple']]
        for i, item in enumerate(quadrants[:4]):
            row, col = divmod(i, 2)
            left = left0 + col * (cell_w + gap)
            top = top0 + row * (cell_h + gap)
            rect = slide.shapes.add_shape(MSO_SHAPE.ROUNDED_RECTANGLE, left, top, cell_w, cell_h)
            self._set_shape_fill(rect, colors[i], theme['line'])
            title, desc, _, _ = self._item_title_desc(item, i + 1)
            self._add_textbox(slide, title, left + Inches(0.28), top + Inches(0.32), cell_w - Inches(0.56), Inches(0.45), 19, True, accent[i], PP_ALIGN.CENTER)
            self._add_textbox(slide, desc, left + Inches(0.4), top + Inches(0.95), cell_w - Inches(0.8), Inches(0.78), 13, False, theme['muted'], PP_ALIGN.CENTER)
        self._add_textbox(slide, data.get('x_label') or '频率 / 可行性', Inches(4.7), Inches(6.08), Inches(3.2), Inches(0.3), 12, False, theme['muted'], PP_ALIGN.CENTER)
        self._add_textbox(slide, data.get('y_label') or '价值 / 影响', Inches(0.35), Inches(3.05), Inches(0.8), Inches(0.4), 12, False, theme['muted'], PP_ALIGN.CENTER)

    def _ppt_add_pyramid(self, prs, data):
        from pptx.enum.shapes import MSO_SHAPE
        from pptx.enum.text import PP_ALIGN
        from pptx.util import Inches
        theme = self._theme()
        slide = self._blank_slide(prs, data.get('title'))
        levels = _as_list(data.get('levels') or data.get('items') or data.get('bullets'))[:5]
        if not levels:
            levels = ['基础能力', '流程编排', '智能决策']
        n = len(levels)
        max_w = Inches(8.8)
        level_h = Inches(0.76)
        gap = Inches(0.08)
        center_x = Inches(6.67)
        bottom = Inches(6.05)
        colors = [theme['primary'], theme['purple'], theme['green'], theme['accent'], theme['primary_dark']]
        for visual_idx, item in enumerate(reversed(levels)):
            # bottom level is widest; upper levels gradually shrink.
            width = max(Inches(2.2), int(max_w * (n - visual_idx) / n))
            left = center_x - width / 2
            top = bottom - (visual_idx + 1) * level_h - visual_idx * gap
            shape = slide.shapes.add_shape(MSO_SHAPE.TRAPEZOID, left, top, width, level_h)
            color = colors[visual_idx % len(colors)]
            self._set_shape_fill(shape, color, theme['card'])
            title, desc, _, _ = self._item_title_desc(item, n - visual_idx)
            text = title if not desc else f'{title}｜{desc}'
            self._add_textbox(slide, text, left + Inches(0.18), top + Inches(0.18), width - Inches(0.36), Inches(0.34), 15, True, theme['card'], PP_ALIGN.CENTER)

    def _ppt_add_cycle(self, prs, data):
        from math import cos, sin, pi
        from pptx.enum.shapes import MSO_SHAPE
        from pptx.enum.text import PP_ALIGN
        from pptx.util import Inches
        theme = self._theme()
        slide = self._blank_slide(prs, data.get('title'))
        steps = _as_list(data.get('steps') or data.get('items') or data.get('bullets'))[:6]
        if not steps:
            steps = ['理解目标', '规划任务', '执行操作', '反馈优化']
        n = len(steps)
        cx, cy = Inches(6.65), Inches(3.75)
        r = Inches(2.0)
        node_w, node_h = Inches(1.55), Inches(0.82)
        colors = [theme['primary'], theme['accent'], theme['green'], theme['purple'], theme['primary_dark'], theme['muted']]
        center = slide.shapes.add_shape(MSO_SHAPE.OVAL, cx - Inches(0.72), cy - Inches(0.72), Inches(1.44), Inches(1.44))
        self._set_shape_fill(center, theme['card'], theme['line'])
        self._add_textbox(slide, data.get('center') or '闭环', cx - Inches(0.62), cy - Inches(0.18), Inches(1.24), Inches(0.36), 17, True, theme['primary'], PP_ALIGN.CENTER)
        for i, item in enumerate(steps):
            angle = -pi / 2 + 2 * pi * i / n
            x = cx + int(r * cos(angle)) - node_w / 2
            y = cy + int(r * sin(angle)) - node_h / 2
            title, desc, _, _ = self._item_title_desc(item, i + 1)
            node = slide.shapes.add_shape(MSO_SHAPE.ROUNDED_RECTANGLE, x, y, node_w, node_h)
            self._set_shape_fill(node, colors[i % len(colors)], theme['card'])
            self._add_textbox(slide, title, x + Inches(0.08), y + Inches(0.2), node_w - Inches(0.16), Inches(0.28), 13, True, theme['card'], PP_ALIGN.CENTER)
            # Tangential arrow between nodes; use non-zero arrow shape to avoid corrupt connector XML.
            next_angle = -pi / 2 + 2 * pi * ((i + 0.5) % n) / n
            ax = cx + int((r + Inches(0.02)) * cos(next_angle)) - Inches(0.28)
            ay = cy + int((r + Inches(0.02)) * sin(next_angle)) - Inches(0.12)
            arrow = slide.shapes.add_shape(MSO_SHAPE.RIGHT_ARROW, ax, ay, Inches(0.56), Inches(0.24))
            self._set_shape_fill(arrow, theme['line'])
            arrow.rotation = int((next_angle + pi / 2) * 180 / pi)

    def _ppt_add_funnel(self, prs, data):
        from pptx.enum.shapes import MSO_SHAPE
        from pptx.enum.text import PP_ALIGN
        from pptx.util import Inches
        theme = self._theme()
        slide = self._blank_slide(prs, data.get('title'))
        stages = _as_list(data.get('stages') or data.get('steps') or data.get('items') or data.get('bullets'))[:5]
        if not stages:
            stages = ['识别场景', '评估价值', '接入工具', '形成闭环']
        n = len(stages)
        center_x = Inches(6.65)
        top0 = Inches(1.45)
        h = Inches(0.82)
        gap = Inches(0.12)
        max_w = Inches(8.6)
        min_w = Inches(3.25)
        colors = [theme['primary'], theme['purple'], theme['accent'], theme['green'], theme['primary_dark']]
        for i, item in enumerate(stages):
            width = max_w - int((max_w - min_w) * i / max(1, n - 1))
            left = center_x - width / 2
            top = top0 + i * (h + gap)
            shape = slide.shapes.add_shape(MSO_SHAPE.TRAPEZOID, left, top, width, h)
            shape.rotation = 180
            self._set_shape_fill(shape, colors[i % len(colors)], theme['card'])
            title, desc, _, _ = self._item_title_desc(item, i + 1)
            text = title if not desc else f'{title}：{desc}'
            self._add_textbox(slide, text, left + Inches(0.25), top + Inches(0.2), width - Inches(0.5), Inches(0.34), 15, True, theme['card'], PP_ALIGN.CENTER)

    def _ppt_add_quote(self, prs, data):
        from pptx.enum.shapes import MSO_SHAPE
        from pptx.enum.text import PP_ALIGN
        from pptx.util import Inches
        theme = self._theme()
        slide = self._blank_slide(prs, data.get('title') or '核心观点')
        quote = data.get('quote') or data.get('content') or data.get('text') or '在这里填写核心观点。'
        author = data.get('author') or data.get('source') or ''
        mark = slide.shapes.add_shape(MSO_SHAPE.OVAL, Inches(1.05), Inches(1.45), Inches(1.0), Inches(1.0))
        self._set_shape_fill(mark, theme['soft_blue'], theme['primary'])
        self._add_textbox(slide, '“', Inches(1.07), Inches(1.43), Inches(0.96), Inches(0.68), 42, True, theme['primary'], PP_ALIGN.CENTER)
        card = slide.shapes.add_shape(MSO_SHAPE.ROUNDED_RECTANGLE, Inches(1.65), Inches(2.0), Inches(10.35), Inches(3.45))
        self._set_shape_fill(card, theme['card'], theme['line'])
        self._add_textbox(slide, quote, Inches(2.2), Inches(2.45), Inches(9.25), Inches(1.65), 26, True, theme['text'], PP_ALIGN.CENTER)
        if author:
            self._add_textbox(slide, f'—— {author}', Inches(6.65), Inches(4.5), Inches(4.6), Inches(0.42), 15, False, theme['muted'], PP_ALIGN.RIGHT)

    def _ppt_add_architecture(self, prs, data):
        from pptx.enum.shapes import MSO_SHAPE
        from pptx.enum.text import PP_ALIGN
        from pptx.util import Inches, Pt
        theme = self._theme()
        slide = self._blank_slide(prs, data.get('title') or '项目架构')
        layers = _as_list(data.get('layers') or data.get('modules') or data.get('items'))[:5]
        if not layers:
            layers = [
                {'title': '数据层', 'items': ['数据采集', '数据标注']},
                {'title': '处理层', 'items': ['清洗去噪', '特征构建']},
                {'title': '模型层', 'items': ['模型训练', '参数优化']},
                {'title': '应用层', 'items': ['结果分析', '系统展示']},
            ]
        n = len(layers)
        left0, top, area_w = Inches(0.75), Inches(1.45), Inches(11.85)
        gap = Inches(0.22)
        box_w = int((area_w - gap * (n - 1)) / n)
        box_h = Inches(4.85)
        colors = [theme['soft_blue'], theme['soft_green'], theme['soft_orange'], theme['soft_purple'], theme['card']]
        accents = [theme['primary'], theme['green'], theme['accent'], theme['purple'], theme['primary_dark']]
        for i, layer in enumerate(layers):
            title, desc, _, _ = self._item_title_desc(layer, i + 1)
            items = _as_list(layer.get('items') or layer.get('modules') or layer.get('components') if isinstance(layer, dict) else [])[:5]
            left = left0 + i * (box_w + gap)
            rect = slide.shapes.add_shape(MSO_SHAPE.ROUNDED_RECTANGLE, left, top, box_w, box_h)
            self._set_shape_fill(rect, colors[i % len(colors)], theme['line'])
            self._add_textbox(slide, title, left + Inches(0.12), top + Inches(0.25), box_w - Inches(0.24), Inches(0.48), 18, True, accents[i % len(accents)], PP_ALIGN.CENTER)
            if desc:
                self._add_textbox(slide, desc, left + Inches(0.2), top + Inches(0.78), box_w - Inches(0.4), Inches(0.38), 11, False, theme['muted'], PP_ALIGN.CENTER)
            tf_box = slide.shapes.add_textbox(left + Inches(0.22), top + Inches(1.3), box_w - Inches(0.44), Inches(2.85))
            tf = tf_box.text_frame; tf.clear(); tf.word_wrap = True
            for j, item in enumerate(items or ['模块说明']):
                p = tf.paragraphs[0] if j == 0 else tf.add_paragraph()
                p.text = '• ' + _item_text(item)
                p.font.size = Pt(12)
                p.font.color.rgb = theme['text']
            if i > 0:
                arrow = slide.shapes.add_shape(MSO_SHAPE.RIGHT_ARROW, left - gap + Inches(0.03), top + Inches(2.22), gap - Inches(0.06), Inches(0.32))
                self._set_shape_fill(arrow, theme['line'])

    def _ppt_add_method_pipeline(self, prs, data):
        from pptx.enum.shapes import MSO_SHAPE
        from pptx.enum.text import PP_ALIGN
        from pptx.util import Inches, Pt
        theme = self._theme()
        slide = self._blank_slide(prs, data.get('title') or '研究方法流程')
        stages = _as_list(data.get('stages') or data.get('steps') or data.get('items'))[:5]
        if not stages:
            stages = [
                {'title': '数据采集', 'items': ['实验数据', '公开数据集']},
                {'title': '预处理', 'items': ['去噪', '归一化']},
                {'title': '模型训练', 'items': ['参数优化', '交叉验证']},
                {'title': '性能评估', 'items': ['指标对比', '消融实验']},
            ]
        n = len(stages)
        left0, top = Inches(0.8), Inches(1.55)
        area_w, gap = Inches(11.75), Inches(0.28)
        card_w = int((area_w - gap * (n - 1)) / n)
        for i, stage in enumerate(stages):
            title, desc, _, _ = self._item_title_desc(stage, i + 1)
            items = _as_list(stage.get('items') or stage.get('substeps') or stage.get('tasks') if isinstance(stage, dict) else [])[:4]
            left = left0 + i * (card_w + gap)
            head = slide.shapes.add_shape(MSO_SHAPE.CHEVRON, left, top, card_w, Inches(0.78))
            self._set_shape_fill(head, theme['primary'] if i % 2 == 0 else theme['accent'], theme['card'])
            self._add_textbox(slide, f'{i + 1}. {title}', left + Inches(0.08), top + Inches(0.22), card_w - Inches(0.18), Inches(0.28), 13, True, theme['card'], PP_ALIGN.CENTER)
            body = slide.shapes.add_shape(MSO_SHAPE.ROUNDED_RECTANGLE, left + Inches(0.05), top + Inches(1.02), card_w - Inches(0.1), Inches(3.75))
            self._set_shape_fill(body, theme['card'], theme['line'])
            if desc:
                self._add_textbox(slide, desc, left + Inches(0.2), top + Inches(1.28), card_w - Inches(0.4), Inches(0.45), 11, False, theme['muted'], PP_ALIGN.CENTER)
            tf_box = slide.shapes.add_textbox(left + Inches(0.22), top + Inches(1.85), card_w - Inches(0.44), Inches(2.4))
            tf = tf_box.text_frame; tf.clear(); tf.word_wrap = True
            for j, item in enumerate(items or ['关键步骤']):
                p = tf.paragraphs[0] if j == 0 else tf.add_paragraph()
                p.text = '• ' + _item_text(item)
                p.font.size = Pt(11)
                p.font.color.rgb = theme['text']

    def _ppt_add_table(self, prs, data):
        from pptx.enum.text import PP_ALIGN
        from pptx.util import Inches, Pt
        theme = self._theme()
        slide = self._blank_slide(prs, data.get('title') or '数据表格')
        headers = _as_list(data.get('headers') or data.get('columns'))[:6]
        rows = _as_list(data.get('rows') or data.get('data'))[:8]
        if not headers:
            headers = ['方法', '指标A', '指标B', '结论']
        if rows and isinstance(rows[0], dict):
            rows = _rows_from_records(rows, headers)
        if not rows:
            rows = [['Baseline', '82.1', '79.4', '参考'], ['Ours', '88.6', '85.2', '最佳']]
        table_shape = slide.shapes.add_table(len(rows) + 1, len(headers), Inches(0.85), Inches(1.55), Inches(11.65), Inches(4.9))
        table = table_shape.table
        for c, h in enumerate(headers):
            cell = table.cell(0, c)
            cell.text = _as_text(h)
            cell.fill.solid(); cell.fill.fore_color.rgb = theme['primary']
            p = cell.text_frame.paragraphs[0]; p.font.size = Pt(12); p.font.bold = True; p.font.color.rgb = theme['card']; p.alignment = PP_ALIGN.CENTER
        for r, row in enumerate(rows, start=1):
            vals = _as_list(row)
            for c in range(len(headers)):
                cell = table.cell(r, c)
                cell.text = _as_text(vals[c] if c < len(vals) else '')
                cell.fill.solid(); cell.fill.fore_color.rgb = theme['soft_blue'] if r % 2 else theme['card']
                p = cell.text_frame.paragraphs[0]; p.font.size = Pt(11); p.font.color.rgb = theme['text']; p.alignment = PP_ALIGN.CENTER

    def _ppt_add_chart(self, prs, data):
        from pptx.enum.shapes import MSO_SHAPE
        from pptx.enum.text import PP_ALIGN
        from pptx.util import Inches
        theme = self._theme()
        slide = self._blank_slide(prs, data.get('title') or '数据图表')
        chart_type = _as_text(data.get('chart_type') or data.get('kind') or 'bar').lower()
        categories = _as_list(data.get('categories') or data.get('labels'))[:8] or ['A', 'B', 'C', 'D']
        values = _as_list(data.get('values') or data.get('data'))[:len(categories)] or [30, 55, 42, 70]
        nums = []
        for v in values:
            try:
                nums.append(float(v.get('value') if isinstance(v, dict) else v))
            except Exception:
                nums.append(0.0)
        max_v = max(nums) if nums else 1.0
        max_v = max_v or 1.0
        if chart_type == 'line':
            self._add_textbox(slide, '折线图（节点表示趋势）', Inches(0.9), Inches(1.35), Inches(3.5), Inches(0.35), 12, False, theme['muted'], PP_ALIGN.LEFT)
            left0, bottom = Inches(1.2), Inches(5.75)
            area_w, area_h = Inches(10.7), Inches(3.9)
            prev = None
            for i, (cat, val) in enumerate(zip(categories, nums)):
                x = left0 + int(area_w * i / max(1, len(nums) - 1))
                y = bottom - int(area_h * val / max_v)
                if prev:
                    px, py = prev
                    dx, dy = x - px, y - py
                    line = slide.shapes.add_shape(MSO_SHAPE.RECTANGLE, px, py, max(Inches(0.02), dx), Inches(0.035))
                    self._set_shape_fill(line, theme['primary'])
                    if dx:
                        line.rotation = 0
                dot = slide.shapes.add_shape(MSO_SHAPE.OVAL, x - Inches(0.08), y - Inches(0.08), Inches(0.16), Inches(0.16))
                self._set_shape_fill(dot, theme['accent'], theme['card'])
                self._add_textbox(slide, _as_text(cat), x - Inches(0.35), bottom + Inches(0.15), Inches(0.7), Inches(0.25), 10, False, theme['muted'], PP_ALIGN.CENTER)
                prev = (x, y)
        else:
            left0, bottom = Inches(1.0), Inches(5.85)
            area_w, area_h = Inches(11.2), Inches(4.0)
            gap = Inches(0.14)
            bar_w = int((area_w - gap * (len(nums) - 1)) / max(1, len(nums)))
            for i, (cat, val) in enumerate(zip(categories, nums)):
                h = max(Inches(0.08), int(area_h * val / max_v))
                left = left0 + i * (bar_w + gap)
                bar = slide.shapes.add_shape(MSO_SHAPE.ROUNDED_RECTANGLE, left, bottom - h, bar_w, h)
                self._set_shape_fill(bar, theme['primary'] if i % 2 == 0 else theme['accent'])
                self._add_textbox(slide, str(val).rstrip('0').rstrip('.'), left, bottom - h - Inches(0.35), bar_w, Inches(0.25), 10, True, theme['text'], PP_ALIGN.CENTER)
                self._add_textbox(slide, _as_text(cat), left, bottom + Inches(0.12), bar_w, Inches(0.28), 10, False, theme['muted'], PP_ALIGN.CENTER)

    def _ppt_add_experiment_design(self, prs, data):
        from pptx.enum.shapes import MSO_SHAPE
        from pptx.enum.text import PP_ALIGN
        from pptx.util import Inches, Pt
        theme = self._theme()
        slide = self._blank_slide(prs, data.get('title') or '实验设计')
        groups = _as_list(data.get('groups') or data.get('experiments') or data.get('items'))[:4]
        if not groups:
            groups = [{'name': 'Baseline', 'method': '传统方法', 'metrics': ['Accuracy', 'F1']}, {'name': 'Ours', 'method': '改进模型', 'metrics': ['Accuracy', 'F1']}]
        n = len(groups)
        left0, top, area_w, gap = Inches(0.9), Inches(1.55), Inches(11.55), Inches(0.25)
        card_w = int((area_w - gap * (n - 1)) / n)
        for i, g in enumerate(groups):
            left = left0 + i * (card_w + gap)
            card = slide.shapes.add_shape(MSO_SHAPE.ROUNDED_RECTANGLE, left, top, card_w, Inches(4.85))
            self._set_shape_fill(card, theme['soft_blue'] if i % 2 == 0 else theme['soft_green'], theme['line'])
            name = _as_text(g.get('name') or g.get('title') or f'实验组 {i + 1}' if isinstance(g, dict) else g)
            method = _as_text(g.get('method') or g.get('desc') or '' if isinstance(g, dict) else '')
            metrics = _as_list(g.get('metrics') or g.get('indicators') or [] if isinstance(g, dict) else [])[:5]
            self._add_textbox(slide, name, left + Inches(0.18), top + Inches(0.28), card_w - Inches(0.36), Inches(0.45), 18, True, theme['primary'], PP_ALIGN.CENTER)
            self._add_textbox(slide, method or '实验方法', left + Inches(0.25), top + Inches(0.95), card_w - Inches(0.5), Inches(0.65), 13, False, theme['text'], PP_ALIGN.CENTER)
            tf_box = slide.shapes.add_textbox(left + Inches(0.35), top + Inches(1.95), card_w - Inches(0.7), Inches(2.1))
            tf = tf_box.text_frame; tf.clear(); tf.word_wrap = True
            for j, m in enumerate(metrics or ['评价指标']):
                p = tf.paragraphs[0] if j == 0 else tf.add_paragraph()
                p.text = '✓ ' + _item_text(m)
                p.font.size = Pt(12)
                p.font.color.rgb = theme['text']

    def _ppt_add_ablation(self, prs, data):
        headers = data.get('headers') or ['实验设置', '去除模块', '关键指标', '结论']
        rows = data.get('rows') or data.get('variants') or data.get('items') or [
            ['Full Model', '-', '88.6', '最佳'],
            ['w/o Module A', '模块A', '84.2', '下降明显'],
            ['w/o Module B', '模块B', '85.1', '有一定影响'],
        ]
        self._ppt_add_table(prs, {'title': data.get('title') or '消融实验', 'headers': headers, 'rows': rows})
