# ============================================================
# server/ppt.py - PPT 生成 mixin
# ============================================================
# 提供 PptMixin，给 Handler 用。MVP 能力：
#   - 接收结构化 deck/slides JSON
#   - 生成 .pptx 文件到 output/ 目录或用户指定的沙箱内路径
#   - 支持 cover / agenda / bullets / section / summary 基础版式
#   - 支持 three_cards / process / timeline / comparison / two_column / matrix
#     / pyramid / cycle / funnel / quote / architecture / method_pipeline
#     / table / chart / experiment_design / ablation 原生 PPT 矢量版式
# ============================================================

import math
import html
import json
import os
import re
import shutil
import time
import zipfile
import hashlib
from collections import Counter
from io import BytesIO
from urllib.parse import quote
from xml.etree import ElementTree as ET

from . import config
from .ppt_pipeline import build_deck_spec
from .sandbox import check_path_or_error


_SAFE_FILENAME_RE = re.compile(r'[\\/:*?"<>|]+')
_STYLE_KEYS = {
    'font', 'font_family', 'font_name', 'font_size', 'size',
    'color', 'font_color', 'text_color',
    'bold', 'italic', 'underline',
    'align', 'alignment',
    'line_spacing',
    'primary_color', 'accent_color', 'background_color', 'bg_color',
    'text_color', 'muted_color', 'card_color', 'line_color'
}

_EMU_PER_INCH = 914400.0


def _emu_to_in(value):
    try:
        return round(int(value) / _EMU_PER_INCH, 3)
    except Exception:
        return 0


def _in_to_emu(value):
    try:
        return int(float(value) * _EMU_PER_INCH)
    except Exception:
        return 0


def _as_text(value, default=''):
    if value is None:
        return default
    return str(value)


def _as_list(value):
    if isinstance(value, list):
        return value
    if value is None:
        return []
    return [value]


def _item_text(value):
    if isinstance(value, dict):
        return _as_text(value.get('text') or value.get('title') or value.get('name') or value.get('desc') or value.get('value'))
    return _as_text(value)


def _rows_from_records(records, headers):
    rows = []
    for rec in _as_list(records):
        rows.append([_as_text(rec.get(h, '') if isinstance(rec, dict) else rec) for h in headers])
    return rows


def _safe_filename(name):
    name = _as_text(name, 'generated.pptx').strip() or 'generated.pptx'
    name = _SAFE_FILENAME_RE.sub('_', name)
    if not name.lower().endswith('.pptx'):
        name += '.pptx'
    return name


class PptMixin:
    """Handler mixin：生成 PPTX 文件。"""

    def handle_generate_ppt(self, body):
        self._ppt_global_style = {}
        self._ppt_current_slide = {}
        self._ppt_pipeline_spec = {}
        try:
            data = body.get('data') or {}
            if not isinstance(data, dict):
                return self._send_json(200, {'ok': False, 'error': 'data 必须是对象'})

            # High-level pipeline mode: allow callers to pass a user request and
            # let the planner produce the structured slides consumed below.
            # Existing structured JSON remains fully supported.
            if (data.get('user_request') or data.get('request') or data.get('prompt')) and not data.get('slides'):
                request_text = data.get('user_request') or data.get('request') or data.get('prompt')
                deck_spec = build_deck_spec(request_text, data)
                data = deck_spec
                self._ppt_pipeline_spec = deck_spec

            try:
                from pptx import Presentation
                from pptx.util import Inches, Pt
            except ImportError:
                return self._send_json(200, {
                    'ok': False,
                    'error': '缺少依赖 python-pptx，请先运行：pip install -r requirements.txt'
                })

            template_path_value = data.get('template_path') or data.get('template')
            template_mode = _as_text(data.get('template_mode') or data.get('template_strategy') or '', '').strip().lower()
            # 用户传入模板 PPT 时，默认复制模板并替换占位文字，保留母版、背景、Logo、页眉页脚和固定装饰。
            # 如只想抽取颜色/字体后重新生成矢量版式，可显式传 template_mode='style'。
            if template_path_value and template_mode not in ('style', 'profile', 'extract'):
                return self._generate_ppt_from_template(data, Presentation)

            template_profile = self._load_ppt_template_profile(data)
            self._ppt_template_profile = template_profile if isinstance(template_profile, dict) else {}
            self._ppt_global_style = {}
            if template_profile:
                self._ppt_global_style.update(self._template_profile_style(template_profile))
            self._ppt_global_style.update(self._style_dict(data.get('theme')))
            self._ppt_global_style.update(self._style_dict(data.get('style')))
            filename = _safe_filename(data.get('filename') or data.get('file_name'))
            output_path = data.get('path') or data.get('output_path') or os.path.join('output', filename)
            if str(output_path).lower().endswith(os.sep) or str(output_path).endswith('/'):
                output_path = os.path.join(output_path, filename)
            if not str(output_path).lower().endswith('.pptx'):
                output_path = os.path.join(str(output_path), filename)

            out_path, err = check_path_or_error(output_path, must_exist=False)
            if err:
                return self._send_json(200, {'ok': False, 'error': err})

            slides = _as_list(data.get('slides'))
            if not slides:
                slides = [{
                    'type': 'cover',
                    'title': data.get('title') or '未命名 PPT',
                    'subtitle': data.get('subtitle') or ''
                }]

            slides = self._auto_visualize_slides(slides)

            prs = Presentation()
            if template_profile and template_profile.get('slide_width') and template_profile.get('slide_height'):
                prs.slide_width = int(template_profile.get('slide_width'))
                prs.slide_height = int(template_profile.get('slide_height'))
            else:
                # python-pptx 默认是 10x7.5（4:3），而矢量版式按 16:9 设计。
                # 不显式设置会导致所有 12.x 英寸坐标横向越界。
                prs.slide_width = Inches(13.333333)
                prs.slide_height = Inches(7.5)
            for slide_data in slides:
                if not isinstance(slide_data, dict):
                    slide_data = {'type': 'bullets', 'title': _as_text(slide_data)}
                slide_type = _as_text(slide_data.get('type') or slide_data.get('layout'), 'bullets').lower()
                if slide_type == 'cover':
                    self._ppt_add_cover(prs, slide_data)
                elif slide_type in ('agenda', 'toc'):
                    bullet_data = dict(slide_data)
                    bullet_data.update({
                        'title': slide_data.get('title') or '目录',
                        'bullets': slide_data.get('items') or slide_data.get('bullets') or []
                    })
                    self._ppt_add_bullets(prs, bullet_data, Pt)
                elif slide_type in ('section', 'divider'):
                    self._ppt_add_section(prs, slide_data)
                elif slide_type == 'summary':
                    bullet_data = dict(slide_data)
                    bullet_data.update({
                        'title': slide_data.get('title') or '总结',
                        'bullets': slide_data.get('bullets') or slide_data.get('items') or []
                    })
                    self._ppt_add_bullets(prs, bullet_data, Pt)
                elif slide_type in ('three_cards', 'cards'):
                    self._ppt_add_three_cards(prs, slide_data)
                elif slide_type in ('process', 'flow', 'workflow'):
                    self._ppt_add_process(prs, slide_data)
                elif slide_type == 'timeline':
                    self._ppt_add_timeline(prs, slide_data)
                elif slide_type in ('comparison', 'compare', 'vs'):
                    self._ppt_add_comparison(prs, slide_data)
                elif slide_type in ('two_column', 'two_columns', 'columns'):
                    self._ppt_add_two_column(prs, slide_data)
                elif slide_type in ('matrix', 'quadrant', 'four_quadrant'):
                    self._ppt_add_matrix(prs, slide_data)
                elif slide_type == 'pyramid':
                    self._ppt_add_pyramid(prs, slide_data)
                elif slide_type in ('cycle', 'loop'):
                    self._ppt_add_cycle(prs, slide_data)
                elif slide_type == 'funnel':
                    self._ppt_add_funnel(prs, slide_data)
                elif slide_type in ('quote', 'quotation'):
                    self._ppt_add_quote(prs, slide_data)
                elif slide_type in ('architecture', 'system_architecture', 'tech_architecture'):
                    self._ppt_add_architecture(prs, slide_data)
                elif slide_type in ('method_pipeline', 'research_pipeline', 'pipeline'):
                    self._ppt_add_method_pipeline(prs, slide_data)
                elif slide_type in ('table', 'data_table'):
                    self._ppt_add_table(prs, slide_data)
                elif slide_type in ('chart', 'bar_chart', 'line_chart'):
                    self._ppt_add_chart(prs, slide_data)
                elif slide_type in ('experiment_design', 'experiment', 'experimental_design'):
                    self._ppt_add_experiment_design(prs, slide_data)
                elif slide_type in ('ablation', 'ablation_study'):
                    self._ppt_add_ablation(prs, slide_data)
                else:
                    self._ppt_add_bullets(prs, slide_data, Pt)

            parent = os.path.dirname(out_path)
            if parent:
                os.makedirs(parent, exist_ok=True)
            prs.save(out_path)

            repair_info = None
            validation = None
            if self._ppt_bool(data.get('auto_repair_text_layout'), True):
                repair_rules = self._ppt_validation_rules(data)
                repair_rules.setdefault('use_powerpoint_com_text_bounds', True)
                repair_info = self._ppt_auto_repair_text_layout(out_path, repair_rules)

            if self._ppt_bool(data.get('validate_after_generate'), True):
                validation_rules = self._ppt_validation_rules(data)
                validation_rules.setdefault('use_powerpoint_com_text_bounds', True)
                validation = self._validate_pptx_file(out_path, validation_rules)
                if (not validation.get('passed')) and self._ppt_bool(data.get('auto_repair_text_layout'), True):
                    retry_rules = dict(validation_rules)
                    retry_rules['auto_repair_passes'] = max(self._ppt_int(retry_rules.get('auto_repair_passes'), 3, 1, 8), 6)
                    retry_rules['auto_repair_min_font_size'] = min(self._ppt_int(retry_rules.get('auto_repair_min_font_size'), 8, 5, 18), 6)
                    retry_info = self._ppt_auto_repair_text_layout(out_path, retry_rules)
                    if repair_info:
                        repair_info['retry'] = retry_info
                    else:
                        repair_info = retry_info
                    validation = self._validate_pptx_file(out_path, validation_rules)

            rel_path = os.path.relpath(out_path, config.WORKSPACE_ROOT).replace(os.sep, '/')
            validation_failed = bool(validation and not validation.get('passed'))
            response = {
                'ok': not validation_failed,
                'path': rel_path,
                'slides': len(slides),
                'template_profile': (template_profile.get('profile_path') or template_profile.get('source')) if template_profile else '',
                'pipeline': self._ppt_pipeline_spec.get('pipeline', {}),
                'message': f'PPT 已生成：{rel_path}'
            }
            if repair_info:
                response['auto_repair'] = repair_info
                if repair_info.get('repaired'):
                    response['message'] = f'PPT 已生成并自动修复文本排版：{rel_path}'
            if validation:
                response['validation'] = {
                    'passed': bool(validation.get('passed')),
                    'score': validation.get('score'),
                    'stats': validation.get('stats', {}),
                    'issues': validation.get('issues', [])[:10],
                    'warnings': validation.get('warnings', [])[:10],
                }
                if validation_failed:
                    response['error'] = 'PPT 生成后验证未通过，请减少内容或更换更宽松版式'
                    response['message'] = f'PPT 已生成但验证未通过：{rel_path}'
            self._send_json(200, response)
        except Exception as e:
            self._send_json(200, {'ok': False, 'error': str(e)})

    def handle_validate_ppt(self, body):
        try:
            data = self._ppt_request_data(body)
            path_value = data.get('path') or data.get('pptx') or data.get('file') or data.get('file_path')
            if not path_value:
                return self._send_json(200, {'ok': False, 'error': 'path is required'})

            ppt_path, err = check_path_or_error(path_value, must_exist=True)
            if err:
                return self._send_json(200, {'ok': False, 'error': err})
            if not os.path.isfile(ppt_path):
                return self._send_json(200, {'ok': False, 'error': f'not a file: {path_value}'})

            rules = self._ppt_validation_rules(data)
            result = self._validate_pptx_file(ppt_path, rules)
            self._send_json(200, result)
        except Exception as e:
            self._send_json(200, {'ok': False, 'error': str(e)})

    def handle_preview_ppt(self, body):
        try:
            data = self._ppt_request_data(body)
            path_value = data.get('path') or data.get('pptx') or data.get('file') or data.get('file_path')
            if not path_value:
                return self._send_json(200, {'ok': False, 'error': 'path is required'})

            ppt_path, err = check_path_or_error(path_value, must_exist=True)
            if err:
                return self._send_json(200, {'ok': False, 'error': err})
            if not os.path.isfile(ppt_path):
                return self._send_json(200, {'ok': False, 'error': f'not a file: {path_value}'})

            base = os.path.splitext(os.path.basename(ppt_path))[0] or 'deck'
            base = _SAFE_FILENAME_RE.sub('_', base).strip('._ ') or 'deck'
            output_dir = data.get('output_dir') or data.get('out_dir') or os.path.join('output', 'ppt_preview', base)
            out_dir, err = check_path_or_error(output_dir, must_exist=False)
            if err:
                return self._send_json(200, {'ok': False, 'error': err})
            os.makedirs(out_dir, exist_ok=True)

            rules = self._ppt_validation_rules(data, include_top_level=False)
            validation = self._validate_pptx_file(ppt_path, rules)
            if not validation.get('ok'):
                return self._send_json(200, validation)

            width = self._ppt_int(data.get('width'), 1280, 320, 4096)
            height = self._ppt_int(data.get('height'), 720, 180, 4096)
            max_slides = self._ppt_int(data.get('max_slides'), 0, 0, 500)
            max_slides = max_slides or None
            prefer_native = self._ppt_bool(data.get('prefer_native'), True)

            renderer = ''
            renderer_note = ''
            image_paths = []
            if prefer_native:
                image_paths, renderer_note = self._render_ppt_preview_powerpoint(
                    ppt_path, out_dir, width, height, max_slides
                )
                if image_paths:
                    renderer = 'powerpoint'

            if not image_paths:
                image_paths, fallback_note = self._render_ppt_preview_pillow(
                    ppt_path, out_dir, width, height, max_slides
                )
                renderer = 'pillow' if image_paths else 'none'
                renderer_note = renderer_note or fallback_note

            images = []
            for idx, img_path in enumerate(image_paths, start=1):
                rel = self._ppt_rel_path(img_path)
                images.append({
                    'slide': idx,
                    'path': rel,
                    'url': self._ppt_preview_url(rel),
                })

            html_path = self._write_ppt_preview_html(
                out_dir,
                self._ppt_rel_path(ppt_path),
                validation,
                images,
                renderer,
                renderer_note,
            )
            html_rel = self._ppt_rel_path(html_path)

            self._send_json(200, {
                'ok': True,
                'passed': validation.get('passed', False),
                'path': self._ppt_rel_path(ppt_path),
                'preview_dir': self._ppt_rel_path(out_dir),
                'html': html_rel,
                'preview_url': self._ppt_preview_url(html_rel),
                'images': images,
                'renderer': renderer,
                'renderer_note': renderer_note,
                'validation': validation,
            })
        except Exception as e:
            self._send_json(200, {'ok': False, 'error': str(e)})

    def handle_analyze_ppt_template(self, body):
        try:
            data = self._ppt_request_data(body)
            path_value = data.get('path') or data.get('template_path') or data.get('pptx') or data.get('file')
            if not path_value:
                return self._send_json(200, {'ok': False, 'error': 'path is required'})

            ppt_path, err = check_path_or_error(path_value, must_exist=True)
            if err:
                return self._send_json(200, {'ok': False, 'error': err})
            if not os.path.isfile(ppt_path):
                return self._send_json(200, {'ok': False, 'error': f'not a file: {path_value}'})

            profile = self._analyze_ppt_template_file(ppt_path)
            save_profile = self._ppt_bool(data.get('save_profile'), True)
            profile_path = ''
            if save_profile:
                profile_path = data.get('profile_path') or data.get('output_path') or self._ppt_default_profile_path(ppt_path)
                profile_path, err = check_path_or_error(profile_path, must_exist=False)
                if err:
                    return self._send_json(200, {'ok': False, 'error': err})
                os.makedirs(os.path.dirname(profile_path), exist_ok=True)
                with open(profile_path, 'w', encoding='utf-8') as f:
                    json.dump(profile, f, ensure_ascii=False, indent=2)

            out = {'ok': True, 'profile': profile}
            if profile_path:
                out['profile_path'] = self._ppt_rel_path(profile_path)
            self._send_json(200, out)
        except Exception as e:
            self._send_json(200, {'ok': False, 'error': str(e)})

    def _generate_ppt_from_template(self, data, Presentation):
        template_value = data.get('template_path') or data.get('template')
        template_path, err = check_path_or_error(template_value, must_exist=True)
        if err:
            return self._send_json(200, {'ok': False, 'error': err})
        if not os.path.isfile(template_path) or os.path.splitext(template_path)[1].lower() != '.pptx':
            return self._send_json(200, {'ok': False, 'error': f'template_path must be a .pptx file: {template_value}'})

        filename = _safe_filename(data.get('filename') or data.get('file_name') or ('generated_from_' + os.path.basename(template_path)))
        output_path = data.get('path') or data.get('output_path') or os.path.join('output', filename)
        if str(output_path).lower().endswith(os.sep) or str(output_path).endswith('/'):
            output_path = os.path.join(output_path, filename)
        if not str(output_path).lower().endswith('.pptx'):
            output_path = os.path.join(str(output_path), filename)
        out_path, err = check_path_or_error(output_path, must_exist=False)
        if err:
            return self._send_json(200, {'ok': False, 'error': err})
        parent = os.path.dirname(out_path)
        if parent:
            os.makedirs(parent, exist_ok=True)

        shutil.copyfile(template_path, out_path)
        prs = Presentation(out_path)
        slides = _as_list(data.get('slides'))
        if not slides:
            slides = [{'title': data.get('title') or '', 'subtitle': data.get('subtitle') or '', 'content': data.get('content') or ''}]

        replacements = data.get('replacements') or data.get('placeholder_values') or data.get('placeholders') or {}
        if not isinstance(replacements, dict):
            replacements = {}
        profile = self._analyze_ppt_template_file(template_path)
        self._ppt_template_profile = profile if isinstance(profile, dict) else {}
        self._ppt_global_style = {}
        if self._ppt_template_profile:
            self._ppt_global_style.update(self._template_profile_style(self._ppt_template_profile))
        self._ppt_global_style.update(self._style_dict(data.get('theme')))
        self._ppt_global_style.update(self._style_dict(data.get('style')))

        trimmed = 0
        target_slide_count = max(1, len(slides))
        if len(prs.slides) > target_slide_count and self._ppt_bool(data.get('trim_extra_template_slides'), True):
            trimmed = self._trim_template_slides(prs, target_slide_count)

        summary = self._fill_template_presentation(prs, slides, data, replacements)
        if trimmed:
            summary['trimmed_template_slides'] = trimmed
        if len(slides) > len(prs.slides):
            extra_slides = slides[len(prs.slides):]
            summary['generated_extra_slides'] = self._append_generated_template_slides(prs, extra_slides)
        prs.save(out_path)

        profile_path = ''
        if self._ppt_bool(data.get('save_profile'), True):
            cache_path = data.get('profile_path') or self._ppt_default_profile_path(template_path)
            abs_cache, cache_err = check_path_or_error(cache_path, must_exist=False)
            if not cache_err:
                os.makedirs(os.path.dirname(abs_cache), exist_ok=True)
                with open(abs_cache, 'w', encoding='utf-8') as f:
                    json.dump(profile, f, ensure_ascii=False, indent=2)
                profile_path = self._ppt_rel_path(abs_cache)

        rel_path = self._ppt_rel_path(out_path)
        return self._send_json(200, {
            'ok': True,
            'path': rel_path,
            'slides': len(prs.slides),
            'template_mode': 'fill',
            'template_path': self._ppt_rel_path(template_path),
            'profile_path': profile_path,
            'filled': summary,
            'message': f'PPT 已基于模板生成：{rel_path}'
        })

    def _trim_template_slides(self, prs, target_count):
        target_count = max(1, int(target_count or 1))
        removed = 0
        while len(prs.slides) > target_count:
            try:
                slide_id = list(prs.slides._sldIdLst)[-1]
                r_id = slide_id.rId
                prs.part.drop_rel(r_id)
                prs.slides._sldIdLst.remove(slide_id)
                removed += 1
            except Exception:
                break
        return removed

    def _append_generated_template_slides(self, prs, slides):
        from pptx.util import Pt
        added = 0
        for slide_data in self._auto_visualize_slides(_as_list(slides)):
            if not isinstance(slide_data, dict):
                slide_data = {'type': 'bullets', 'title': _as_text(slide_data)}
            slide_type = _as_text(slide_data.get('type') or slide_data.get('layout'), 'bullets').lower()
            if slide_type == 'cover':
                self._ppt_add_cover(prs, slide_data)
            elif slide_type in ('agenda', 'toc'):
                bullet_data = dict(slide_data)
                bullet_data.update({
                    'title': slide_data.get('title') or '目录',
                    'bullets': slide_data.get('items') or slide_data.get('bullets') or []
                })
                self._ppt_add_bullets(prs, bullet_data, Pt)
            elif slide_type in ('section', 'divider'):
                self._ppt_add_section(prs, slide_data)
            elif slide_type == 'summary':
                bullet_data = dict(slide_data)
                bullet_data.update({
                    'title': slide_data.get('title') or '总结',
                    'bullets': slide_data.get('bullets') or slide_data.get('items') or []
                })
                self._ppt_add_bullets(prs, bullet_data, Pt)
            elif slide_type in ('three_cards', 'cards'):
                self._ppt_add_three_cards(prs, slide_data)
            elif slide_type in ('process', 'flow', 'workflow'):
                self._ppt_add_process(prs, slide_data)
            elif slide_type == 'timeline':
                self._ppt_add_timeline(prs, slide_data)
            elif slide_type in ('comparison', 'compare', 'vs'):
                self._ppt_add_comparison(prs, slide_data)
            elif slide_type in ('two_column', 'two_columns', 'columns'):
                self._ppt_add_two_column(prs, slide_data)
            elif slide_type in ('matrix', 'quadrant', 'four_quadrant'):
                self._ppt_add_matrix(prs, slide_data)
            elif slide_type == 'pyramid':
                self._ppt_add_pyramid(prs, slide_data)
            elif slide_type in ('cycle', 'loop'):
                self._ppt_add_cycle(prs, slide_data)
            elif slide_type == 'funnel':
                self._ppt_add_funnel(prs, slide_data)
            elif slide_type in ('quote', 'quotation'):
                self._ppt_add_quote(prs, slide_data)
            elif slide_type in ('architecture', 'system_architecture', 'tech_architecture'):
                self._ppt_add_architecture(prs, slide_data)
            elif slide_type in ('method_pipeline', 'research_pipeline', 'pipeline'):
                self._ppt_add_method_pipeline(prs, slide_data)
            elif slide_type in ('table', 'data_table'):
                self._ppt_add_table(prs, slide_data)
            elif slide_type in ('chart', 'bar_chart', 'line_chart'):
                self._ppt_add_chart(prs, slide_data)
            elif slide_type in ('experiment_design', 'experiment', 'experimental_design'):
                self._ppt_add_experiment_design(prs, slide_data)
            elif slide_type in ('ablation', 'ablation_study'):
                self._ppt_add_ablation(prs, slide_data)
            else:
                self._ppt_add_bullets(prs, slide_data, Pt)
            added += 1
        return added

    def _fill_template_presentation(self, prs, slides, data, replacements):
        stats = {
            'slides_used': 0,
            'text_replaced': 0,
            'placeholders_replaced': 0,
            'template_text_replaced': 0,
            'fallback_textboxes_added': 0,
        }
        count = min(len(prs.slides), max(1, len(slides)))
        for idx in range(count):
            slide_data = slides[idx] if isinstance(slides[idx], dict) else {'title': _as_text(slides[idx])}
            slide = prs.slides[idx]
            mapping = self._template_slide_mapping(slide_data, data, idx, replacements)
            used_shapes = set()
            for shape in self._iter_ppt_shapes(slide.shapes):
                if self._replace_template_shape_text(shape, mapping, used_shapes):
                    stats['text_replaced'] += 1
                    if getattr(shape, 'is_placeholder', False):
                        stats['placeholders_replaced'] += 1
            slot_replaced = self._replace_template_text_slots(slide, prs, slide_data, data, used_shapes)
            stats['template_text_replaced'] += slot_replaced
            stats['text_replaced'] += slot_replaced
            stats['fallback_textboxes_added'] += self._fill_unmatched_template_content(slide, prs, slide_data, used_shapes)
            cleanup_replaced = self._clear_unused_template_sample_text(slide, prs, used_shapes)
            stats['template_text_replaced'] += cleanup_replaced
            stats['text_replaced'] += cleanup_replaced
            stats['slides_used'] += 1
        return stats

    def _template_slide_mapping(self, slide_data, deck_data, idx, replacements):
        title = _as_text(slide_data.get('title') or (deck_data.get('title') if idx == 0 else '') or '')
        subtitle = _as_text(slide_data.get('subtitle') or slide_data.get('desc') or (deck_data.get('subtitle') if idx == 0 else '') or '')
        body_items = slide_data.get('bullets') or slide_data.get('items') or slide_data.get('content') or slide_data.get('cards') or slide_data.get('steps') or ''
        body = self._template_body_text(body_items)
        mapping = {
            'title': title,
            'subtitle': subtitle,
            'sub_title': subtitle,
            'body': body,
            'content': body,
            'text': body,
            'date': _as_text(deck_data.get('date') or ''),
            'author': _as_text(deck_data.get('author') or ''),
            'footer': _as_text(deck_data.get('footer') or ''),
        }
        for k, v in replacements.items():
            mapping[_as_text(k).strip().lower()] = self._template_body_text(v)
        return mapping

    def _template_body_text(self, value):
        if isinstance(value, list):
            lines = []
            for item in value:
                if isinstance(item, dict):
                    title = _as_text(item.get('title') or item.get('name') or item.get('text') or '')
                    desc = _as_text(item.get('desc') or item.get('description') or item.get('content') or '')
                    lines.append((title + ('：' + desc if desc and title else desc)).strip())
                else:
                    lines.append(_as_text(item))
            return '\n'.join(x for x in lines if x)
        if isinstance(value, dict):
            return '\n'.join(f'{k}: {v}' for k, v in value.items())
        return _as_text(value)

    def _replace_template_shape_text(self, shape, mapping, used_shapes):
        if not self._ppt_has(shape, 'has_text_frame'):
            return False
        original = self._ppt_text_frame_text(shape.text_frame)
        if not original:
            return False
        key = self._placeholder_key(shape, original)
        replacement = None
        if key and key in mapping and mapping[key]:
            replacement = mapping[key]
        else:
            replaced = original
            for name, value in mapping.items():
                if value is None:
                    continue
                for token in (f'{{{{{name}}}}}', f'{{{name}}}', f'《{name}》', f'[{name}]', f'<{name}>'):
                    replaced = replaced.replace(token, _as_text(value))
            if replaced != original:
                replacement = replaced
        if replacement is None:
            return False
        try:
            tf = shape.text_frame
            self._configure_text_frame(tf, normalize_margins=False)
            tf.clear()
            lines = _as_text(replacement).splitlines() or ['']
            for i, line in enumerate(lines):
                p = tf.paragraphs[0] if i == 0 else tf.add_paragraph()
                p.text = line
                if i > 0:
                    p.level = 0
            used_shapes.add(id(shape))
            return True
        except Exception:
            return False

    def _placeholder_key(self, shape, text):
        lower = _as_text(text).strip().lower()
        m = re.fullmatch(r'[\[{<《【（(]*\s*([a-zA-Z_][\w\- ]{0,30})\s*[\]}>》】）)]*', lower)
        if m:
            token = m.group(1).strip().replace(' ', '_')
            aliases = {'heading': 'title', 'subject': 'title', 'subheading': 'subtitle', 'subtitle': 'subtitle', '正文': 'body'}
            return aliases.get(token, token)
        try:
            if getattr(shape, 'is_placeholder', False):
                ph = str(shape.placeholder_format.type).lower()
                if 'title' in ph or 'center_title' in ph:
                    return 'title'
                if 'subtitle' in ph or 'sub_title' in ph:
                    return 'subtitle'
                if any(x in ph for x in ('body', 'object', 'content', 'text')):
                    return 'body'
        except Exception:
            pass
        if lower in ('click to add title', '单击此处添加标题'):
            return 'title'
        if lower in ('click to add subtitle', '单击此处添加副标题'):
            return 'subtitle'
        if lower in ('click to add text', '单击此处添加文本'):
            return 'body'
        return ''

    def _replace_template_text_slots(self, slide, prs, slide_data, deck_data, used_shapes):
        title = _as_text(slide_data.get('title') or '')
        subtitle = _as_text(slide_data.get('subtitle') or slide_data.get('desc') or '')
        body = self._template_body_text(
            slide_data.get('bullets') or slide_data.get('items') or slide_data.get('content') or
            slide_data.get('cards') or slide_data.get('steps') or ''
        )
        if not title and slide == prs.slides[0]:
            title = _as_text(deck_data.get('title') or '')
        if not subtitle and slide == prs.slides[0]:
            subtitle = _as_text(deck_data.get('subtitle') or '')

        body_lines = [x.strip() for x in _as_text(body).splitlines() if x.strip()]
        if not (title or subtitle or body_lines):
            return 0

        existing_text = '\n'.join(self._ppt_shape_text(s) for s in self._iter_ppt_shapes(slide.shapes))
        candidates = {'title': [], 'subtitle': [], 'body': []}
        for shape in self._iter_ppt_shapes(slide.shapes):
            if id(shape) in used_shapes or not self._ppt_has(shape, 'has_text_frame'):
                continue
            text = self._ppt_text_frame_text(shape.text_frame)
            if not text:
                continue
            role, score = self._template_text_slot_role(shape, text, prs)
            if role in candidates:
                candidates[role].append((score, self._template_shape_order(shape), shape))

        replaced = 0
        title_shape = self._pop_template_slot(candidates['title'])
        if title and title not in existing_text and title_shape is not None:
            if self._write_template_shape_text(title_shape, title, used_shapes):
                replaced += 1

        subtitle_shape = self._pop_template_slot(candidates['subtitle'])
        if subtitle and subtitle not in existing_text and subtitle_shape is not None:
            if self._write_template_shape_text(subtitle_shape, subtitle, used_shapes):
                replaced += 1

        if body_lines and body not in existing_text:
            body_shapes = [self._pop_template_slot(candidates['body']) for _ in range(len(candidates['body']))]
            body_shapes = [shape for shape in body_shapes if shape is not None and id(shape) not in used_shapes]
            if not body_shapes:
                body_shape = self._pop_template_slot(candidates['title'])
                if body_shape is not None and id(body_shape) not in used_shapes:
                    body_shapes = [body_shape]
            if body_shapes:
                if len(body_shapes) == 1 or len(body_lines) == 1:
                    if self._write_template_shape_text(body_shapes[0], body, used_shapes):
                        replaced += 1
                else:
                    for idx, shape in enumerate(body_shapes):
                        if idx >= len(body_lines):
                            break
                        value = body_lines[idx] if idx < len(body_shapes) - 1 else '\n'.join(body_lines[idx:])
                        if self._write_template_shape_text(shape, value, used_shapes):
                            replaced += 1
        replaced += self._clear_unused_template_sample_text(slide, prs, used_shapes)
        return replaced

    def _template_text_slot_role(self, shape, text, prs):
        lower = re.sub(r'\s+', ' ', _as_text(text).strip().lower())
        if not lower or self._template_text_is_decorative(shape, text, prs):
            return '', 0

        key = self._placeholder_key(shape, text)
        if key in ('title', 'subtitle'):
            return key, 100
        if key in ('body', 'content', 'text'):
            return 'body', 100

        try:
            top = int(getattr(shape, 'top', 0) or 0)
            height = int(getattr(shape, 'height', 0) or 0)
            slide_h = max(1, int(prs.slide_height))
        except Exception:
            top = 0
            height = 0
            slide_h = 1

        promptish = self._template_text_is_promptish(lower)
        looks_title = self._ppt_shape_looks_like_title(shape, text, slide_h)
        if promptish:
            if looks_title or top <= slide_h * 0.24:
                return 'title', 90
            return 'body', 85

        if lower in ('contents', 'content', 'agenda', '目录', '目 录'):
            return 'title', 80
        if re.fullmatch(r'(part|chapter|section)\s*\d{1,2}', lower):
            return 'title', 70
        if re.search(r'(标题|title|heading|topic)$', lower) and len(lower) <= 40:
            return ('title', 70) if top <= slide_h * 0.4 else ('body', 55)

        if looks_title and len(lower) <= 120:
            return 'title', 45
        if self._ppt_shape_area_sq_in(shape) >= 0.5 and top >= slide_h * 0.18 and height >= slide_h * 0.04:
            return 'body', 35
        return '', 0

    def _template_text_is_promptish(self, lower):
        compact = re.sub(r'\s+', '', lower)
        prompt_tokens = (
            'clicktoaddtext', 'clicktoaddtitle', 'clicktoaddsubtitle',
            'addtext', 'addtitle', 'placeholder', 'loremipsum',
            '单击此处添加文字', '单击此处添加文本', '单击此处添加标题', '单击此处添加副标题',
            '点击此处添加文字', '点击此处添加文本', '点击此处添加标题', '请在此处输入',
            '请输入文字', '请输入文本', '输入标题', '输入文字', '输入文本',
        )
        return any(token in compact for token in prompt_tokens)

    def _template_text_is_sample_label(self, text):
        lower = re.sub(r'\s+', ' ', _as_text(text).strip().lower())
        compact = re.sub(r'\s+', '', lower)
        if self._template_text_is_promptish(lower):
            return True
        sample_labels = {
            '图标标题', '标题', '小标题', '文本标题', '正文', '正文内容',
            '示例标题', '示例文本', '样例标题', '样例文本',
            'sampletitle', 'sampletext', 'textplaceholder', 'titleplaceholder',
        }
        return compact in sample_labels

    def _template_text_is_decorative(self, shape, text, prs):
        try:
            left = int(getattr(shape, 'left', 0) or 0)
            top = int(getattr(shape, 'top', 0) or 0)
            width = int(getattr(shape, 'width', 0) or 0)
            height = int(getattr(shape, 'height', 0) or 0)
            slide_w = max(1, int(prs.slide_width))
            slide_h = max(1, int(prs.slide_height))
        except Exception:
            return False
        lower = re.sub(r'\s+', ' ', _as_text(text).strip().lower())
        if not lower:
            return True
        if self._template_text_is_promptish(lower):
            return False
        if top >= slide_h * 0.84:
            return True
        if left >= slide_w * 0.68 and top >= slide_h * 0.55 and len(lower) <= 80:
            return True
        if width <= slide_w * 0.08 or height <= slide_h * 0.025:
            return True
        if re.search(r'(copyright|confidential|footer|\b\d{4}[/-]\d{1,2}\b|\b20\d{2}\b|页脚|保密|版权所有)', lower):
            if top >= slide_h * 0.55 or len(lower) <= 80:
                return True
        return False

    def _template_shape_order(self, shape):
        try:
            return (int(getattr(shape, 'top', 0) or 0), int(getattr(shape, 'left', 0) or 0))
        except Exception:
            return (0, 0)

    def _pop_template_slot(self, candidates):
        if not candidates:
            return None
        candidates.sort(key=lambda item: (-item[0], item[1][0], item[1][1]))
        return candidates.pop(0)[2]

    def _write_template_shape_text(self, shape, text, used_shapes):
        try:
            tf = shape.text_frame
            self._configure_text_frame(tf, normalize_margins=False)
            tf.clear()
            lines = _as_text(text).splitlines() or ['']
            for i, line in enumerate(lines):
                p = tf.paragraphs[0] if i == 0 else tf.add_paragraph()
                p.text = line
                if i > 0:
                    p.level = 0
            used_shapes.add(id(shape))
            return True
        except Exception:
            return False

    def _clear_unused_template_sample_text(self, slide, prs, used_shapes):
        cleared = 0
        for shape in self._iter_ppt_shapes(slide.shapes):
            if id(shape) in used_shapes or not self._ppt_has(shape, 'has_text_frame'):
                continue
            text = self._ppt_text_frame_text(shape.text_frame)
            if not text or not self._template_text_is_sample_label(text):
                continue
            if self._template_text_is_decorative(shape, text, prs):
                continue
            if self._write_template_shape_text(shape, '', used_shapes):
                cleared += 1
        return cleared

    def _fill_unmatched_template_content(self, slide, prs, slide_data, used_shapes):
        added = 0
        existing_text = '\n'.join(self._ppt_shape_text(s) for s in self._iter_ppt_shapes(slide.shapes))
        title = _as_text(slide_data.get('title') or '')
        body = self._template_body_text(slide_data.get('bullets') or slide_data.get('items') or slide_data.get('content') or '')
        body_lines = [x.strip() for x in _as_text(body).splitlines() if x.strip()]
        from pptx.enum.text import PP_ALIGN
        if title and title not in existing_text:
            box = self._template_box('title_box') or (_in_to_emu(0.6), _in_to_emu(0.35), int(prs.slide_width) - _in_to_emu(1.2), _in_to_emu(0.6))
            self._add_textbox(slide, title, *box, size=28, bold=True, color=self._theme()['text'], align=PP_ALIGN.LEFT, style=self._style_for(slide_data, 'title'))
            added += 1
        body_present = bool(body) and (body in existing_text or (body_lines and all(line in existing_text for line in body_lines)))
        if body and not body_present:
            box = self._template_box('content_box') or (_in_to_emu(0.8), _in_to_emu(1.3), int(prs.slide_width) - _in_to_emu(1.6), int(prs.slide_height) - _in_to_emu(2.0))
            self._add_textbox(slide, body, *box, size=16, color=self._theme()['text'], align=PP_ALIGN.LEFT, style=self._style_for(slide_data, 'body'))
            added += 1
        return added

    def _ppt_request_data(self, body):
        if not isinstance(body, dict):
            return {}
        data = body.get('data')
        if isinstance(data, dict):
            return data
        return body

    def _ppt_validation_rules(self, data, include_top_level=True):
        rules = data.get('rules') if isinstance(data, dict) else {}
        rules = dict(rules) if isinstance(rules, dict) else {}
        if not include_top_level:
            return rules
        for key in (
            'min_slides', 'max_slides', 'expected_text', 'require_chinese',
            'max_question_marks', 'fail_on_warnings'
        ):
            if isinstance(data, dict) and key in data and key not in rules:
                rules[key] = data.get(key)
        return rules

    def _validate_pptx_file(self, ppt_path, rules=None):
        rules = rules or {}
        rel_path = self._ppt_rel_path(ppt_path)
        try:
            from pptx import Presentation
        except ImportError:
            return {
                'ok': False,
                'passed': False,
                'path': rel_path,
                'error': 'missing dependency: python-pptx',
            }

        if os.path.splitext(ppt_path)[1].lower() != '.pptx':
            return {
                'ok': False,
                'passed': False,
                'path': rel_path,
                'error': 'only .pptx files are supported',
            }

        try:
            prs = Presentation(ppt_path)
        except Exception as e:
            return {
                'ok': False,
                'passed': False,
                'path': rel_path,
                'error': f'failed to open pptx: {e}',
            }

        issues = []
        warnings = []
        slide_summaries = []
        all_text = []
        font_counts = {}
        font_size_counts = {}
        color_counts = {}
        total_shapes = 0
        total_pictures = 0
        total_tables = 0
        total_charts = 0
        total_text_overlaps = 0
        total_text_overflows = 0
        total_text_graphic_overlaps = 0

        slide_count = len(prs.slides)
        if slide_count <= 0:
            issues.append({'severity': 'error', 'message': 'deck has no slides'})

        min_slides = self._ppt_int(rules.get('min_slides'), 0, 0, 500)
        max_slides = self._ppt_int(rules.get('max_slides'), 0, 0, 500)
        if min_slides and slide_count < min_slides:
            issues.append({'severity': 'error', 'message': f'slide count {slide_count} is below min_slides {min_slides}'})
        if max_slides and slide_count > max_slides:
            issues.append({'severity': 'error', 'message': f'slide count {slide_count} is above max_slides {max_slides}'})

        com_text_bounds_by_slide = {}
        com_warning = None
        use_com_text_bounds = self._ppt_bool(rules.get('use_powerpoint_com_text_bounds'), True)
        if use_com_text_bounds:
            com_text_bounds_by_slide, com_warning = self._ppt_powerpoint_com_text_bounds(ppt_path)
            if com_warning:
                warnings.append({
                    'severity': 'warning',
                    'message': f'PowerPoint COM text bounds unavailable: {com_warning}; falling back to estimated text bounds',
                })

        com_layout_issues_by_slide = {}
        if use_com_text_bounds and not com_warning:
            com_layout_issues_by_slide, com_layout_warning = self._ppt_powerpoint_com_layout_issues(ppt_path)
            if com_layout_warning:
                warnings.append({
                    'severity': 'warning',
                    'message': f'PowerPoint COM layout issue scan unavailable: {com_layout_warning}',
                })

        for idx, slide in enumerate(prs.slides, start=1):
            summary = self._ppt_slide_summary(slide, idx, prs, font_counts, font_size_counts, color_counts, com_text_bounds_by_slide.get(idx))
            slide_summaries.append(summary)
            total_shapes += summary['shape_count']
            total_pictures += summary['picture_count']
            total_tables += summary['table_count']
            total_charts += summary['chart_count']
            total_text_overlaps += len(summary.get('text_overlaps', []))
            total_text_overflows += len(summary.get('text_overflow', []))
            if summary['text']:
                all_text.append(summary['text'])
            if not summary['title'] and idx > 1:
                warnings.append({'severity': 'warning', 'slide': idx, 'message': 'slide has no clear title'})
            if summary['text_count'] == 0 and summary['picture_count'] == 0 and summary['table_count'] == 0 and summary['chart_count'] == 0:
                warnings.append({'severity': 'warning', 'slide': idx, 'message': 'slide appears empty'})
            if summary['char_count'] > 900:
                warnings.append({'severity': 'warning', 'slide': idx, 'message': f'slide is text-heavy ({summary["char_count"]} chars)'})
            for item in summary.get('overcrowded', [])[:3]:
                warnings.append({
                    'severity': 'warning',
                    'slide': idx,
                    'message': f'possible text overflow: {item["chars"]} chars in {item["area_sq_in"]:.2f} sq.in',
                    'text': item.get('text', ''),
                })
            for item in summary.get('off_slide', [])[:5]:
                warnings.append({
                    'severity': 'warning',
                    'slide': idx,
                    'message': f'shape may be outside slide bounds: {item["bounds"]}',
                })
            for item in summary.get('text_overflow', [])[:5]:
                issues.append({
                    'severity': 'error',
                    'slide': idx,
                    'message': f'possible text box overflow: estimated {item["estimated_lines"]} lines > capacity {item["capacity_lines"]}',
                    'text': item.get('text', ''),
                    'bounds': item.get('bounds', {}),
                })
            for item in summary.get('text_overlaps', [])[:5]:
                issues.append({
                    'severity': 'error',
                    'slide': idx,
                    'message': f'overlapping text boxes: {item["overlap_pct"]}% of smaller text area overlaps',
                    'text_a': item.get('text_a', ''),
                    'text_b': item.get('text_b', ''),
                    'bounds_a': item.get('bounds_a', {}),
                    'bounds_b': item.get('bounds_b', {}),
                    'text_bounds_a': item.get('text_bounds_a', {}),
                    'text_bounds_b': item.get('text_bounds_b', {}),
                })
            for item in com_layout_issues_by_slide.get(idx, [])[:8]:
                total_text_graphic_overlaps += 1
                issues.append({
                    'severity': 'error',
                    'slide': idx,
                    'message': item.get('message', 'text/graphic layout conflict'),
                    'kind': item.get('kind', 'text_graphic_overlap'),
                    'text': item.get('text', ''),
                    'text_bounds': item.get('text_bounds', {}),
                    'graphic_bounds': item.get('graphic_bounds', {}),
                    'overlap_pct': item.get('overlap_pct'),
                })

        visible_text = '\n'.join(all_text)
        question_marks = visible_text.count('?')
        cjk_chars = len(re.findall(r'[\u4e00-\u9fff]', visible_text))
        mojibake_hits = len(re.findall(r'[锛鑰鈥馃�]', visible_text))
        if re.search(r'\?{3,}', visible_text):
            issues.append({'severity': 'error', 'message': 'visible text contains repeated question marks; possible font or encoding loss'})
        max_question_marks = self._ppt_int(rules.get('max_question_marks'), -1, -1, 100000)
        if max_question_marks >= 0 and question_marks > max_question_marks:
            issues.append({'severity': 'error', 'message': f'question mark count {question_marks} exceeds max_question_marks {max_question_marks}'})
        if mojibake_hits >= 3:
            issues.append({'severity': 'error', 'message': 'visible text contains likely mojibake characters'})
        if self._ppt_bool(rules.get('require_chinese'), False) and cjk_chars <= 0:
            issues.append({'severity': 'error', 'message': 'require_chinese is true but no CJK text was found'})

        expected_text = _as_list(rules.get('expected_text'))
        for expected in expected_text:
            expected = _item_text(expected).strip()
            if expected and expected not in visible_text:
                issues.append({'severity': 'error', 'message': f'expected text not found: {expected[:80]}'})

        fail_on_warnings = self._ppt_bool(rules.get('fail_on_warnings'), False)
        passed = len(issues) == 0 and (not fail_on_warnings or len(warnings) == 0)
        score = max(0, 100 - len(issues) * 25 - len(warnings) * 5)
        return {
            'ok': True,
            'passed': passed,
            'score': score,
            'path': rel_path,
            'slide_count': slide_count,
            'stats': {
                'shape_count': total_shapes,
                'text_char_count': len(visible_text),
                'text_block_count': sum(x['text_count'] for x in slide_summaries),
                'picture_count': total_pictures,
                'table_count': total_tables,
                'chart_count': total_charts,
                'text_overlap_count': total_text_overlaps,
                'text_overflow_count': total_text_overflows,
                'text_graphic_overlap_count': total_text_graphic_overlaps,
                'visible_question_marks': question_marks,
                'cjk_char_count': cjk_chars,
                'mojibake_hit_count': mojibake_hits,
                'fonts': self._ppt_top_counts(font_counts),
                'font_sizes': self._ppt_top_counts(font_size_counts),
                'font_colors': self._ppt_top_counts(color_counts),
            },
            'issues': issues,
            'warnings': warnings,
            'slides': slide_summaries,
        }

    def _ppt_slide_summary(self, slide, idx, prs, font_counts=None, font_size_counts=None, color_counts=None, com_text_bounds=None):
        font_counts = font_counts if isinstance(font_counts, dict) else {}
        font_size_counts = font_size_counts if isinstance(font_size_counts, dict) else {}
        color_counts = color_counts if isinstance(color_counts, dict) else {}
        slide_w = max(1, int(getattr(prs, 'slide_width', 1) or 1))
        slide_h = max(1, int(getattr(prs, 'slide_height', 1) or 1))
        texts = []
        text_blocks = []
        title = ''
        shape_count = 0
        picture_count = 0
        table_count = 0
        chart_count = 0
        overcrowded = []
        bullet_like = 0
        off_slide = []
        text_overflow = []
        text_overlap_candidates = []

        for shape in self._iter_ppt_shapes(slide.shapes):
            shape_count += 1
            if self._ppt_shape_has_image(shape):
                picture_count += 1
            if self._ppt_has(shape, 'has_table'):
                table_count += 1
            if self._ppt_has(shape, 'has_chart'):
                chart_count += 1
            bounds = self._ppt_shape_bounds(shape)
            if bounds and self._ppt_bounds_outside_slide(bounds, slide_w, slide_h):
                off_slide.append({
                    'bounds': self._ppt_format_bounds(bounds),
                    'slide_size': self._ppt_format_bounds((0, 0, slide_w, slide_h)),
                    'text': self._ppt_shape_text(shape).strip().replace('\n', ' ')[:120],
                })

            text = self._ppt_shape_text(shape).strip()
            if not text:
                continue
            texts.append(text)
            preview = text.replace('\n', ' ')[:160]
            text_blocks.append(preview)
            bullet_like += len(re.findall(r'(^|\n)\s*([•\-\*\u2022]|\d+[.)])\s+', text))
            if not title and self._ppt_shape_looks_like_title(shape, text, slide_h):
                title = text.splitlines()[0][:120]
            overlap_candidate = self._ppt_text_overlap_candidate(shape, text, bounds, slide_w, slide_h, com_text_bounds)
            if overlap_candidate:
                text_overlap_candidates.append(overlap_candidate)

            self._ppt_collect_text_styles(shape, font_counts, font_size_counts, color_counts)
            area = self._ppt_shape_area_sq_in(shape)
            if area > 0:
                cjk_count = len(re.findall(r'[\u4e00-\u9fff]', text))
                density_limit = 90 if cjk_count > max(0, len(text) * 0.3) else 180
                if len(text) > 120 and (len(text) / area) > density_limit:
                    overcrowded.append({
                        'chars': len(text),
                        'area_sq_in': area,
                        'text': preview,
                    })
            overflow = self._ppt_estimate_text_overflow(shape, text)
            if overflow:
                text_overflow.append({
                    'estimated_lines': overflow['estimated_lines'],
                    'capacity_lines': overflow['capacity_lines'],
                    'font_size': overflow['font_size'],
                    'bounds': self._ppt_format_bounds(bounds) if bounds else {},
                    'text': preview,
                })

        if not title and texts:
            title = texts[0].splitlines()[0][:120]

        full_text = '\n'.join(texts)
        text_overlaps = self._ppt_find_text_overlaps(text_overlap_candidates)
        return {
            'slide': idx,
            'title': title,
            'shape_count': shape_count,
            'text_count': len(text_blocks),
            'char_count': len(full_text),
            'picture_count': picture_count,
            'table_count': table_count,
            'chart_count': chart_count,
            'bullet_like_count': bullet_like,
            'text_preview': text_blocks[:8],
            'overcrowded': overcrowded,
            'off_slide': off_slide,
            'text_overflow': text_overflow,
            'text_overlaps': text_overlaps,
            'text': full_text,
        }

    def _ppt_shape_bounds(self, shape):
        try:
            left = int(getattr(shape, 'left', 0) or 0)
            top = int(getattr(shape, 'top', 0) or 0)
            width = int(getattr(shape, 'width', 0) or 0)
            height = int(getattr(shape, 'height', 0) or 0)
            return (left, top, left + width, top + height)
        except Exception:
            return None

    def _ppt_bounds_outside_slide(self, bounds, slide_w, slide_h, tolerance=91440):
        left, top, right, bottom = bounds
        return left < -tolerance or top < -tolerance or right > slide_w + tolerance or bottom > slide_h + tolerance

    def _ppt_auto_repair_text_layout(self, ppt_path, rules=None):
        """Use PowerPoint's renderer to repair overflowing/overlapping text.

        python-pptx can set text frame metadata, but it cannot perform the final
        layout pass. On Windows with PowerPoint installed, use COM after saving
        the deck so generated files are not merely diagnosed as invalid. Prefer
        spatial repairs first: enlarge overflowing text boxes when there is room
        and move lower colliding text boxes downward to open whitespace. Only
        when the slide has insufficient room for a safe spatial repair do we fall
        back to PowerPoint's text-fit / font-size reduction behavior.
        """
        if os.name != 'nt':
            return {'enabled': False, 'repaired': False, 'reason': 'not running on Windows'}
        rules = rules or {}
        max_passes = self._ppt_int(rules.get('auto_repair_passes'), 3, 1, 8)
        min_font_size = self._ppt_int(rules.get('auto_repair_min_font_size'), 8, 5, 18)
        try:
            import pythoncom
            import win32com.client
        except Exception as e:
            return {'enabled': False, 'repaired': False, 'reason': f'pywin32 is not available ({e})'}

        app = None
        pres = None
        repaired = False
        repaired_shapes = 0
        layout_repairs = 0
        shrink_repairs = 0
        last_problem_count = 0
        pythoncom.CoInitialize()
        try:
            app = win32com.client.DispatchEx('PowerPoint.Application')
            try:
                app.DisplayAlerts = 0
            except Exception:
                pass
            pres = app.Presentations.Open(os.path.abspath(ppt_path), ReadOnly=False, Untitled=False, WithWindow=False)

            for pass_idx in range(max_passes):
                problem_shapes = {}
                for slide in pres.Slides:
                    items = []
                    graphics = []
                    self._ppt_collect_com_text_shape_items(slide.Shapes, items)
                    self._ppt_collect_com_graphic_shape_items(slide.Shapes, graphics)
                    problem_shapes.update(self._ppt_com_problem_shape_map_with_graphics(items, graphics))

                last_problem_count = len(problem_shapes)
                if not problem_shapes:
                    break

                changed_this_pass = 0
                for slide in pres.Slides:
                    items = []
                    graphics = []
                    self._ppt_collect_com_text_shape_items(slide.Shapes, items)
                    self._ppt_collect_com_graphic_shape_items(slide.Shapes, graphics)
                    changed_this_pass += self._ppt_com_reflow_text_slide(slide, items, graphics=graphics, pass_idx=pass_idx)
                if changed_this_pass > 0:
                    layout_repairs += changed_this_pass
                else:
                    for shape in list(problem_shapes.values()):
                        if self._ppt_com_shrink_text_shape(shape, min_font_size=min_font_size, pass_idx=pass_idx):
                            changed_this_pass += 1
                    shrink_repairs += changed_this_pass
                if changed_this_pass <= 0:
                    break
                repaired = True
                repaired_shapes += changed_this_pass
                try:
                    pres.Save()
                except Exception:
                    pass

            if repaired:
                pres.Save()
            remaining_problem_count = 0
            for slide in pres.Slides:
                items = []
                graphics = []
                self._ppt_collect_com_text_shape_items(slide.Shapes, items)
                self._ppt_collect_com_graphic_shape_items(slide.Shapes, graphics)
                remaining_problem_count += len(self._ppt_com_problem_shape_map_with_graphics(items, graphics))
            return {
                'enabled': True,
                'repaired': bool(repaired),
                'repaired_shapes': int(repaired_shapes),
                'remaining_problem_shapes': int(remaining_problem_count),
                'layout_repairs': int(layout_repairs),
                'shrink_repairs': int(shrink_repairs),
                'method': 'powerpoint_com_layout_first_then_autofit',
            }
        except Exception as e:
            return {'enabled': False, 'repaired': False, 'reason': str(e)}
        finally:
            try:
                if pres is not None:
                    pres.Close()
            except Exception:
                pass
            try:
                if app is not None:
                    app.Quit()
            except Exception:
                pass
            try:
                pythoncom.CoUninitialize()
            except Exception:
                pass

    def _ppt_collect_com_text_shape_items(self, shapes, out):
        try:
            count = int(shapes.Count)
        except Exception:
            return
        for i in range(1, count + 1):
            try:
                shape = shapes.Item(i)
            except Exception:
                continue
            try:
                if int(getattr(shape, 'Type', 0)) == 6:  # msoGroup
                    self._ppt_collect_com_text_shape_items(shape.GroupItems, out)
                    continue
            except Exception:
                pass
            try:
                if not bool(shape.HasTextFrame) or not bool(shape.TextFrame.HasText):
                    continue
                text = _as_text(shape.TextFrame2.TextRange.Text).strip()
                if not text:
                    continue
                tr = shape.TextFrame2.TextRange
                text_bounds = self._ppt_com_rect_to_emu(tr.BoundLeft, tr.BoundTop, tr.BoundWidth, tr.BoundHeight)
                shape_bounds = self._ppt_com_rect_to_emu(shape.Left, shape.Top, shape.Width, shape.Height)
                if not text_bounds or not shape_bounds or self._ppt_bounds_area(text_bounds) <= 0:
                    continue
                try:
                    shape_key = f'{int(shape.Parent.SlideIndex)}:{int(shape.Id)}'
                except Exception:
                    shape_key = str(id(shape))
                out.append({
                    'shape': shape, 'shape_key': shape_key, 'text': text, 'text_bounds': text_bounds, 'shape_bounds': shape_bounds
                })
            except Exception:
                continue

    def _ppt_collect_com_graphic_shape_items(self, shapes, out):
        """Collect non-text visual shapes that may collide with text."""
        try:
            count = int(shapes.Count)
        except Exception:
            return
        for i in range(1, count + 1):
            try:
                shape = shapes.Item(i)
            except Exception:
                continue
            try:
                if int(getattr(shape, 'Type', 0)) == 6:  # msoGroup
                    self._ppt_collect_com_graphic_shape_items(shape.GroupItems, out)
                    continue
            except Exception:
                pass
            try:
                has_text = bool(shape.HasTextFrame) and bool(shape.TextFrame.HasText)
            except Exception:
                has_text = False
            if has_text:
                continue
            try:
                bounds = self._ppt_com_rect_to_emu(shape.Left, shape.Top, shape.Width, shape.Height)
                if not bounds or self._ppt_bounds_area(bounds) <= 0:
                    continue
                try:
                    slide_w_pt = float(shape.Parent.Parent.PageSetup.SlideWidth)
                    slide_h_pt = float(shape.Parent.Parent.PageSetup.SlideHeight)
                    slide_area = max(1.0, slide_w_pt * slide_h_pt)
                    if (float(shape.Width) * float(shape.Height)) / slide_area > 0.70:
                        continue
                except Exception:
                    pass
                try:
                    shape_key = f'{int(shape.Parent.SlideIndex)}:{int(shape.Id)}'
                except Exception:
                    shape_key = str(id(shape))
                out.append({
                    'shape': shape,
                    'shape_key': shape_key,
                    'bounds': bounds,
                    'z_order': self._ppt_com_z_order(shape),
                    'type': int(getattr(shape, 'Type', 0) or 0),
                })
            except Exception:
                continue

    def _ppt_com_z_order(self, shape):
        try:
            return int(shape.ZOrderPosition)
        except Exception:
            return 0

    def _ppt_com_problem_shape_map(self, items):
        problem_shapes = {}
        for item in items:
            if self._ppt_com_text_item_overflows(item):
                problem_shapes[item.get('shape_key') or id(item.get('shape'))] = item.get('shape')
        for a_idx, first in enumerate(items):
            for second in items[a_idx + 1:]:
                if self._ppt_com_text_items_overlap(first, second):
                    problem_shapes[first.get('shape_key') or id(first.get('shape'))] = first.get('shape')
                    problem_shapes[second.get('shape_key') or id(second.get('shape'))] = second.get('shape')
        return {k: v for k, v in problem_shapes.items() if v is not None}

    def _ppt_com_problem_shape_map_with_graphics(self, text_items, graphic_items):
        problem_shapes = self._ppt_com_problem_shape_map(text_items)
        for issue in self._ppt_com_text_graphic_issues(text_items, graphic_items):
            shape = issue.get('text_shape')
            if shape is not None:
                problem_shapes[issue.get('text_key') or id(shape)] = shape
        return {k: v for k, v in problem_shapes.items() if v is not None}

    def _ppt_com_text_item_overflows(self, item):
        return self._ppt_bounds_extends_outside(item.get('text_bounds'), item.get('shape_bounds'), tolerance=int(_EMU_PER_INCH * 0.02))

    def _ppt_com_text_items_overlap(self, first, second):
        inter = self._ppt_bounds_intersection(first.get('text_bounds'), second.get('text_bounds'))
        if not inter:
            return False
        inter_area = self._ppt_bounds_area(inter)
        smaller = max(1, min(self._ppt_bounds_area(first.get('text_bounds')), self._ppt_bounds_area(second.get('text_bounds'))))
        overlap_ratio = inter_area / smaller
        left, top, right, bottom = inter
        overlap_w_in = (right - left) / _EMU_PER_INCH
        overlap_h_in = (bottom - top) / _EMU_PER_INCH
        overlap_area_in = inter_area / (_EMU_PER_INCH * _EMU_PER_INCH)
        return overlap_ratio >= 0.08 and overlap_area_in >= 0.01 and overlap_w_in >= 0.06 and overlap_h_in >= 0.03

    def _ppt_com_text_graphic_issues(self, text_items, graphic_items):
        issues = []
        for text_item in text_items or []:
            text_bounds = text_item.get('text_bounds')
            shape_bounds = text_item.get('shape_bounds')
            if not text_bounds or not shape_bounds:
                continue
            if self._ppt_com_text_item_overflows(text_item):
                issues.append({
                    'kind': 'text_outside_container',
                    'message': 'text extends outside its containing shape',
                    'text': text_item.get('text', ''),
                    'text_shape': text_item.get('shape'),
                    'text_key': text_item.get('shape_key'),
                    'text_bounds': text_bounds,
                    'shape_bounds': shape_bounds,
                    'graphic_bounds': shape_bounds,
                    'overlap_pct': 100,
                })
            for graphic in graphic_items or []:
                graphic_bounds = graphic.get('bounds')
                inter = self._ppt_bounds_intersection(text_bounds, graphic_bounds)
                if not inter:
                    continue
                inter_area = self._ppt_bounds_area(inter)
                text_area = max(1, self._ppt_bounds_area(text_bounds))
                graphic_area = max(1, self._ppt_bounds_area(graphic_bounds))
                text_ratio = inter_area / text_area
                graphic_ratio = inter_area / graphic_area
                left, top, right, bottom = inter
                overlap_w_in = (right - left) / _EMU_PER_INCH
                overlap_h_in = (bottom - top) / _EMU_PER_INCH
                overlap_area_in = inter_area / (_EMU_PER_INCH * _EMU_PER_INCH)
                if overlap_area_in < 0.015 or overlap_w_in < 0.05 or overlap_h_in < 0.04:
                    continue
                if text_ratio < 0.08 and graphic_ratio < 0.08:
                    continue
                text_above = self._ppt_com_z_order(text_item.get('shape')) >= int(graphic.get('z_order') or 0)
                kind = 'text_covers_graphic' if text_above else 'graphic_covers_text'
                issues.append({
                    'kind': kind,
                    'message': 'text overlaps a non-text graphic shape' if text_above else 'non-text graphic shape overlaps text',
                    'text': text_item.get('text', '')[:160],
                    'text_shape': text_item.get('shape'),
                    'text_key': text_item.get('shape_key'),
                    'text_bounds': text_bounds,
                    'shape_bounds': shape_bounds,
                    'graphic_bounds': graphic_bounds,
                    'graphic_shape': graphic.get('shape'),
                    'overlap_pct': int(round(max(text_ratio, graphic_ratio) * 100)),
                })
        issues.sort(key=lambda item: int(item.get('overlap_pct') or 0), reverse=True)
        return issues[:20]

    def _ppt_com_reflow_text_slide(self, slide, items, graphics=None, pass_idx=0):
        """Repair a slide by changing geometry before changing font size."""
        if not items:
            return 0
        graphics = graphics or []
        try:
            slide_h_pt = float(slide.Parent.PageSetup.SlideHeight)
        except Exception:
            slide_h_pt = 540.0
        slide_h_emu = int(round(slide_h_pt / 72.0 * _EMU_PER_INCH))
        bottom_margin_emu = int(_EMU_PER_INCH * 0.25)
        gap_emu = int(_EMU_PER_INCH * (0.08 + min(max(pass_idx, 0), 3) * 0.02))
        changed = 0

        for item in sorted(items, key=lambda it: (it.get('shape_bounds') or (0, 0, 0, 0))[1]):
            if not self._ppt_com_text_item_overflows(item):
                continue
            text_bounds = item.get('text_bounds')
            shape_bounds = item.get('shape_bounds')
            shape = item.get('shape')
            if not text_bounds or not shape_bounds or shape is None:
                continue
            needed_emu = max(0, text_bounds[3] - shape_bounds[3] + gap_emu)
            if needed_emu <= 0:
                continue
            max_growth = max(0, slide_h_emu - bottom_margin_emu - shape_bounds[3])
            grow_emu = min(needed_emu, max_growth)
            if grow_emu < int(_EMU_PER_INCH * 0.03):
                continue
            try:
                shape.Height = max(float(shape.Height), float(shape.Height) + (grow_emu / _EMU_PER_INCH * 72.0))
                changed += 1
            except Exception:
                pass

        if changed:
            return changed

        moves = {}
        for a_idx, first in enumerate(items):
            for second in items[a_idx + 1:]:
                if not self._ppt_com_text_items_overlap(first, second):
                    continue
                first_tb = first.get('text_bounds')
                second_tb = second.get('text_bounds')
                if not first_tb or not second_tb:
                    continue
                upper, lower = (first, second) if first_tb[1] <= second_tb[1] else (second, first)
                upper_tb = upper.get('text_bounds')
                lower_tb = lower.get('text_bounds')
                lower_shape_bounds = lower.get('shape_bounds')
                if not upper_tb or not lower_tb or not lower_shape_bounds:
                    continue
                delta_emu = max(0, upper_tb[3] - lower_tb[1] + gap_emu)
                if delta_emu < int(_EMU_PER_INCH * 0.03):
                    continue
                available_emu = slide_h_emu - bottom_margin_emu - lower_shape_bounds[3]
                if available_emu < delta_emu:
                    continue
                key = lower.get('shape_key') or id(lower.get('shape'))
                previous = moves.get(key)
                if previous is None or delta_emu > previous[0]:
                    moves[key] = (delta_emu, lower.get('shape'))

        for delta_emu, shape in moves.values():
            if shape is None:
                continue
            try:
                shape.Top = float(shape.Top) + (delta_emu / _EMU_PER_INCH * 72.0)
                changed += 1
            except Exception:
                pass
        if changed:
            return changed

        for issue in self._ppt_com_text_graphic_issues(items, graphics):
            text_shape = issue.get('text_shape')
            text_bounds = issue.get('text_bounds')
            graphic_bounds = issue.get('graphic_bounds')
            shape_bounds = issue.get('shape_bounds')
            if text_shape is None or not text_bounds or not shape_bounds or not graphic_bounds:
                continue
            if issue.get('kind') == 'text_outside_container':
                # Horizontal text overflow is usually caused by an over-wide line
                # or missing wrap. Growing the box to fit the rendered text can
                # make it cover neighboring icons/diagrams, so leave this case to
                # the shrink/autofit fallback. Vertical growth is already handled
                # above by _ppt_com_text_item_overflows.
                needed_h = max(0, text_bounds[3] - shape_bounds[3] + gap_emu)
                if needed_h <= int(_EMU_PER_INCH * 0.03):
                    continue
                continue
            try:
                if text_bounds[0] < graphic_bounds[0] and shape_bounds[2] > graphic_bounds[0]:
                    new_w_emu = max(int(_EMU_PER_INCH * 0.35), graphic_bounds[0] - shape_bounds[0] - gap_emu)
                    if new_w_emu < shape_bounds[2] - shape_bounds[0] - int(_EMU_PER_INCH * 0.03):
                        text_shape.Width = new_w_emu / _EMU_PER_INCH * 72.0
                        changed += 1
                elif text_bounds[2] > graphic_bounds[2] and shape_bounds[0] < graphic_bounds[2]:
                    shift_emu = graphic_bounds[2] - shape_bounds[0] + gap_emu
                    available = shape_bounds[2] - shape_bounds[0] - shift_emu
                    if available > int(_EMU_PER_INCH * 0.35):
                        text_shape.Left = float(text_shape.Left) + (shift_emu / _EMU_PER_INCH * 72.0)
                        text_shape.Width = available / _EMU_PER_INCH * 72.0
                        changed += 1
            except Exception:
                pass
        return changed

    def _ppt_com_shrink_text_shape(self, shape, min_font_size=8, pass_idx=0):
        changed = False
        try:
            shape.TextFrame.WordWrap = True
        except Exception:
            pass
        try:
            shape.TextFrame2.WordWrap = -1  # msoTrue
        except Exception:
            pass
        try:
            shape.TextFrame2.AutoSize = 2  # msoAutoSizeTextToFitShape
            changed = True
        except Exception:
            pass
        try:
            tf2 = shape.TextFrame2
            tf2.MarginLeft = min(float(getattr(tf2, 'MarginLeft', 0) or 0), 2.0)
            tf2.MarginRight = min(float(getattr(tf2, 'MarginRight', 0) or 0), 2.0)
            tf2.MarginTop = min(float(getattr(tf2, 'MarginTop', 0) or 0), 1.0)
            tf2.MarginBottom = min(float(getattr(tf2, 'MarginBottom', 0) or 0), 1.0)
            changed = True
        except Exception:
            pass
        try:
            font = shape.TextFrame2.TextRange.Font
            current = float(font.Size)
            factor = 0.88 if pass_idx == 0 else 0.82
            target = max(float(min_font_size), current * factor)
            if target < current - 0.2:
                font.Size = target
                changed = True
        except Exception:
            pass
        return changed

    def _ppt_powerpoint_com_layout_issues(self, ppt_path):
        """Return rendered text-vs-graphic conflicts by slide using PowerPoint COM."""
        if os.name != 'nt':
            return {}, 'not running on Windows'
        try:
            import pythoncom
            import win32com.client
        except Exception as e:
            return {}, f'pywin32 is not available ({e})'

        app = None
        pres = None
        pythoncom.CoInitialize()
        try:
            app = win32com.client.DispatchEx('PowerPoint.Application')
            try:
                app.DisplayAlerts = 0
            except Exception:
                pass
            pres = app.Presentations.Open(os.path.abspath(ppt_path), ReadOnly=True, Untitled=False, WithWindow=False)
            by_slide = {}
            for slide in pres.Slides:
                text_items = []
                graphic_items = []
                self._ppt_collect_com_text_shape_items(slide.Shapes, text_items)
                self._ppt_collect_com_graphic_shape_items(slide.Shapes, graphic_items)
                slide_issues = []
                for issue in self._ppt_com_text_graphic_issues(text_items, graphic_items):
                    slide_issues.append({
                        'kind': issue.get('kind'),
                        'message': issue.get('message'),
                        'text': issue.get('text', ''),
                        'text_bounds': self._ppt_format_bounds(issue.get('text_bounds')),
                        'shape_bounds': self._ppt_format_bounds(issue.get('shape_bounds')),
                        'graphic_bounds': self._ppt_format_bounds(issue.get('graphic_bounds')),
                        'overlap_pct': issue.get('overlap_pct'),
                    })
                if slide_issues:
                    by_slide[int(slide.SlideIndex)] = slide_issues
            return by_slide, None
        except Exception as e:
            return {}, str(e)
        finally:
            try:
                if pres is not None:
                    pres.Close()
            except Exception:
                pass
            try:
                if app is not None:
                    app.Quit()
            except Exception:
                pass
            try:
                pythoncom.CoUninitialize()
            except Exception:
                pass

    def _ppt_powerpoint_com_text_bounds(self, ppt_path):
        """Return rendered PowerPoint text bounds by slide, in EMU units."""
        if os.name != 'nt':
            return {}, 'not running on Windows'
        try:
            import pythoncom
            import win32com.client
        except Exception as e:
            return {}, f'pywin32 is not available ({e})'

        app = None
        pres = None
        pythoncom.CoInitialize()
        try:
            app = win32com.client.DispatchEx('PowerPoint.Application')
            try:
                app.DisplayAlerts = 0
            except Exception:
                pass
            pres = app.Presentations.Open(os.path.abspath(ppt_path), ReadOnly=True, Untitled=False, WithWindow=False)
            by_slide = {}
            for slide in pres.Slides:
                slide_items = []
                self._ppt_collect_com_text_bounds(slide.Shapes, slide_items)
                if slide_items:
                    by_slide[int(slide.SlideIndex)] = slide_items
            return by_slide, None
        except Exception as e:
            return {}, str(e)
        finally:
            try:
                if pres is not None:
                    pres.Close()
            except Exception:
                pass
            try:
                if app is not None:
                    app.Quit()
            except Exception:
                pass
            try:
                pythoncom.CoUninitialize()
            except Exception:
                pass

    def _ppt_collect_com_text_bounds(self, shapes, out):
        try:
            count = int(shapes.Count)
        except Exception:
            return
        for i in range(1, count + 1):
            try:
                shape = shapes.Item(i)
            except Exception:
                continue
            try:
                if int(getattr(shape, 'Type', 0)) == 6:  # msoGroup
                    self._ppt_collect_com_text_bounds(shape.GroupItems, out)
                    continue
            except Exception:
                pass
            try:
                if not bool(shape.HasTextFrame) or not bool(shape.TextFrame.HasText):
                    continue
            except Exception:
                continue
            try:
                text = _as_text(shape.TextFrame2.TextRange.Text).strip()
                if not text:
                    continue
                rendered = shape.TextFrame2.TextRange
                text_bounds = self._ppt_com_rect_to_emu(rendered.BoundLeft, rendered.BoundTop, rendered.BoundWidth, rendered.BoundHeight)
                shape_bounds = self._ppt_com_rect_to_emu(shape.Left, shape.Top, shape.Width, shape.Height)
                if not text_bounds or self._ppt_bounds_area(text_bounds) <= 0:
                    continue
                out.append({
                    'text': text,
                    'text_norm': self._ppt_normalize_text_for_match(text),
                    'text_bounds': text_bounds,
                    'shape_bounds': shape_bounds,
                    'matched': False,
                })
            except Exception:
                continue

    def _ppt_com_rect_to_emu(self, left, top, width, height):
        try:
            l = int(round(float(left) / 72.0 * _EMU_PER_INCH))
            t = int(round(float(top) / 72.0 * _EMU_PER_INCH))
            r = int(round((float(left) + float(width)) / 72.0 * _EMU_PER_INCH))
            b = int(round((float(top) + float(height)) / 72.0 * _EMU_PER_INCH))
            return (l, t, r, b)
        except Exception:
            return None

    def _ppt_normalize_text_for_match(self, text):
        return re.sub(r'\s+', '', _as_text(text)).lower()

    def _ppt_text_overlap_candidate(self, shape, text, bounds, slide_w, slide_h, com_text_bounds=None):
        if not bounds or not self._ppt_has(shape, 'has_text_frame'):
            return None
        stripped = _as_text(text).strip()
        if not stripped:
            return None
        if len(stripped) <= 2 and not re.search(r'[A-Za-z\u4e00-\u9fff]', stripped):
            return None
        left, top, right, bottom = bounds
        width = max(0, right - left)
        height = max(0, bottom - top)
        if width < slide_w * 0.035 or height < slide_h * 0.025:
            return None
        lower = re.sub(r'\s+', ' ', stripped.lower())
        if top >= slide_h * 0.86 and len(lower) <= 80:
            return None
        if re.search(r'(copyright|footer|\b20\d{2}\b|页脚|版权所有)', lower) and len(lower) <= 100:
            return None

        source = 'estimated'
        com_item = self._ppt_match_com_text_bounds(stripped, bounds, com_text_bounds)
        occupied = self._ppt_text_occupied_bounds(shape, stripped, bounds)
        if com_item:
            occupied = com_item.get('text_bounds') or occupied
            source = 'powerpoint_com'
            if com_item.get('shape_bounds'):
                bounds = com_item['shape_bounds']
        if not occupied:
            return None
        if self._ppt_bounds_area(occupied) <= 0:
            return None
        # Keep both the text's estimated visual footprint and the original
        # textbox bounds.  Some decks use very small textboxes with autofit off;
        # the rendered glyphs then extend outside the shape even though the
        # shape rectangles themselves do not overlap.
        return {
            'bounds': occupied,
            'shape_bounds': bounds,
            'text': stripped.replace('\n', ' ')[:120],
            'bounds_source': source,
        }

    def _ppt_match_com_text_bounds(self, text, shape_bounds, com_text_bounds):
        if not com_text_bounds:
            return None
        norm = self._ppt_normalize_text_for_match(text)
        best = None
        best_score = -1.0
        for item in com_text_bounds:
            if item.get('matched'):
                continue
            item_norm = item.get('text_norm') or ''
            if not item_norm:
                continue
            if not (norm == item_norm or norm in item_norm or item_norm in norm):
                continue
            inter = self._ppt_bounds_intersection(shape_bounds, item.get('shape_bounds'))
            shape_score = 0.0
            if inter:
                shape_score = self._ppt_bounds_area(inter) / max(1, min(self._ppt_bounds_area(shape_bounds), self._ppt_bounds_area(item.get('shape_bounds'))))
            if shape_score < 0.20:
                continue
            score = (1.0 if norm == item_norm else 0.7) + shape_score
            if score > best_score:
                best = item
                best_score = score
        if best:
            best['matched'] = True
        return best

    def _ppt_text_occupied_bounds(self, shape, text, bounds):
        left, top, right, bottom = bounds
        inner_left, inner_top, inner_right, inner_bottom = self._ppt_inner_text_bounds(shape, bounds)
        inner_width = max(1, inner_right - inner_left)
        inner_height = max(1, inner_bottom - inner_top)

        layout = self._ppt_estimate_text_layout(shape, text, bounds)
        font_size = layout['font_size'] if layout else (self._ppt_shape_font_size(shape) or 14)
        estimated_lines = layout['estimated_lines'] if layout else 1
        line_height_in = (layout.get('line_height_in') if layout else None) or max(0.12, font_size * 1.34 / 72.0)
        line_height = int(line_height_in * _EMU_PER_INCH)
        # Add a small leading/descent safety margin.  Rendering engines reserve
        # glyph ascent/descent differently; without this, a visibly touching or
        # overlapping pair may be underestimated as separated by a few pixels.
        text_height = max(line_height, int((estimated_lines * line_height) + (font_size * 0.10 / 72.0 * _EMU_PER_INCH)))
        # When the estimate already says the text does not fit, include the
        # clipped textbox height as part of the occupied region.  PowerPoint
        # paints the first line inside the shape and the remaining lines below
        # it, so the visual footprint is closer to shape-height + overflow.
        if layout and not self._ppt_text_frame_autofit_enabled(shape) and layout.get('estimated_lines', 0) > layout.get('capacity_lines', 0):
            text_height = max(text_height, inner_height + max(1, estimated_lines - layout['capacity_lines']) * line_height + int(line_height * 0.35))
        if self._ppt_text_frame_autofit_enabled(shape):
            text_height = min(inner_height, text_height)

        text_width = inner_width
        if layout:
            estimated_width = int(max(0.05, layout.get('max_line_width_in', 0.0)) * _EMU_PER_INCH)
            # If wrapping is disabled (or the XML explicitly allows overflow),
            # glyphs can extend horizontally beyond the textbox.  Even with wrap
            # enabled, keep a modest overhang tolerance for long CJK/Latin runs
            # because python-pptx cannot provide exact rendered glyph metrics.
            if not self._ppt_text_frame_word_wrap_enabled(shape):
                text_width = max(text_width, estimated_width)
            else:
                text_width = max(text_width, min(estimated_width, int(inner_width * 1.18)))

        x = inner_left
        try:
            paragraph = shape.text_frame.paragraphs[0] if shape.text_frame.paragraphs else None
            alignment = str(getattr(paragraph, 'alignment', '') or '').lower()
            if 'center' in alignment or 'centre' in alignment:
                x = inner_left + (inner_width - text_width) // 2
            elif 'right' in alignment:
                x = inner_right - text_width
        except Exception:
            pass

        y = inner_top
        try:
            anchor = str(getattr(shape.text_frame, 'vertical_anchor', '')).lower()
            if 'middle' in anchor or 'mid' in anchor:
                y = inner_top + max(0, (inner_height - text_height) // 2)
            elif 'bottom' in anchor:
                y = inner_bottom - text_height
        except Exception:
            pass
        return (x, y, x + text_width, y + text_height)

    def _ppt_find_text_overlaps(self, candidates):
        overlaps = []
        for i, first in enumerate(candidates):
            for second in candidates[i + 1:]:
                inter = self._ppt_bounds_intersection(first['bounds'], second['bounds'])
                if not inter:
                    continue
                inter_area = self._ppt_bounds_area(inter)
                if inter_area <= 0:
                    continue
                first_area = self._ppt_bounds_area(first['bounds'])
                second_area = self._ppt_bounds_area(second['bounds'])
                smaller = max(1, min(first_area, second_area))
                overlap_ratio = inter_area / smaller
                overflow_related = self._ppt_bounds_extends_outside(first['bounds'], first['shape_bounds']) or self._ppt_bounds_extends_outside(second['bounds'], second['shape_bounds'])
                left, top, right, bottom = inter
                overlap_w_in = (right - left) / _EMU_PER_INCH
                overlap_h_in = (bottom - top) / _EMU_PER_INCH
                overlap_area_in = inter_area / (_EMU_PER_INCH * _EMU_PER_INCH)
                if overlap_ratio < (0.12 if overflow_related else 0.22) or overlap_area_in < (0.015 if overflow_related else 0.035):
                    continue
                if overlap_w_in < (0.08 if overflow_related else 0.12) or overlap_h_in < (0.035 if overflow_related else 0.08):
                    continue
                overlaps.append({
                    'overlap_pct': int(round(overlap_ratio * 100)),
                    'overlap_area_sq_in': round(overlap_area_in, 3),
                    'bounds_a': self._ppt_format_bounds(first['shape_bounds']),
                    'bounds_b': self._ppt_format_bounds(second['shape_bounds']),
                    'overlap_bounds': self._ppt_format_bounds(inter),
                    'text_bounds_a': self._ppt_format_bounds(first['bounds']),
                    'text_bounds_b': self._ppt_format_bounds(second['bounds']),
                    'text_a': first.get('text', ''),
                    'text_b': second.get('text', ''),
                    'bounds_source_a': first.get('bounds_source', 'estimated'),
                    'bounds_source_b': second.get('bounds_source', 'estimated'),
                })
        overlaps.sort(key=lambda item: (-item['overlap_pct'], -item['overlap_area_sq_in']))
        return overlaps[:10]

    def _ppt_bounds_extends_outside(self, outer, inner, tolerance=9144):
        try:
            if not outer or not inner:
                return False
            return (
                outer[0] < inner[0] - tolerance or outer[1] < inner[1] - tolerance or
                outer[2] > inner[2] + tolerance or outer[3] > inner[3] + tolerance
            )
        except Exception:
            return False

    def _ppt_bounds_intersection(self, a, b):
        left = max(a[0], b[0])
        top = max(a[1], b[1])
        right = min(a[2], b[2])
        bottom = min(a[3], b[3])
        if right <= left or bottom <= top:
            return None
        return (left, top, right, bottom)

    def _ppt_bounds_area(self, bounds):
        try:
            return max(0, bounds[2] - bounds[0]) * max(0, bounds[3] - bounds[1])
        except Exception:
            return 0

    def _ppt_format_bounds(self, bounds):
        try:
            left, top, right, bottom = bounds
            return {
                'left_in': round(left / 914400.0, 2),
                'top_in': round(top / 914400.0, 2),
                'right_in': round(right / 914400.0, 2),
                'bottom_in': round(bottom / 914400.0, 2),
            }
        except Exception:
            return {}

    def _ppt_text_frame_autofit_enabled(self, shape):
        try:
            auto_size = getattr(shape.text_frame, 'auto_size', None)
            return 'TEXT_TO_FIT_SHAPE' in str(auto_size).upper()
        except Exception:
            return False

    def _ppt_text_frame_word_wrap_enabled(self, shape):
        try:
            word_wrap = getattr(shape.text_frame, 'word_wrap', None)
            return word_wrap is not False
        except Exception:
            return True

    def _ppt_inner_text_bounds(self, shape, bounds):
        left, top, right, bottom = bounds
        try:
            tf = shape.text_frame
            ml = int(getattr(tf, 'margin_left', 0) or 0)
            mr = int(getattr(tf, 'margin_right', 0) or 0)
            mt = int(getattr(tf, 'margin_top', 0) or 0)
            mb = int(getattr(tf, 'margin_bottom', 0) or 0)
        except Exception:
            ml = mr = mt = mb = 0

        inner_left = min(right, left + max(0, ml))
        inner_right = max(inner_left, right - max(0, mr))
        inner_top = min(bottom, top + max(0, mt))
        inner_bottom = max(inner_top, bottom - max(0, mb))
        return inner_left, inner_top, inner_right, inner_bottom

    def _ppt_text_visual_units(self, text):
        units = 0.0
        for ch in _as_text(text):
            if ch in '\r\n':
                continue
            if ch.isspace():
                units += 0.35
            elif '\u2e80' <= ch <= '\u9fff' or '\uff00' <= ch <= '\uffef':
                units += 1.85
            elif ch in '.,;:!?，。；：！？、()[]{}<>《》“”‘’"\'`':
                units += 0.6
            elif ch.isupper():
                units += 1.08
            else:
                units += 1.0
        return max(1.0, units)

    def _ppt_estimate_text_layout(self, shape, text, bounds=None):
        if not text or not self._ppt_has(shape, 'has_text_frame'):
            return None
        bounds = bounds or self._ppt_shape_bounds(shape)
        if not bounds:
            return None
        inner_left, inner_top, inner_right, inner_bottom = self._ppt_inner_text_bounds(shape, bounds)
        width_in = max(0.05, (inner_right - inner_left) / _EMU_PER_INCH)
        height_in = max(0.05, (inner_bottom - inner_top) / _EMU_PER_INCH)
        font_size = self._ppt_shape_font_size(shape) or 14
        base_char_width_in = max(0.045, font_size * 0.0062)
        units_per_line = max(1.0, width_in / base_char_width_in)
        estimated_lines = 0
        max_line_units = 1.0
        wrap_enabled = self._ppt_text_frame_word_wrap_enabled(shape)
        for raw in _as_text(text).splitlines() or ['']:
            raw = raw.strip()
            visual_units = self._ppt_text_visual_units(raw)
            max_line_units = max(max_line_units, visual_units if not wrap_enabled else min(visual_units, units_per_line))
            estimated_lines += max(1, math.ceil(visual_units / units_per_line)) if wrap_enabled else 1
        # Use a slightly conservative line-height estimate.  The validator must
        # catch visual collisions; underestimating by a few points can hide the
        # common case where glyphs spill below a too-short textbox.
        line_height_in = max(0.12, font_size * 1.34 / 72.0)
        capacity_lines = max(1, int(height_in / line_height_in))
        return {
            'estimated_lines': int(estimated_lines),
            'capacity_lines': int(capacity_lines),
            'font_size': float(font_size),
            'line_height_in': line_height_in,
            'height_in': height_in,
            'max_line_width_in': max_line_units * base_char_width_in,
        }

    def _ppt_estimate_text_overflow(self, shape, text):
        if not text or not self._ppt_has(shape, 'has_text_frame'):
            return None
        if self._ppt_text_frame_autofit_enabled(shape):
            return None
        layout = self._ppt_estimate_text_layout(shape, text)
        if not layout:
            return None
        estimated_lines = layout['estimated_lines']
        capacity_lines = layout['capacity_lines']
        estimated_height = estimated_lines * layout['line_height_in']
        if len(_as_text(text).strip()) > 4 and estimated_lines > capacity_lines and estimated_height > layout['height_in'] * 1.08:
            return {
                'estimated_lines': estimated_lines,
                'capacity_lines': capacity_lines,
                'font_size': layout['font_size'],
            }
        return None

    def _ppt_shape_font_size(self, shape):
        try:
            for paragraph in shape.text_frame.paragraphs:
                if paragraph.font.size is not None:
                    return float(paragraph.font.size.pt)
                for run in paragraph.runs:
                    if run.font.size is not None:
                        return float(run.font.size.pt)
        except Exception:
            pass
        return None

    def _iter_ppt_shapes(self, shapes):
        for shape in shapes:
            yield shape
            nested = getattr(shape, 'shapes', None)
            if nested is not None:
                yield from self._iter_ppt_shapes(nested)

    def _ppt_shape_text(self, shape):
        parts = []
        if self._ppt_has(shape, 'has_text_frame'):
            parts.append(self._ppt_text_frame_text(shape.text_frame))
        if self._ppt_has(shape, 'has_table'):
            try:
                for row in shape.table.rows:
                    for cell in row.cells:
                        text = self._ppt_text_frame_text(cell.text_frame)
                        if text:
                            parts.append(text)
            except Exception:
                pass
        return '\n'.join(x for x in parts if x)

    def _ppt_text_frame_text(self, text_frame):
        lines = []
        try:
            for p in text_frame.paragraphs:
                text = _as_text(getattr(p, 'text', '')).strip()
                if text:
                    lines.append(text)
        except Exception:
            pass
        return '\n'.join(lines)

    def _ppt_collect_text_styles(self, shape, font_counts, font_size_counts, color_counts):
        frames = []
        if self._ppt_has(shape, 'has_text_frame'):
            frames.append(shape.text_frame)
        if self._ppt_has(shape, 'has_table'):
            try:
                for row in shape.table.rows:
                    for cell in row.cells:
                        frames.append(cell.text_frame)
            except Exception:
                pass
        for frame in frames:
            try:
                for paragraph in frame.paragraphs:
                    self._ppt_count_font(paragraph.font, font_counts, font_size_counts, color_counts)
                    for run in paragraph.runs:
                        self._ppt_count_font(run.font, font_counts, font_size_counts, color_counts)
            except Exception:
                continue

    def _ppt_count_font(self, font, font_counts, font_size_counts, color_counts):
        try:
            name = getattr(font, 'name', None)
            if name:
                font_counts[_as_text(name)] = font_counts.get(_as_text(name), 0) + 1
        except Exception:
            pass
        try:
            size = getattr(font, 'size', None)
            if size is not None:
                key = str(int(round(size.pt)))
                font_size_counts[key] = font_size_counts.get(key, 0) + 1
        except Exception:
            pass
        try:
            rgb = font.color.rgb
            color = self._ppt_rgb_hex(rgb)
            if color:
                color_counts[color] = color_counts.get(color, 0) + 1
        except Exception:
            pass

    def _ppt_default_profile_path(self, ppt_path):
        base = os.path.splitext(os.path.basename(ppt_path))[0] or 'template'
        base = _SAFE_FILENAME_RE.sub('_', base).strip('._ ') or 'template'
        digest = hashlib.sha1(os.path.abspath(ppt_path).encode('utf-8', 'ignore')).hexdigest()[:8]
        return os.path.join('output', 'ppt_templates', f'{base}_{digest}.profile.json')

    def _analyze_ppt_template_file(self, ppt_path):
        from pptx import Presentation
        prs = Presentation(ppt_path)
        slide_w = int(prs.slide_width)
        slide_h = int(prs.slide_height)
        color_counts = Counter()
        fill_counts = Counter()
        font_color_counts = Counter()
        font_counts = Counter()
        font_size_counts = Counter()
        title_boxes = []
        content_boxes = []
        footer_boxes = []
        repeated = {}
        decorations = []
        warnings = []

        theme_colors = self._extract_ppt_theme_colors(ppt_path)
        for c in theme_colors.values():
            if c:
                color_counts[c] += 3

        for slide in prs.slides:
            for shape in self._iter_ppt_shapes(slide.shapes):
                bounds = self._ppt_shape_bounds(shape)
                if not bounds:
                    continue
                left, top, right, bottom = bounds
                width, height = right - left, bottom - top
                text = self._ppt_shape_text(shape).strip()
                font_size = self._ppt_shape_font_size(shape)
                if font_size:
                    font_size_counts[str(int(round(font_size)))] += 1
                self._collect_template_shape_colors(shape, color_counts, fill_counts, font_color_counts)
                self._collect_template_shape_fonts(shape, font_counts)

                rel = {'left': _emu_to_in(left), 'top': _emu_to_in(top), 'width': _emu_to_in(width), 'height': _emu_to_in(height)}
                if text and (top < slide_h * 0.24 or (font_size and font_size >= 22)):
                    title_boxes.append(rel)
                elif top > slide_h * 0.82:
                    footer_boxes.append(rel)
                elif width > slide_w * 0.12 and height > slide_h * 0.04:
                    content_boxes.append(rel)

                sig = self._template_shape_signature(shape, bounds, text)
                if sig:
                    repeated.setdefault(sig, {'count': 0, 'shape': shape, 'bounds': bounds, 'text': text})['count'] += 1

        slide_count = max(1, len(prs.slides))
        for item in repeated.values():
            if item['count'] >= max(2, min(slide_count, 3)) or (slide_count == 1 and self._is_template_decoration(item['shape'], item['bounds'], slide_w, slide_h, item['text'])):
                deco = self._extract_template_decoration(item['shape'], item['bounds'], item['text'])
                if deco:
                    decorations.append(deco)

        primary = self._pick_color(color_counts, exclude_light=True) or theme_colors.get('accent1') or '#2563eb'
        bg = self._pick_color(fill_counts, prefer_light=True) or theme_colors.get('lt1') or '#f8fafc'
        text_color = self._pick_color(font_color_counts, exclude_light=True) or theme_colors.get('dk1') or '#0f172a'
        accent = self._pick_color(color_counts, exclude=[primary, bg, text_color], exclude_light=True) or theme_colors.get('accent2') or '#f97316'
        font_family = font_counts.most_common(1)[0][0] if font_counts else 'Microsoft YaHei'
        title_size = self._pick_font_size(font_size_counts, default=28, high=True)
        body_size = self._pick_font_size(font_size_counts, default=16, high=False)
        layout = {
            'title_box': self._median_box(title_boxes) or {'left': 0.55, 'top': 0.3, 'width': max(1, _emu_to_in(slide_w) - 1.1), 'height': 0.6},
            'content_box': self._median_box(content_boxes) or {'left': 0.8, 'top': 1.35, 'width': max(1, _emu_to_in(slide_w) - 1.6), 'height': max(1, _emu_to_in(slide_h) - 2.0)},
            'footer_box': self._median_box(footer_boxes) or {'left': 0.8, 'top': max(0, _emu_to_in(slide_h) - 0.45), 'width': max(1, _emu_to_in(slide_w) - 1.6), 'height': 0.25},
        }
        if not decorations:
            warnings.append('no repeated header/footer/logo decorations detected; style colors and layout will still be applied')
        return {
            'version': 1,
            'source': self._ppt_rel_path(ppt_path),
            'slide_width': slide_w,
            'slide_height': slide_h,
            'size': {'width_in': _emu_to_in(slide_w), 'height_in': _emu_to_in(slide_h), 'aspect_ratio': round(slide_w / float(slide_h), 4) if slide_h else 0},
            'theme': {'primary_color': primary, 'accent_color': accent, 'background_color': bg, 'text_color': text_color, 'muted_color': text_color, 'line_color': self._pick_color(color_counts, exclude=[primary, bg, text_color, accent]) or '#cbd5e1'},
            'font': {'font_family': font_family, 'title_font_size': title_size, 'body_font_size': body_size},
            'layout': layout,
            'decorations': decorations[:24],
            'stats': {'colors': self._counter_top(color_counts), 'fonts': self._counter_top(font_counts), 'font_sizes': self._counter_top(font_size_counts)},
            'warnings': warnings,
        }

    def _extract_ppt_theme_colors(self, ppt_path):
        out = {}
        ns = {'a': 'http://schemas.openxmlformats.org/drawingml/2006/main'}
        try:
            with zipfile.ZipFile(ppt_path) as zf:
                names = [n for n in zf.namelist() if n.startswith('ppt/theme/theme') and n.endswith('.xml')]
                if not names:
                    return out
                root = ET.fromstring(zf.read(names[0]))
            for node in root.findall('.//a:clrScheme/*', ns):
                srgb = node.find('.//a:srgbClr', ns)
                if srgb is not None and srgb.get('val'):
                    out[node.tag.split('}')[-1]] = '#' + srgb.get('val').lower()
        except Exception:
            pass
        return out

    def _collect_template_shape_colors(self, shape, all_counts, fill_counts, font_color_counts):
        for getter, counts in ((lambda s: s.fill.fore_color.rgb, fill_counts), (lambda s: s.line.color.rgb, all_counts)):
            try:
                color = self._ppt_rgb_hex(getter(shape))
                if color:
                    counts[color] += 1
                    all_counts[color] += 1
            except Exception:
                pass
        try:
            if self._ppt_has(shape, 'has_text_frame'):
                for p in shape.text_frame.paragraphs:
                    for font in [p.font] + [r.font for r in p.runs]:
                        try:
                            color = self._ppt_rgb_hex(font.color.rgb)
                            if color:
                                font_color_counts[color] += 1
                                all_counts[color] += 1
                        except Exception:
                            pass
        except Exception:
            pass

    def _collect_template_shape_fonts(self, shape, font_counts):
        try:
            if not self._ppt_has(shape, 'has_text_frame'):
                return
            for p in shape.text_frame.paragraphs:
                for font in [p.font] + [r.font for r in p.runs]:
                    name = getattr(font, 'name', None)
                    if name:
                        font_counts[_as_text(name)] += 1
        except Exception:
            pass

    def _counter_top(self, counter, limit=8):
        return [{'value': k, 'count': v} for k, v in counter.most_common(limit)]

    def _color_luma(self, color):
        rgb = self._ppt_rgb_tuple(color)
        if not rgb:
            return 128
        r, g, b = rgb
        return 0.2126 * r + 0.7152 * g + 0.0722 * b

    def _pick_color(self, counter, exclude=None, exclude_light=False, prefer_light=False):
        exclude = {x.lower() for x in _as_list(exclude) if x}
        for color, _ in counter.most_common(20):
            if color.lower() in exclude:
                continue
            luma = self._color_luma(color)
            if exclude_light and luma > 238:
                continue
            if prefer_light and luma < 210:
                continue
            return color
        return ''

    def _pick_font_size(self, counter, default=16, high=False):
        vals = []
        for k, count in counter.items():
            try:
                vals.extend([int(k)] * min(8, count))
            except Exception:
                pass
        if not vals:
            return default
        vals.sort()
        idx = int(len(vals) * (0.8 if high else 0.45))
        return max(8, min(48, vals[min(len(vals) - 1, idx)]))

    def _median_box(self, boxes):
        if not boxes:
            return None
        out = {}
        for key in ('left', 'top', 'width', 'height'):
            vals = sorted(float(b.get(key, 0)) for b in boxes)
            out[key] = round(vals[len(vals) // 2], 3)
        return out

    def _template_shape_signature(self, shape, bounds, text):
        left, top, right, bottom = bounds
        kind = 'pic' if self._ppt_shape_has_image(shape) else ('text' if text else 'shape')
        fill = ''
        try:
            fill = self._ppt_rgb_hex(shape.fill.fore_color.rgb)
        except Exception:
            pass
        return '|'.join([kind, str(round(left / 20000)), str(round(top / 20000)), str(round((right-left) / 20000)), str(round((bottom-top) / 20000)), text[:40], fill])

    def _is_template_decoration(self, shape, bounds, slide_w, slide_h, text):
        left, top, right, bottom = bounds
        width, height = right - left, bottom - top
        edge = top < slide_h * 0.18 or bottom > slide_h * 0.82 or left < slide_w * 0.08 or right > slide_w * 0.92
        small_text = bool(text) and len(text) <= 80
        return self._ppt_shape_has_image(shape) or edge or small_text or height < slide_h * 0.08 or width < slide_w * 0.08

    def _extract_template_decoration(self, shape, bounds, text):
        left, top, right, bottom = bounds
        item = {'kind': 'text' if text else 'shape', 'bounds': {'left': _emu_to_in(left), 'top': _emu_to_in(top), 'width': _emu_to_in(right-left), 'height': _emu_to_in(bottom-top)}}
        if self._ppt_shape_has_image(shape):
            try:
                digest = hashlib.sha1(shape.image.blob).hexdigest()[:12]
                ext = shape.image.ext or 'png'
                img_dir = os.path.join(config.WORKSPACE_ROOT, 'output', 'ppt_templates', 'assets')
                os.makedirs(img_dir, exist_ok=True)
                img_path = os.path.join(img_dir, f'{digest}.{ext}')
                if not os.path.exists(img_path):
                    with open(img_path, 'wb') as f:
                        f.write(shape.image.blob)
                item.update({'kind': 'image', 'path': self._ppt_rel_path(img_path)})
                return item
            except Exception:
                return None
        item['text'] = text
        try:
            item['fill'] = self._ppt_rgb_hex(shape.fill.fore_color.rgb)
        except Exception:
            pass
        try:
            item['line'] = self._ppt_rgb_hex(shape.line.color.rgb)
        except Exception:
            pass
        try:
            item['font_size'] = self._ppt_shape_font_size(shape)
        except Exception:
            pass
        return item if (item.get('text') or item.get('fill') or item.get('line')) else None

    def _render_ppt_preview_powerpoint(self, ppt_path, out_dir, width, height, max_slides=None):
        if os.name != 'nt':
            return [], 'native PowerPoint preview is only available on Windows'
        app = None
        pres = None
        coinit = False
        try:
            import pythoncom
            import win32com.client
            pythoncom.CoInitialize()
            coinit = True
            app = win32com.client.DispatchEx('PowerPoint.Application')
            try:
                app.DisplayAlerts = 0
            except Exception:
                pass
            pres = app.Presentations.Open(os.path.abspath(ppt_path), ReadOnly=True, Untitled=False, WithWindow=False)
            count = int(pres.Slides.Count)
            if max_slides:
                count = min(count, int(max_slides))
            images = []
            for i in range(1, count + 1):
                out = os.path.join(out_dir, f'slide_{i:03d}.png')
                pres.Slides(i).Export(out, 'PNG', int(width), int(height))
                if os.path.exists(out):
                    images.append(out)
            return images, ''
        except Exception as e:
            return [], f'PowerPoint native render unavailable: {e}'
        finally:
            try:
                if pres is not None:
                    pres.Close()
            except Exception:
                pass
            try:
                if app is not None:
                    app.Quit()
            except Exception:
                pass
            if coinit:
                try:
                    pythoncom.CoUninitialize()
                except Exception:
                    pass

    def _render_ppt_preview_pillow(self, ppt_path, out_dir, width, height, max_slides=None):
        try:
            from PIL import Image, ImageDraw
            from pptx import Presentation
        except ImportError as e:
            return [], f'pillow fallback unavailable: {e}'

        try:
            prs = Presentation(ppt_path)
            slide_w = max(1, int(prs.slide_width))
            slide_h = max(1, int(prs.slide_height))
            if not height:
                height = max(180, int(width * slide_h / slide_w))
            sx = width / float(slide_w)
            sy = height / float(slide_h)
            px_per_pt = width / (slide_w / 914400.0) / 72.0
            images = []
            for idx, slide in enumerate(prs.slides, start=1):
                if max_slides and idx > max_slides:
                    break
                img = Image.new('RGB', (width, height), (255, 255, 255))
                draw = ImageDraw.Draw(img)
                for shape in self._iter_ppt_shapes(slide.shapes):
                    self._draw_ppt_preview_shape(img, draw, shape, sx, sy, px_per_pt)
                out = os.path.join(out_dir, f'slide_{idx:03d}.png')
                img.save(out, 'PNG')
                images.append(out)
            return images, ''
        except Exception as e:
            return [], f'pillow fallback failed: {e}'

    def _draw_ppt_preview_shape(self, canvas, draw, shape, sx, sy, px_per_pt):
        x = int((getattr(shape, 'left', 0) or 0) * sx)
        y = int((getattr(shape, 'top', 0) or 0) * sy)
        w = int((getattr(shape, 'width', 0) or 0) * sx)
        h = int((getattr(shape, 'height', 0) or 0) * sy)
        if w <= 1 or h <= 1:
            return

        if self._ppt_shape_has_image(shape):
            if self._draw_ppt_preview_image(canvas, draw, shape, x, y, w, h):
                return

        if self._ppt_has(shape, 'has_table'):
            self._draw_ppt_preview_table(draw, shape, x, y, w, h, sx, sy, px_per_pt)
            return

        fill = self._ppt_shape_fill_tuple(shape)
        line = self._ppt_shape_line_tuple(shape)
        if fill is not None:
            try:
                draw.rectangle([x, y, x + w, y + h], fill=fill, outline=line)
            except Exception:
                pass
        elif line is not None:
            try:
                draw.rectangle([x, y, x + w, y + h], outline=line)
            except Exception:
                pass

        if self._ppt_has(shape, 'has_chart'):
            font = self._ppt_preview_font(18)
            draw.text((x + 8, y + 8), 'CHART', fill=(71, 85, 105), font=font)

        if self._ppt_has(shape, 'has_text_frame'):
            self._draw_ppt_preview_text_frame(draw, shape.text_frame, x, y, w, h, px_per_pt)

    def _draw_ppt_preview_image(self, canvas, draw, shape, x, y, w, h):
        try:
            from PIL import Image
            img = Image.open(BytesIO(shape.image.blob)).convert('RGB')
            img.thumbnail((max(1, w), max(1, h)))
            ox = x + max(0, int((w - img.width) / 2))
            oy = y + max(0, int((h - img.height) / 2))
            canvas.paste(img, (ox, oy))
            return True
        except Exception:
            try:
                draw.rectangle([x, y, x + w, y + h], fill=(226, 232, 240), outline=(148, 163, 184))
                draw.text((x + 8, y + 8), 'IMAGE', fill=(71, 85, 105), font=self._ppt_preview_font(16))
            except Exception:
                pass
            return False

    def _draw_ppt_preview_table(self, draw, shape, x, y, w, h, sx, sy, px_per_pt):
        try:
            table = shape.table
            rows = list(table.rows)
            cols = list(table.columns)
            if not rows or not cols:
                return
            col_widths = [max(1, int(col.width * sx)) for col in cols]
            row_heights = [max(1, int(row.height * sy)) for row in rows]
            total_w = sum(col_widths) or 1
            total_h = sum(row_heights) or 1
            if abs(total_w - w) > 3:
                col_widths = [max(1, int(v * w / total_w)) for v in col_widths]
            if abs(total_h - h) > 3:
                row_heights = [max(1, int(v * h / total_h)) for v in row_heights]
            cy = y
            for r, row in enumerate(rows):
                cx = x
                rh = row_heights[r]
                for c, cell in enumerate(row.cells):
                    cw = col_widths[c]
                    fill = self._ppt_cell_fill_tuple(cell) or ((226, 232, 240) if r == 0 else (255, 255, 255))
                    draw.rectangle([cx, cy, cx + cw, cy + rh], fill=fill, outline=(203, 213, 225))
                    self._draw_ppt_preview_text_frame(draw, cell.text_frame, cx + 3, cy + 2, max(1, cw - 6), max(1, rh - 4), px_per_pt, default_pt=10)
                    cx += cw
                cy += rh
        except Exception:
            draw.rectangle([x, y, x + w, y + h], fill=(248, 250, 252), outline=(203, 213, 225))

    def _draw_ppt_preview_text_frame(self, draw, text_frame, x, y, w, h, px_per_pt, default_pt=14):
        cursor_y = y + 4
        max_y = y + h - 2
        try:
            paragraphs = list(text_frame.paragraphs)
        except Exception:
            return
        for paragraph in paragraphs:
            text = _as_text(getattr(paragraph, 'text', '')).strip()
            if not text:
                cursor_y += 4
                continue
            pt, family, color = self._ppt_paragraph_font_info(paragraph, default_pt)
            font_size = max(7, min(64, int(round(pt * px_per_pt))))
            font = self._ppt_preview_font(font_size, family)
            color = color or (15, 23, 42)
            level = max(0, min(6, int(getattr(paragraph, 'level', 0) or 0)))
            indent = level * max(10, int(font_size * 0.8))
            available_w = max(20, w - indent - 8)
            line_h = max(font_size + 3, int(font_size * 1.25))
            for line in self._ppt_wrap_text(draw, text, font, available_w):
                if cursor_y + line_h > max_y:
                    draw.text((x + indent + 4, max(y, max_y - line_h)), '...', fill=color, font=font)
                    return
                draw.text((x + indent + 4, cursor_y), line, fill=color, font=font)
                cursor_y += line_h
            cursor_y += max(1, int(line_h * 0.25))

    def _ppt_paragraph_font_info(self, paragraph, default_pt=14):
        fonts = []
        try:
            fonts.append(paragraph.font)
        except Exception:
            pass
        try:
            fonts.extend(run.font for run in paragraph.runs)
        except Exception:
            pass
        size_pt = None
        family = None
        color = None
        for font in fonts:
            if size_pt is None:
                try:
                    if font.size is not None:
                        size_pt = float(font.size.pt)
                except Exception:
                    pass
            if family is None:
                try:
                    if font.name:
                        family = _as_text(font.name)
                except Exception:
                    pass
            if color is None:
                try:
                    color = self._ppt_rgb_tuple(font.color.rgb)
                except Exception:
                    pass
        return size_pt or default_pt, family, color

    def _write_ppt_preview_html(self, out_dir, ppt_rel, validation, images, renderer, renderer_note):
        path = os.path.join(out_dir, 'index.html')
        issue_items = validation.get('issues') or []
        warning_items = validation.get('warnings') or []
        slide_items = validation.get('slides') or []

        def esc(value):
            return html.escape(_as_text(value), quote=True)

        issue_html = ''.join(
            f'<li><strong>{esc(x.get("severity", "issue"))}</strong> {esc(x.get("message", ""))}</li>'
            for x in issue_items
        ) or '<li>None</li>'
        warning_html = ''.join(
            f'<li><strong>slide {esc(x.get("slide", ""))}</strong> {esc(x.get("message", ""))}</li>'
            for x in warning_items
        ) or '<li>None</li>'
        image_html = ''.join(
            '<section class="slide">'
            f'<h2>Slide {img["slide"]}</h2>'
            f'<img src="{esc(img["url"])}" alt="Slide {img["slide"]} preview">'
            '</section>'
            for img in images
        ) or '<p>No preview images were generated.</p>'
        toc_html = ''.join(
            f'<li>{esc(s.get("slide"))}. {esc(s.get("title") or "(untitled)")} '
            f'<span>{esc(s.get("char_count"))} chars</span></li>'
            for s in slide_items
        )
        status = 'PASS' if validation.get('passed') else 'FAIL'
        content = f'''<!doctype html>
<html lang="en">
<head>
  <meta charset="utf-8">
  <meta name="viewport" content="width=device-width, initial-scale=1">
  <title>PPT Preview - {esc(os.path.basename(ppt_rel))}</title>
  <style>
    body {{ margin: 0; font-family: system-ui, -apple-system, "Segoe UI", sans-serif; background: #f8fafc; color: #0f172a; }}
    header {{ padding: 20px 24px; background: #0f172a; color: white; }}
    main {{ max-width: 1180px; margin: 0 auto; padding: 20px; }}
    .meta, .report {{ background: white; border: 1px solid #e2e8f0; border-radius: 8px; padding: 16px; margin-bottom: 16px; }}
    .status {{ display: inline-block; padding: 3px 10px; border-radius: 999px; background: {'#dcfce7' if validation.get('passed') else '#fee2e2'}; color: {'#166534' if validation.get('passed') else '#991b1b'}; font-weight: 700; }}
    .slide {{ margin: 18px 0 28px; }}
    .slide h2 {{ font-size: 15px; margin: 0 0 8px; color: #334155; }}
    .slide img {{ width: 100%; border: 1px solid #cbd5e1; border-radius: 6px; background: white; display: block; }}
    ul {{ margin: 8px 0 0; padding-left: 22px; }}
    li span {{ color: #64748b; margin-left: 8px; }}
    code {{ background: #e2e8f0; padding: 1px 5px; border-radius: 4px; }}
  </style>
</head>
<body>
  <header>
    <h1>PPT Preview</h1>
    <div><code>{esc(ppt_rel)}</code></div>
  </header>
  <main>
    <section class="meta">
      <p><span class="status">{status}</span> Score: {esc(validation.get('score'))} / Renderer: {esc(renderer or 'none')}</p>
      <p>{esc(renderer_note or '')}</p>
      <ul>{toc_html}</ul>
    </section>
    <section class="report">
      <h2>Issues</h2>
      <ul>{issue_html}</ul>
      <h2>Warnings</h2>
      <ul>{warning_html}</ul>
    </section>
    {image_html}
  </main>
</body>
</html>'''
        with open(path, 'w', encoding='utf-8') as f:
            f.write(content)
        return path

    def _ppt_wrap_text(self, draw, text, font, max_width):
        lines = []
        for raw in _as_text(text).splitlines() or ['']:
            if not raw:
                lines.append('')
                continue
            tokens = re.findall(r'\S+\s*', raw) if re.search(r'\s', raw) else list(raw)
            current = ''
            for token in tokens:
                candidate = current + token
                if self._ppt_text_width(draw, candidate, font) <= max_width:
                    current = candidate
                    continue
                if current:
                    lines.append(current.rstrip())
                    current = ''
                if self._ppt_text_width(draw, token, font) <= max_width:
                    current = token.lstrip()
                    continue
                for ch in token:
                    candidate = current + ch
                    if self._ppt_text_width(draw, candidate, font) <= max_width:
                        current = candidate
                    else:
                        if current:
                            lines.append(current.rstrip())
                        current = ch
            if current:
                lines.append(current.rstrip())
        return lines

    def _ppt_text_width(self, draw, text, font):
        try:
            box = draw.textbbox((0, 0), text, font=font)
            return box[2] - box[0]
        except Exception:
            return len(_as_text(text)) * 8

    def _ppt_preview_font(self, size, family=None):
        try:
            from PIL import ImageFont
        except Exception:
            return None
        size = max(7, min(96, int(size or 14)))
        cache = getattr(self, '_ppt_preview_font_cache', None)
        if cache is None:
            cache = {}
            self._ppt_preview_font_cache = cache
        key = (size, family or '')
        if key in cache:
            return cache[key]
        candidates = []
        if os.name == 'nt':
            fonts_dir = os.path.join(os.environ.get('WINDIR', r'C:\Windows'), 'Fonts')
            candidates.extend([
                os.path.join(fonts_dir, 'msyh.ttc'),
                os.path.join(fonts_dir, 'simhei.ttf'),
                os.path.join(fonts_dir, 'simsun.ttc'),
                os.path.join(fonts_dir, 'arial.ttf'),
            ])
        candidates.extend([
            '/usr/share/fonts/truetype/noto/NotoSansCJK-Regular.ttc',
            '/usr/share/fonts/opentype/noto/NotoSansCJK-Regular.ttc',
            '/System/Library/Fonts/PingFang.ttc',
            '/Library/Fonts/Arial Unicode.ttf',
        ])
        for candidate in candidates:
            try:
                if os.path.exists(candidate):
                    font = ImageFont.truetype(candidate, size)
                    cache[key] = font
                    return font
            except Exception:
                continue
        font = ImageFont.load_default()
        cache[key] = font
        return font

    def _ppt_has(self, shape, attr):
        try:
            return bool(getattr(shape, attr, False))
        except Exception:
            return False

    def _ppt_shape_has_image(self, shape):
        try:
            return hasattr(shape, 'image') and shape.image is not None
        except Exception:
            return False

    def _ppt_shape_looks_like_title(self, shape, text, slide_h):
        try:
            if getattr(shape, 'is_placeholder', False):
                ph_type = str(shape.placeholder_format.type).lower()
                if 'title' in ph_type:
                    return True
        except Exception:
            pass
        try:
            top = int(getattr(shape, 'top', 0) or 0)
        except Exception:
            top = 0
        return top <= slide_h * 0.22 and len(text) <= 140

    def _ppt_shape_area_sq_in(self, shape):
        try:
            w = max(0, int(getattr(shape, 'width', 0) or 0))
            h = max(0, int(getattr(shape, 'height', 0) or 0))
            return (w / 914400.0) * (h / 914400.0)
        except Exception:
            return 0

    def _ppt_shape_fill_tuple(self, shape):
        try:
            return self._ppt_rgb_tuple(shape.fill.fore_color.rgb)
        except Exception:
            return None

    def _ppt_shape_line_tuple(self, shape):
        try:
            return self._ppt_rgb_tuple(shape.line.color.rgb)
        except Exception:
            return None

    def _ppt_cell_fill_tuple(self, cell):
        try:
            return self._ppt_rgb_tuple(cell.fill.fore_color.rgb)
        except Exception:
            return None

    def _ppt_rgb_tuple(self, rgb):
        if rgb is None:
            return None
        try:
            values = tuple(int(x) for x in rgb)
            if len(values) >= 3:
                return values[:3]
        except Exception:
            pass
        text = str(rgb).strip().lstrip('#')
        if re.fullmatch(r'[0-9a-fA-F]{6}', text):
            return tuple(int(text[i:i + 2], 16) for i in (0, 2, 4))
        return None

    def _ppt_rgb_hex(self, rgb):
        values = self._ppt_rgb_tuple(rgb)
        if not values:
            return ''
        return '#%02x%02x%02x' % values

    def _ppt_top_counts(self, counts, limit=8):
        return [
            {'value': key, 'count': value}
            for key, value in sorted(counts.items(), key=lambda item: item[1], reverse=True)[:limit]
        ]

    def _ppt_int(self, value, default=0, low=None, high=None):
        try:
            out = int(value)
        except Exception:
            out = int(default)
        if low is not None:
            out = max(low, out)
        if high is not None:
            out = min(high, out)
        return out

    def _ppt_bool(self, value, default=False):
        if value is None:
            return default
        if isinstance(value, bool):
            return value
        if isinstance(value, (int, float)):
            return bool(value)
        text = _as_text(value).strip().lower()
        if text in ('1', 'true', 'yes', 'y', 'on'):
            return True
        if text in ('0', 'false', 'no', 'n', 'off'):
            return False
        return default

    def _ppt_rel_path(self, abs_path):
        return os.path.relpath(abs_path, config.WORKSPACE_ROOT).replace(os.sep, '/')

    def _ppt_preview_url(self, rel_path):
        return '/preview-file?path=' + quote(_as_text(rel_path).replace('\\', '/'))

    def _load_ppt_template_profile(self, data):
        if not isinstance(data, dict):
            return {}
        profile = data.get('template_profile')
        if isinstance(profile, dict):
            return dict(profile)
        profile_path = data.get('template_profile_path') or (profile if isinstance(profile, str) else '')
        if profile_path:
            path, err = check_path_or_error(profile_path, must_exist=True)
            if err:
                raise ValueError(err)
            with open(path, 'r', encoding='utf-8') as f:
                loaded = json.load(f)
            if isinstance(loaded, dict):
                loaded['profile_path'] = self._ppt_rel_path(path)
                return loaded
        template_path = data.get('template_path') or data.get('template')
        if template_path:
            path, err = check_path_or_error(template_path, must_exist=True)
            if err:
                raise ValueError(err)
            cache_path = self._ppt_default_profile_path(path)
            abs_cache, cache_err = check_path_or_error(cache_path, must_exist=False)
            if not cache_err and os.path.exists(abs_cache):
                try:
                    with open(abs_cache, 'r', encoding='utf-8') as f:
                        loaded = json.load(f)
                    if isinstance(loaded, dict):
                        loaded['profile_path'] = self._ppt_rel_path(abs_cache)
                        return loaded
                except Exception:
                    pass
            loaded = self._analyze_ppt_template_file(path)
            if not cache_err:
                os.makedirs(os.path.dirname(abs_cache), exist_ok=True)
                with open(abs_cache, 'w', encoding='utf-8') as f:
                    json.dump(loaded, f, ensure_ascii=False, indent=2)
                loaded['profile_path'] = self._ppt_rel_path(abs_cache)
            return loaded
        return {}

    def _template_profile_style(self, profile):
        out = {}
        if not isinstance(profile, dict):
            return out
        out.update(self._style_dict(profile.get('theme')))
        font = self._style_dict(profile.get('font'))
        if font.get('font_family'):
            out['font_family'] = font.get('font_family')
        if font.get('body_font_size'):
            out['font_size'] = font.get('body_font_size')
        return out

    def _template_box(self, name):
        profile = getattr(self, '_ppt_template_profile', {})
        if not isinstance(profile, dict):
            return None
        box = self._style_dict(self._style_dict(profile.get('layout')).get(name))
        if not box:
            return None
        return (_in_to_emu(box.get('left')), _in_to_emu(box.get('top')), _in_to_emu(box.get('width')), _in_to_emu(box.get('height')))

    def _apply_template_decorations(self, slide, prs, include_title=False):
        profile = getattr(self, '_ppt_template_profile', {})
        if not isinstance(profile, dict):
            return
        for item in _as_list(profile.get('decorations')):
            if not isinstance(item, dict):
                continue
            bounds = self._style_dict(item.get('bounds'))
            left, top, width, height = _in_to_emu(bounds.get('left')), _in_to_emu(bounds.get('top')), _in_to_emu(bounds.get('width')), _in_to_emu(bounds.get('height'))
            if width <= 0 or height <= 0:
                continue
            if not include_title and top < int(prs.slide_height * 0.18) and item.get('text'):
                continue
            kind = item.get('kind')
            try:
                if kind == 'image' and item.get('path'):
                    img_path, err = check_path_or_error(item.get('path'), must_exist=True)
                    if not err and os.path.exists(img_path):
                        slide.shapes.add_picture(img_path, left, top, width, height)
                elif item.get('text'):
                    box = self._add_textbox(slide, item.get('text'), left, top, width, height, size=item.get('font_size') or 9, color=self._parse_color(item.get('line') or item.get('fill'), self._theme()['muted']), style={'font_size': item.get('font_size') or 9})
                    try:
                        box.fill.background(); box.line.fill.background()
                    except Exception:
                        pass
                else:
                    from pptx.enum.shapes import MSO_SHAPE
                    shape = slide.shapes.add_shape(MSO_SHAPE.RECTANGLE, left, top, width, height)
                    fill = self._parse_color(item.get('fill'), None)
                    line = self._parse_color(item.get('line'), None)
                    if fill:
                        shape.fill.solid(); shape.fill.fore_color.rgb = fill
                    else:
                        shape.fill.background()
                    if line:
                        shape.line.color.rgb = line
                    else:
                        shape.line.fill.background()
            except Exception:
                continue

    def _ppt_add_cover(self, prs, data):
        self._ppt_current_slide = data if isinstance(data, dict) else {}
        slide = prs.slides.add_slide(prs.slide_layouts[0])
        self._apply_template_decorations(slide, prs, include_title=True)
        slide.shapes.title.text = _as_text(data.get('title'), '未命名 PPT')
        self._configure_text_frame(slide.shapes.title.text_frame)
        self._apply_text_style(slide.shapes.title.text_frame.paragraphs[0], self._style_for(data, 'title'))
        if len(slide.placeholders) > 1:
            slide.placeholders[1].text = _as_text(data.get('subtitle') or data.get('desc'))
            self._configure_text_frame(slide.placeholders[1].text_frame)
            self._apply_text_style(slide.placeholders[1].text_frame.paragraphs[0], self._style_for(data, 'subtitle'))

    def _ppt_add_section(self, prs, data):
        self._ppt_current_slide = data if isinstance(data, dict) else {}
        layout = prs.slide_layouts[2] if len(prs.slide_layouts) > 2 else prs.slide_layouts[0]
        slide = prs.slides.add_slide(layout)
        self._apply_template_decorations(slide, prs)
        slide.shapes.title.text = _as_text(data.get('title'), '章节')
        self._configure_text_frame(slide.shapes.title.text_frame)
        self._apply_text_style(slide.shapes.title.text_frame.paragraphs[0], self._style_for(data, 'title'))
        if len(slide.placeholders) > 1:
            slide.placeholders[1].text = _as_text(data.get('subtitle') or data.get('desc'))
            self._configure_text_frame(slide.placeholders[1].text_frame)
            self._apply_text_style(slide.placeholders[1].text_frame.paragraphs[0], self._style_for(data, 'subtitle'))

    def _ppt_add_bullets(self, prs, data, Pt):
        self._ppt_current_slide = data if isinstance(data, dict) else {}
        slide = prs.slides.add_slide(prs.slide_layouts[1])
        self._apply_template_decorations(slide, prs)
        slide.shapes.title.text = _as_text(data.get('title'), '未命名页面')
        self._configure_text_frame(slide.shapes.title.text_frame)
        self._apply_text_style(slide.shapes.title.text_frame.paragraphs[0], self._style_for(data, 'title'))

        body = slide.placeholders[1]
        tf = body.text_frame
        self._configure_text_frame(tf)
        tf.clear()

        bullets = data.get('bullets') or data.get('items') or []
        bullets = _as_list(bullets)
        if not bullets and data.get('content'):
            bullets = _as_list(data.get('content'))

        if not bullets:
            bullets = ['']

        for i, item in enumerate(bullets):
            if isinstance(item, dict):
                text = item.get('text') or item.get('title') or item.get('desc') or ''
                level = int(item.get('level') or 0)
            else:
                text = _as_text(item)
                level = 0
            p = tf.paragraphs[0] if i == 0 else tf.add_paragraph()
            p.text = text
            p.level = max(0, min(level, 4))
            self._apply_text_style(
                p,
                self._style_for(data, 'body', item),
                default_size=20 if p.level == 0 else 16,
                default_color=self._theme()['text']
            )

    # ============ 原生 PPT 矢量版式 ============

    def _theme(self):
        from pptx.dml.color import RGBColor
        theme = {
            'primary': RGBColor(37, 99, 235),
            'primary_dark': RGBColor(30, 64, 175),
            'accent': RGBColor(249, 115, 22),
            'bg': RGBColor(248, 250, 252),
            'card': RGBColor(255, 255, 255),
            'line': RGBColor(203, 213, 225),
            'text': RGBColor(15, 23, 42),
            'muted': RGBColor(71, 85, 105),
            'soft_blue': RGBColor(219, 234, 254),
            'soft_orange': RGBColor(255, 237, 213),
            'soft_green': RGBColor(220, 252, 231),
            'soft_purple': RGBColor(237, 233, 254),
            'green': RGBColor(22, 163, 74),
            'purple': RGBColor(124, 58, 237),
        }
        style = self._style_dict(getattr(self, '_ppt_global_style', {}))
        colors = style.get('colors') if isinstance(style.get('colors'), dict) else {}
        color_source = {}
        color_source.update(colors)
        color_source.update(style)
        profile_theme = self._style_dict(getattr(self, '_ppt_template_profile', {}).get('theme') if isinstance(getattr(self, '_ppt_template_profile', {}), dict) else {})
        color_source.update(profile_theme)
        aliases = {
            'primary': ('primary', 'primary_color'),
            'primary_dark': ('primary_dark', 'primary_dark_color'),
            'accent': ('accent', 'accent_color'),
            'bg': ('bg', 'bg_color', 'background_color'),
            'card': ('card', 'card_color'),
            'line': ('line', 'line_color'),
            'text': ('text', 'text_color'),
            'muted': ('muted', 'muted_color'),
            'soft_blue': ('soft_blue', 'soft_blue_color'),
            'soft_orange': ('soft_orange', 'soft_orange_color'),
            'soft_green': ('soft_green', 'soft_green_color'),
            'soft_purple': ('soft_purple', 'soft_purple_color'),
            'green': ('green', 'green_color'),
            'purple': ('purple', 'purple_color'),
        }
        for key, names in aliases.items():
            for name in names:
                if name in color_source:
                    theme[key] = self._parse_color(color_source.get(name), theme[key], theme)
                    break
        return theme

    def _style_dict(self, value):
        return dict(value) if isinstance(value, dict) else {}

    def _direct_style(self, value):
        if not isinstance(value, dict):
            return {}
        return {k: v for k, v in value.items() if k in _STYLE_KEYS}

    def _role_style(self, value, role):
        if not isinstance(value, dict) or not role:
            return {}
        aliases = [role]
        if role == 'body':
            aliases += ['text', 'desc', 'description', 'item']
        elif role == 'subtitle':
            aliases += ['sub_title', 'desc', 'description']
        elif role == 'quote':
            aliases += ['body', 'text']
        elif role == 'author':
            aliases += ['subtitle', 'caption']
        elif role == 'label':
            aliases += ['time', 'axis', 'caption']
        out = {}
        for alias in aliases:
            nested = value.get(f'{alias}_style')
            if isinstance(nested, dict):
                out.update(nested)
        attrs = ('font', 'font_family', 'font_name', 'font_size', 'size', 'color',
                 'font_color', 'bold', 'italic', 'underline', 'align', 'alignment')
        for alias in aliases:
            for attr in attrs:
                key = f'{alias}_{attr}'
                if key in value:
                    out[attr] = value[key]
        return out

    def _style_for(self, slide_data=None, role=None, item=None):
        out = {}
        profile = getattr(self, '_ppt_template_profile', {})
        if isinstance(profile, dict):
            out.update(self._template_profile_style(profile))
            font = self._style_dict(profile.get('font'))
            if role == 'title' and font.get('title_font_size'):
                out['font_size'] = font.get('title_font_size')
            elif role in ('body', 'subtitle') and font.get('body_font_size'):
                out['font_size'] = font.get('body_font_size')
        global_style = self._style_dict(getattr(self, '_ppt_global_style', {}))
        out.update(self._direct_style(global_style))
        out.update(self._role_style(global_style, role))
        if isinstance(slide_data, dict):
            slide_style = {}
            slide_style.update(self._style_dict(slide_data.get('theme')))
            slide_style.update(self._style_dict(slide_data.get('style')))
            out.update(self._direct_style(slide_style))
            out.update(self._role_style(slide_style, role))
            out.update(self._direct_style(slide_data))
            out.update(self._role_style(slide_data, role))
        if isinstance(item, dict):
            item_style = self._style_dict(item.get('style'))
            out.update(self._direct_style(item_style))
            out.update(self._role_style(item_style, role))
            out.update(self._direct_style(item))
            out.update(self._role_style(item, role))
        return out

    def _parse_color(self, value, default=None, theme=None):
        from pptx.dml.color import RGBColor
        if value is None or value == '':
            return default
        if isinstance(value, (list, tuple)) and len(value) >= 3:
            try:
                return RGBColor(*[max(0, min(255, int(x))) for x in value[:3]])
            except Exception:
                return default
        text = _as_text(value).strip()
        lower = text.lower()
        theme = theme if isinstance(theme, dict) else self._theme()
        if lower in theme:
            return theme[lower]
        named = {
            'white': RGBColor(255, 255, 255),
            'black': RGBColor(0, 0, 0),
            'red': RGBColor(220, 38, 38),
            'green': RGBColor(22, 163, 74),
            'blue': RGBColor(37, 99, 235),
            'orange': RGBColor(249, 115, 22),
            'purple': RGBColor(124, 58, 237),
            'gray': RGBColor(100, 116, 139),
            'grey': RGBColor(100, 116, 139),
        }
        if lower in named:
            return named[lower]
        m = re.fullmatch(r'#?([0-9a-fA-F]{6})', text)
        if m:
            raw = m.group(1)
            return RGBColor(int(raw[0:2], 16), int(raw[2:4], 16), int(raw[4:6], 16))
        m = re.fullmatch(r'rgba?\(\s*(\d{1,3})\s*,\s*(\d{1,3})\s*,\s*(\d{1,3})(?:\s*,\s*[\d.]+)?\s*\)', lower)
        if m:
            return RGBColor(*[max(0, min(255, int(x))) for x in m.groups()[:3]])
        return default

    def _style_number(self, style, keys, default=None, min_value=6, max_value=72):
        for key in keys:
            if key in style:
                try:
                    value = float(style.get(key))
                    return max(min_value, min(max_value, value))
                except Exception:
                    return default
        return default

    def _style_bool(self, style, keys, default=None):
        for key in keys:
            if key in style:
                value = style.get(key)
                if isinstance(value, bool):
                    return value
                if isinstance(value, (int, float)):
                    return bool(value)
                text = _as_text(value).strip().lower()
                if text in ('1', 'true', 'yes', 'y', 'on', 'bold'):
                    return True
                if text in ('0', 'false', 'no', 'n', 'off', 'normal'):
                    return False
        return default

    def _parse_align(self, value, default=None):
        if value is None:
            return default
        if not isinstance(value, str):
            return value
        from pptx.enum.text import PP_ALIGN
        mapping = {
            'left': PP_ALIGN.LEFT,
            'center': PP_ALIGN.CENTER,
            'centre': PP_ALIGN.CENTER,
            'right': PP_ALIGN.RIGHT,
            'justify': PP_ALIGN.JUSTIFY,
            'justified': PP_ALIGN.JUSTIFY,
        }
        return mapping.get(value.strip().lower(), default)

    def _apply_text_style(self, paragraph, style=None, default_size=None, default_bold=None, default_color=None, default_align=None):
        from pptx.util import Pt
        style = self._style_dict(style)
        size = self._style_number(style, ('font_size', 'size'), default_size)
        if size is not None:
            paragraph.font.size = Pt(size)
        family = style.get('font_family') or style.get('font_name') or style.get('font')
        if family:
            paragraph.font.name = _as_text(family)
        bold = self._style_bool(style, ('bold',), default_bold)
        if bold is not None:
            paragraph.font.bold = bold
        italic = self._style_bool(style, ('italic',), None)
        if italic is not None:
            paragraph.font.italic = italic
        underline = self._style_bool(style, ('underline',), None)
        if underline is not None:
            paragraph.font.underline = underline
        color = None
        for key in ('color', 'font_color', 'text_color'):
            if key in style:
                color = self._parse_color(style.get(key), default_color)
                break
        if color is None:
            color = default_color
        if color is not None:
            paragraph.font.color.rgb = color
        align = self._parse_align(style.get('align') or style.get('alignment'), default_align)
        if align:
            paragraph.alignment = align
        line_spacing = self._style_number(style, ('line_spacing',), None, 0.5, 3)
        if line_spacing is not None:
            paragraph.line_spacing = line_spacing

    def _blank_slide(self, prs, title, data=None):
        from pptx.util import Inches, Pt
        self._ppt_current_slide = data if isinstance(data, dict) else {}
        slide = prs.slides.add_slide(prs.slide_layouts[6])
        self._add_background(slide, prs)
        title_box = self._template_box('title_box')
        if title_box:
            box = slide.shapes.add_textbox(*title_box)
        else:
            box = slide.shapes.add_textbox(Inches(0.55), Inches(0.3), Inches(12.2), Inches(0.55))
        self._configure_text_frame(box.text_frame)
        p = box.text_frame.paragraphs[0]
        p.text = _as_text(title, '未命名页面')
        self._apply_text_style(
            p,
            self._style_for(data, 'title'),
            default_size=28,
            default_bold=True,
            default_color=self._theme()['text']
        )
        return slide

    def _add_background(self, slide, prs):
        from pptx.enum.shapes import MSO_SHAPE
        from pptx.util import Inches
        theme = self._theme()
        bg = slide.shapes.add_shape(MSO_SHAPE.RECTANGLE, 0, 0, prs.slide_width, prs.slide_height)
        bg.fill.solid()
        bg.fill.fore_color.rgb = theme['bg']
        bg.line.fill.background()
        if getattr(self, '_ppt_template_profile', None):
            self._apply_template_decorations(slide, prs)
            return
        bar = slide.shapes.add_shape(MSO_SHAPE.RECTANGLE, 0, 0, Inches(0.11), prs.slide_height)
        bar.fill.solid()
        bar.fill.fore_color.rgb = theme['primary']
        bar.line.fill.background()

    def _set_shape_fill(self, shape, color, line_color=None):
        shape.fill.solid()
        shape.fill.fore_color.rgb = color
        if line_color is None:
            shape.line.fill.background()
        else:
            shape.line.color.rgb = line_color

    def _configure_text_frame(self, text_frame, *, margin_left=2, margin_right=2, margin_top=1, margin_bottom=1, auto_fit=True, normalize_margins=True):
        try:
            from pptx.enum.text import MSO_AUTO_SIZE
            from pptx.util import Pt
            text_frame.word_wrap = True
            if normalize_margins:
                text_frame.margin_left = Pt(margin_left)
                text_frame.margin_right = Pt(margin_right)
                text_frame.margin_top = Pt(margin_top)
                text_frame.margin_bottom = Pt(margin_bottom)
            if auto_fit:
                text_frame.auto_size = MSO_AUTO_SIZE.TEXT_TO_FIT_SHAPE
        except Exception:
            try:
                text_frame.word_wrap = True
            except Exception:
                pass

    def _add_textbox(self, slide, text, left, top, width, height, size=16, bold=False, color=None, align=None, style=None, role=None):
        theme = self._theme()
        box = slide.shapes.add_textbox(left, top, width, height)
        tf = box.text_frame
        self._configure_text_frame(tf)
        p = tf.paragraphs[0]
        p.text = _as_text(text)
        if style is None:
            style = self._style_for(getattr(self, '_ppt_current_slide', {}), role)
        self._apply_text_style(
            p,
            style,
            default_size=size,
            default_bold=bold,
            default_color=color or theme['text'],
            default_align=align
        )
        return box

    def _auto_visualize_slides(self, slides):
        """把模型偷懒生成的纯 bullets，按标题关键词自动提升成矢量版式。

        这不是替代模型设计，而是兜底：用户明明要卡片/流程/时间轴/对比时，
        模型经常仍只给 bullets。这里在不破坏 cover/agenda/summary 的前提下，
        将常见页面转换为更像 PPT 的原生形状页面。学术页会优先提升为架构、
        研究流程、表格、图表、实验设计、消融实验等版式。
        """
        out = []
        for idx, slide in enumerate(_as_list(slides)):
            if not isinstance(slide, dict):
                out.append(slide)
                continue
            s = dict(slide)
            slide_type = _as_text(s.get('type') or s.get('layout'), 'bullets').lower()
            if slide_type not in ('bullets', 'content', 'text'):
                out.append(s)
                continue

            title = _as_text(s.get('title')).lower()
            bullets = _as_list(s.get('bullets') or s.get('items') or [])
            if len(bullets) < 2:
                out.append(s)
                continue

            if self._auto_visualize_academic_slide(s, title, bullets):
                out.append(s)
                continue

            if any(k in title for k in ('对比', '比较', '区别', '差异', ' vs ', 'vs')) and len(bullets) >= 2:
                mid = max(1, len(bullets) // 2)
                s['type'] = 'comparison'
                s.setdefault('left', {'title': '传统方式', 'items': bullets[:mid]})
                s.setdefault('right', {'title': 'Agent 方式', 'items': bullets[mid:] or bullets[:mid]})
            elif any(k in title for k in ('流程', '步骤', '链路', '过程', '工作流', '方法')):
                s['type'] = 'process'
                s['steps'] = [self._bullet_to_item(x, i + 1) for i, x in enumerate(bullets[:5])]
            elif any(k in title for k in ('路线', '时间轴', '阶段', '计划', '排期', '里程碑', '演进')):
                s['type'] = 'timeline'
                s['events'] = [self._bullet_to_item(x, i + 1, with_time=True) for i, x in enumerate(bullets[:5])]
            elif any(k in title for k in ('矩阵', '四象限', '象限')):
                s['type'] = 'matrix'
                s['quadrants'] = [self._bullet_to_item(x, i + 1) for i, x in enumerate(bullets[:4])]
            elif any(k in title for k in ('金字塔', '层级', '分层')):
                s['type'] = 'pyramid'
                s['levels'] = bullets[:5]
            elif any(k in title for k in ('闭环', '循环', '飞轮')):
                s['type'] = 'cycle'
                s['steps'] = [self._bullet_to_item(x, i + 1) for i, x in enumerate(bullets[:5])]
            elif any(k in title for k in ('能力', '优势', '价值', '场景', '模块', '特点', '特性')) or (idx > 0 and 3 <= len(bullets) <= 4):
                s['type'] = 'three_cards'
                s['cards'] = [self._bullet_to_item(x, i + 1) for i, x in enumerate(bullets[:4])]
            out.append(s)
        return out

    def _auto_visualize_academic_slide(self, slide, title, bullets):
        text = f'{title} ' + ' '.join(_item_text(x).lower() for x in bullets[:8])
        numeric_points = self._numeric_points_from_bullets(bullets)
        table_data = self._table_from_bullets(bullets, title)

        if self._has_any(text, ('消融', 'ablation', '移除模块', '去除模块', 'w/o', 'without')):
            slide['type'] = 'ablation'
            headers, rows = self._ablation_from_bullets(bullets)
            slide.setdefault('headers', headers)
            slide.setdefault('rows', rows)
            return True

        if self._has_any(title, ('实验设计', '实验设置', '实验方案', '对照实验', '评估方案', '研究设计')):
            slide['type'] = 'experiment_design'
            slide.setdefault('groups', [self._bullet_to_experiment_group(x, i + 1) for i, x in enumerate(bullets[:4])])
            return True

        if self._has_any(title, ('系统架构', '项目架构', '技术架构', '模型架构', '整体架构', '网络结构', '模型结构', '框架设计', '模块组成')):
            slide['type'] = 'architecture'
            slide.setdefault('layers', [self._bullet_to_layer(x, i + 1) for i, x in enumerate(bullets[:5])])
            return True

        if self._has_any(title, ('研究流程', '研究方法', '方法流程', '技术路线', '实验流程', '训练流程', '推理流程', '处理流程', 'pipeline', 'method pipeline')):
            slide['type'] = 'method_pipeline'
            slide.setdefault('stages', [self._bullet_to_stage(x, i + 1) for i, x in enumerate(bullets[:5])])
            return True

        if numeric_points and self._has_any(title, ('图表', '曲线', '趋势', '实验结果', '性能', '指标', '准确率', '精度', '召回率', 'f1', 'auc', 'loss', '结果对比', '性能对比')):
            slide['type'] = 'chart'
            slide.setdefault('categories', [p['category'] for p in numeric_points[:8]])
            slide.setdefault('values', [p['value'] for p in numeric_points[:8]])
            if self._has_any(title, ('曲线', '趋势', '变化', 'epoch', 'loss', '收敛')):
                slide.setdefault('chart_type', 'line')
            else:
                slide.setdefault('chart_type', 'bar')
            return True

        if table_data and self._has_any(title, ('表格', '数据表', '结果表', '指标表', '对比表', '实验结果', '性能对比', '结果对比')):
            headers, rows = table_data
            slide['type'] = 'table'
            slide.setdefault('headers', headers)
            slide.setdefault('rows', rows)
            return True

        return False

    def _has_any(self, text, keywords):
        text = _as_text(text).lower()
        return any(k in text for k in keywords)

    def _split_bullet_detail(self, item, idx=1):
        if isinstance(item, dict):
            title = item.get('title') or item.get('name') or item.get('text') or f'项目 {idx}'
            detail = item.get('desc') or item.get('description') or item.get('content') or ''
            return _as_text(title).strip(), _as_text(detail).strip()
        text = _as_text(item).strip()
        parts = re.split(r'[：:；;]\s*|\s+-\s+|\s+—\s+', text, maxsplit=1)
        title = parts[0].strip() or f'项目 {idx}'
        detail = parts[1].strip() if len(parts) > 1 else ''
        return title, detail

    def _split_detail_items(self, text):
        items = [x.strip() for x in re.split(r'[、，,；;|｜/]+', _as_text(text)) if x.strip()]
        return items

    def _bullet_to_layer(self, item, idx=1):
        if isinstance(item, dict):
            layer = dict(item)
            layer.setdefault('title', layer.get('name') or layer.get('text') or f'架构层 {idx}')
            layer.setdefault('items', _as_list(layer.get('items') or layer.get('modules') or layer.get('components')))
            return layer
        title, detail = self._split_bullet_detail(item, idx)
        parts = self._split_detail_items(detail)
        return {'title': title, 'items': parts or [_as_text(detail) or '模块说明']}

    def _bullet_to_stage(self, item, idx=1):
        if isinstance(item, dict):
            stage = dict(item)
            stage.setdefault('title', stage.get('name') or stage.get('text') or f'阶段 {idx}')
            stage.setdefault('items', _as_list(stage.get('items') or stage.get('substeps') or stage.get('tasks')))
            return stage
        title, detail = self._split_bullet_detail(item, idx)
        parts = self._split_detail_items(detail)
        stage = {'title': title}
        if detail and not parts:
            stage['desc'] = detail
        if parts:
            stage['items'] = parts[:4]
        return stage

    def _bullet_to_experiment_group(self, item, idx=1):
        if isinstance(item, dict):
            group = dict(item)
            group.setdefault('name', group.get('title') or group.get('text') or f'实验组 {idx}')
            group.setdefault('method', group.get('desc') or group.get('description') or '')
            group.setdefault('metrics', _as_list(group.get('metrics') or group.get('indicators')))
            return group
        name, detail = self._split_bullet_detail(item, idx)
        parts = self._split_detail_items(detail)
        return {
            'name': name,
            'method': parts[0] if parts else (detail or '实验方法'),
            'metrics': self._metrics_from_text(detail) or parts[1:4] or ['Accuracy', 'F1']
        }

    def _metrics_from_text(self, text):
        found = []
        metric_names = ('accuracy', 'acc', 'precision', 'recall', 'f1', 'auc', 'map', 'mae', 'rmse', 'loss', '准确率', '精度', '召回率')
        lower = _as_text(text).lower()
        for name in metric_names:
            if re.fullmatch(r'[a-z0-9]+', name):
                matched = re.search(rf'(?<![a-z0-9]){re.escape(name)}(?![a-z0-9])', lower)
            else:
                matched = name in lower
            if matched and name not in found:
                found.append(name.upper() if re.fullmatch(r'[a-z0-9]+', name) else name)
        return found[:5]

    def _numeric_points_from_bullets(self, bullets):
        points = []
        for idx, item in enumerate(_as_list(bullets)):
            if isinstance(item, dict):
                raw_value = item.get('value')
                label = item.get('title') or item.get('name') or item.get('text') or f'项目 {idx + 1}'
                if raw_value is not None:
                    try:
                        points.append({'category': _as_text(label), 'value': float(raw_value)})
                        continue
                    except Exception:
                        pass
                text = _item_text(item)
            else:
                text = _as_text(item)
            nums = re.findall(r'[-+]?\d+(?:\.\d+)?\s*%?', text)
            if not nums:
                continue
            number = nums[-1].replace('%', '').strip()
            try:
                value = float(number)
            except Exception:
                continue
            title, detail = self._split_bullet_detail(text, idx + 1)
            if detail:
                category = title
            else:
                category = re.sub(r'[-+]?\d+(?:\.\d+)?\s*%?.*$', '', title).strip(' ：:，,；;-—') or title
            category = category or f'项目 {idx + 1}'
            points.append({'category': category[:24], 'value': value})
        return points if len(points) >= 2 else []

    def _table_from_bullets(self, bullets, title=''):
        rows = []
        max_cols = 0
        for idx, item in enumerate(_as_list(bullets)[:8]):
            if isinstance(item, dict):
                keys = [k for k in item.keys() if k not in ('type', 'layout')]
                if len(keys) >= 2:
                    row = [_as_text(item.get(k, '')) for k in keys[:6]]
                else:
                    row = [_item_text(item)]
            else:
                label, detail = self._split_bullet_detail(item, idx + 1)
                parts = self._split_detail_items(detail)
                if not parts and '|' in _as_text(item):
                    parts = [x.strip() for x in _as_text(item).split('|') if x.strip()]
                    row = parts[:6]
                else:
                    row = ([label] + parts)[:6] if parts else [label]
            if len(row) >= 2:
                rows.append(row)
                max_cols = max(max_cols, len(row))
        if len(rows) < 2 or max_cols < 2:
            return None
        headers = self._default_table_headers(max_cols, title)
        padded = [row + [''] * (max_cols - len(row)) for row in rows]
        return headers, padded

    def _default_table_headers(self, count, title=''):
        if self._has_any(title, ('性能', '指标', '结果')):
            base = ['对象', '指标1', '指标2', '结论', '备注', '说明']
        elif self._has_any(title, ('对比', '比较')):
            base = ['对象', '对比项1', '对比项2', '结论', '备注', '说明']
        else:
            base = ['项目', '内容1', '内容2', '结论', '备注', '说明']
        return base[:count]

    def _ablation_from_bullets(self, bullets):
        rows = []
        for idx, item in enumerate(_as_list(bullets)[:8]):
            if isinstance(item, dict):
                name = item.get('name') or item.get('title') or item.get('setting') or f'实验设置 {idx + 1}'
                removed = item.get('removed') or item.get('module') or item.get('change') or '-'
                metric = item.get('metric') or item.get('score') or item.get('value') or ''
                conclusion = item.get('conclusion') or item.get('desc') or item.get('description') or ''
                rows.append([_as_text(name), _as_text(removed), _as_text(metric), _as_text(conclusion)])
                continue
            name, detail = self._split_bullet_detail(item, idx + 1)
            parts = self._split_detail_items(detail)
            metric = ''
            for part in parts:
                if re.search(r'[-+]?\d+(?:\.\d+)?\s*%?', part):
                    metric = part
                    break
            removed = '-'
            lower_name = name.lower()
            if lower_name.startswith(('w/o', 'without')):
                removed = re.sub(r'^(w/o|without)\s+', '', name, flags=re.I).strip() or '-'
            conclusion = '；'.join([p for p in parts if p != metric]) or detail
            rows.append([name, removed, metric, conclusion])
        return ['实验设置', '去除模块', '关键指标', '结论'], rows

    def _bullet_to_item(self, item, idx=1, with_time=False):
        if isinstance(item, dict):
            return dict(item)
        text = _as_text(item)
        parts = re.split(r'[：:，,。；;\-—]', text, maxsplit=1)
        result = {'title': parts[0].strip() or f'项目 {idx}'}
        if len(parts) > 1 and parts[1].strip():
            result['desc'] = parts[1].strip()
        if with_time:
            result['time'] = f'阶段 {idx}'
        return result

    def _item_title_desc(self, item, idx=None):
        if isinstance(item, dict):
            title = item.get('title') or item.get('text') or item.get('name') or (f'步骤 {idx}' if idx else '')
            desc = item.get('desc') or item.get('description') or item.get('content') or ''
            icon = item.get('icon') or ''
            time = item.get('time') or item.get('date') or item.get('phase') or ''
        else:
            title, desc, icon, time = _as_text(item), '', '', ''
        return _as_text(title), _as_text(desc), _as_text(icon), _as_text(time)

    def _ppt_add_three_cards(self, prs, data):
        from pptx.enum.shapes import MSO_SHAPE
        from pptx.enum.text import PP_ALIGN
        from pptx.util import Inches
        theme = self._theme()
        slide = self._blank_slide(prs, data.get('title'), data)
        cards = _as_list(data.get('cards') or data.get('items') or data.get('bullets'))[:4]
        if not cards:
            cards = [{'title': '能力', 'desc': '在此补充说明'}]
        n = len(cards)
        gap = Inches(0.28)
        left0 = Inches(0.85)
        top = Inches(1.55)
        total_w = Inches(11.6)
        card_w = int((total_w - gap * (n - 1)) / n)
        card_h = Inches(4.55)
        colors = [theme['soft_blue'], theme['soft_orange'], theme['card'], theme['soft_blue']]
        icons = ['①', '②', '③', '④']
        for i, item in enumerate(cards):
            title, desc, icon, _ = self._item_title_desc(item, i + 1)
            icon = icon or icons[i]
            left = left0 + i * (card_w + gap)
            card = slide.shapes.add_shape(MSO_SHAPE.ROUNDED_RECTANGLE, left, top, card_w, card_h)
            self._set_shape_fill(card, colors[i % len(colors)], theme['line'])
            self._add_textbox(slide, icon, left + Inches(0.2), top + Inches(0.28), card_w - Inches(0.4), Inches(0.62), 30, True, theme['primary'], PP_ALIGN.CENTER, self._style_for(data, 'icon', item), 'icon')
            self._add_textbox(slide, title, left + Inches(0.22), top + Inches(1.15), card_w - Inches(0.44), Inches(0.55), 20, True, theme['text'], PP_ALIGN.CENTER, self._style_for(data, 'title', item), 'title')
            self._add_textbox(slide, desc, left + Inches(0.28), top + Inches(2.05), card_w - Inches(0.56), Inches(1.65), 14, False, theme['muted'], PP_ALIGN.CENTER, self._style_for(data, 'body', item), 'body')

    def _ppt_add_process(self, prs, data):
        from pptx.enum.shapes import MSO_SHAPE
        from pptx.enum.text import PP_ALIGN
        from pptx.util import Inches
        theme = self._theme()
        slide = self._blank_slide(prs, data.get('title'), data)
        steps = _as_list(data.get('steps') or data.get('items') or data.get('bullets'))[:5]
        if not steps:
            steps = ['理解目标', '规划步骤', '调用工具', '检查结果']
        n = len(steps)
        left0 = Inches(0.85)
        top = Inches(2.0)
        area_w = Inches(11.45)
        node = Inches(1.18)
        step_gap = int((area_w - node * n) / max(1, n - 1)) if n > 1 else 0
        for i, item in enumerate(steps):
            title, desc, _, _ = self._item_title_desc(item, i + 1)
            x = left0 + i * (node + step_gap)
            if i > 0:
                x1 = left0 + (i - 1) * (node + step_gap) + node
                y = top + node / 2
                arrow = slide.shapes.add_shape(MSO_SHAPE.RIGHT_ARROW, x1 + Inches(0.14), y - Inches(0.08), max(Inches(0.2), x - x1 - Inches(0.28)), Inches(0.16))
                arrow.fill.solid(); arrow.fill.fore_color.rgb = theme['line']
                arrow.line.fill.background()
            circle = slide.shapes.add_shape(MSO_SHAPE.OVAL, x, top, node, node)
            self._set_shape_fill(circle, theme['primary'], theme['primary_dark'])
            self._add_textbox(slide, str(i + 1), x, top + Inches(0.28), node, Inches(0.42), 22, True, theme['card'], PP_ALIGN.CENTER, self._style_for(data, 'label', item), 'label')
            self._add_textbox(slide, title, x - Inches(0.28), top + Inches(1.45), node + Inches(0.56), Inches(0.48), 17, True, theme['text'], PP_ALIGN.CENTER, self._style_for(data, 'title', item), 'title')
            self._add_textbox(slide, desc, x - Inches(0.42), top + Inches(2.02), node + Inches(0.84), Inches(0.9), 12, False, theme['muted'], PP_ALIGN.CENTER, self._style_for(data, 'body', item), 'body')

    def _ppt_add_timeline(self, prs, data):
        from pptx.enum.shapes import MSO_SHAPE
        from pptx.enum.text import PP_ALIGN
        from pptx.util import Inches
        theme = self._theme()
        slide = self._blank_slide(prs, data.get('title'), data)
        events = _as_list(data.get('events') or data.get('steps') or data.get('items'))[:5]
        if not events:
            events = [{'time': '阶段一', 'title': '启动'}, {'time': '阶段二', 'title': '落地'}, {'time': '阶段三', 'title': '优化'}]
        n = len(events)
        left0 = Inches(1.05)
        right = Inches(12.1)
        y = Inches(3.05)
        line = slide.shapes.add_shape(MSO_SHAPE.RECTANGLE, left0, y - Inches(0.015), right - left0, Inches(0.03))
        line.fill.solid(); line.fill.fore_color.rgb = theme['line']
        line.line.fill.background()
        span = int((right - left0) / max(1, n - 1)) if n > 1 else 0
        for i, item in enumerate(events):
            title, desc, _, time = self._item_title_desc(item, i + 1)
            x = left0 + i * span
            dot = slide.shapes.add_shape(MSO_SHAPE.OVAL, x - Inches(0.16), y - Inches(0.16), Inches(0.32), Inches(0.32))
            self._set_shape_fill(dot, theme['accent'] if i % 2 else theme['primary'], theme['card'])
            top = Inches(1.45) if i % 2 == 0 else Inches(3.45)
            self._add_textbox(slide, time or f'阶段 {i + 1}', x - Inches(0.75), top, Inches(1.5), Inches(0.35), 13, True, theme['primary'], PP_ALIGN.CENTER, self._style_for(data, 'label', item), 'label')
            self._add_textbox(slide, title, x - Inches(0.95), top + Inches(0.42), Inches(1.9), Inches(0.48), 15, True, theme['text'], PP_ALIGN.CENTER, self._style_for(data, 'title', item), 'title')
            self._add_textbox(slide, desc, x - Inches(1.05), top + Inches(0.92), Inches(2.1), Inches(0.68), 11, False, theme['muted'], PP_ALIGN.CENTER, self._style_for(data, 'body', item), 'body')

    def _ppt_add_comparison(self, prs, data):
        from pptx.enum.shapes import MSO_SHAPE
        from pptx.enum.text import PP_ALIGN
        from pptx.util import Inches, Pt
        theme = self._theme()
        slide = self._blank_slide(prs, data.get('title'), data)
        left_data = data.get('left') if isinstance(data.get('left'), dict) else {'title': '方案 A', 'items': data.get('left') or []}
        right_data = data.get('right') if isinstance(data.get('right'), dict) else {'title': '方案 B', 'items': data.get('right') or []}
        panels = [(left_data, Inches(0.95), theme['soft_blue'], theme['primary']), (right_data, Inches(7.0), theme['soft_orange'], theme['accent'])]
        for panel, left, fill, color in panels:
            box = slide.shapes.add_shape(MSO_SHAPE.ROUNDED_RECTANGLE, left, Inches(1.45), Inches(5.25), Inches(4.75))
            self._set_shape_fill(box, fill, theme['line'])
            self._add_textbox(slide, panel.get('title') or '对比项', left + Inches(0.28), Inches(1.75), Inches(4.69), Inches(0.55), 22, True, color, PP_ALIGN.CENTER, self._style_for(data, 'title', panel), 'title')
            items = _as_list(panel.get('items') or panel.get('bullets') or panel.get('points'))[:6]
            tf_box = slide.shapes.add_textbox(left + Inches(0.52), Inches(2.65), Inches(4.25), Inches(2.9))
            tf = tf_box.text_frame
            self._configure_text_frame(tf)
            tf.clear()
            for i, item in enumerate(items or ['']):
                p = tf.paragraphs[0] if i == 0 else tf.add_paragraph()
                p.text = _as_text(item.get('text') if isinstance(item, dict) else item)
                self._apply_text_style(p, self._style_for(data, 'body', item), default_size=15, default_color=theme['text'])
                p.level = 0
        vs = slide.shapes.add_shape(MSO_SHAPE.OVAL, Inches(6.05), Inches(3.08), Inches(0.82), Inches(0.82))
        self._set_shape_fill(vs, theme['primary'], theme['card'])
        self._add_textbox(slide, 'VS', Inches(6.05), Inches(3.29), Inches(0.82), Inches(0.3), 18, True, theme['card'], PP_ALIGN.CENTER, self._style_for(data, 'label'), 'label')

    def _ppt_add_two_column(self, prs, data):
        from pptx.enum.shapes import MSO_SHAPE
        from pptx.enum.text import PP_ALIGN
        from pptx.util import Inches, Pt
        theme = self._theme()
        slide = self._blank_slide(prs, data.get('title'), data)

        left_data = data.get('left') if isinstance(data.get('left'), dict) else {'title': data.get('left_title') or '左侧', 'items': data.get('left') or []}
        right_data = data.get('right') if isinstance(data.get('right'), dict) else {'title': data.get('right_title') or '右侧', 'items': data.get('right') or []}
        panels = [
            (left_data, Inches(0.9), theme['soft_blue'], theme['primary']),
            (right_data, Inches(6.85), theme['soft_green'], theme['green']),
        ]
        for panel, left, fill, color in panels:
            box = slide.shapes.add_shape(MSO_SHAPE.ROUNDED_RECTANGLE, left, Inches(1.45), Inches(5.45), Inches(4.85))
            self._set_shape_fill(box, fill, theme['line'])
            self._add_textbox(slide, panel.get('title') or '栏目', left + Inches(0.35), Inches(1.78), Inches(4.75), Inches(0.5), 22, True, color, PP_ALIGN.CENTER, self._style_for(data, 'title', panel), 'title')
            items = _as_list(panel.get('items') or panel.get('bullets') or panel.get('points') or panel.get('content'))[:7]
            tf_box = slide.shapes.add_textbox(left + Inches(0.55), Inches(2.65), Inches(4.5), Inches(3.15))
            tf = tf_box.text_frame
            self._configure_text_frame(tf)
            tf.clear()
            for i, item in enumerate(items or ['']) :
                p = tf.paragraphs[0] if i == 0 else tf.add_paragraph()
                p.text = '• ' + _as_text(item.get('text') if isinstance(item, dict) else item)
                self._apply_text_style(p, self._style_for(data, 'body', item), default_size=15, default_color=theme['text'])

    def _ppt_add_matrix(self, prs, data):
        from pptx.enum.shapes import MSO_SHAPE
        from pptx.enum.text import PP_ALIGN
        from pptx.util import Inches
        theme = self._theme()
        slide = self._blank_slide(prs, data.get('title'), data)
        quadrants = _as_list(data.get('quadrants') or data.get('items') or data.get('cards') or data.get('bullets'))[:4]
        if not quadrants:
            quadrants = ['高价值高频', '高价值低频', '低价值高频', '低价值低频']
        while len(quadrants) < 4:
            quadrants.append({'title': f'象限 {len(quadrants) + 1}', 'desc': ''})

        left0, top0 = Inches(1.45), Inches(1.55)
        cell_w, cell_h = Inches(4.95), Inches(2.15)
        gap = Inches(0.12)
        colors = [theme['soft_blue'], theme['soft_orange'], theme['soft_green'], theme['soft_purple']]
        accent = [theme['primary'], theme['accent'], theme['green'], theme['purple']]
        for i, item in enumerate(quadrants[:4]):
            row, col = divmod(i, 2)
            left = left0 + col * (cell_w + gap)
            top = top0 + row * (cell_h + gap)
            rect = slide.shapes.add_shape(MSO_SHAPE.ROUNDED_RECTANGLE, left, top, cell_w, cell_h)
            self._set_shape_fill(rect, colors[i], theme['line'])
            title, desc, _, _ = self._item_title_desc(item, i + 1)
            self._add_textbox(slide, title, left + Inches(0.28), top + Inches(0.32), cell_w - Inches(0.56), Inches(0.45), 19, True, accent[i], PP_ALIGN.CENTER, self._style_for(data, 'title', item), 'title')
            self._add_textbox(slide, desc, left + Inches(0.4), top + Inches(0.95), cell_w - Inches(0.8), Inches(0.78), 13, False, theme['muted'], PP_ALIGN.CENTER, self._style_for(data, 'body', item), 'body')
        self._add_textbox(slide, data.get('x_label') or '频率 / 可行性', Inches(4.7), Inches(6.08), Inches(3.2), Inches(0.3), 12, False, theme['muted'], PP_ALIGN.CENTER, self._style_for(data, 'label'), 'label')
        self._add_textbox(slide, data.get('y_label') or '价值 / 影响', Inches(0.35), Inches(3.05), Inches(0.8), Inches(0.4), 12, False, theme['muted'], PP_ALIGN.CENTER, self._style_for(data, 'label'), 'label')

    def _ppt_add_pyramid(self, prs, data):
        from pptx.enum.shapes import MSO_SHAPE
        from pptx.enum.text import PP_ALIGN
        from pptx.util import Inches
        theme = self._theme()
        slide = self._blank_slide(prs, data.get('title'), data)
        levels = _as_list(data.get('levels') or data.get('items') or data.get('bullets'))[:5]
        if not levels:
            levels = ['基础能力', '流程编排', '智能决策']
        n = len(levels)
        max_w = Inches(8.8)
        level_h = Inches(0.76)
        gap = Inches(0.08)
        center_x = Inches(6.67)
        bottom = Inches(6.05)
        colors = [theme['primary'], theme['purple'], theme['green'], theme['accent'], theme['primary_dark']]
        for visual_idx, item in enumerate(reversed(levels)):
            # bottom level is widest; upper levels gradually shrink.
            width = max(Inches(2.2), int(max_w * (n - visual_idx) / n))
            left = center_x - width / 2
            top = bottom - (visual_idx + 1) * level_h - visual_idx * gap
            shape = slide.shapes.add_shape(MSO_SHAPE.TRAPEZOID, left, top, width, level_h)
            color = colors[visual_idx % len(colors)]
            self._set_shape_fill(shape, color, theme['card'])
            title, desc, _, _ = self._item_title_desc(item, n - visual_idx)
            text = title if not desc else f'{title}｜{desc}'
            self._add_textbox(slide, text, left + Inches(0.18), top + Inches(0.18), width - Inches(0.36), Inches(0.34), 15, True, theme['card'], PP_ALIGN.CENTER, self._style_for(data, 'title', item), 'title')

    def _ppt_add_cycle(self, prs, data):
        from math import cos, sin, pi
        from pptx.enum.shapes import MSO_SHAPE
        from pptx.enum.text import PP_ALIGN
        from pptx.util import Inches
        theme = self._theme()
        slide = self._blank_slide(prs, data.get('title'), data)
        steps = _as_list(data.get('steps') or data.get('items') or data.get('bullets'))[:6]
        if not steps:
            steps = ['理解目标', '规划任务', '执行操作', '反馈优化']
        n = len(steps)
        cx, cy = Inches(6.65), Inches(3.75)
        r = Inches(2.0)
        node_w, node_h = Inches(1.55), Inches(0.82)
        colors = [theme['primary'], theme['accent'], theme['green'], theme['purple'], theme['primary_dark'], theme['muted']]
        center = slide.shapes.add_shape(MSO_SHAPE.OVAL, cx - Inches(0.72), cy - Inches(0.72), Inches(1.44), Inches(1.44))
        self._set_shape_fill(center, theme['card'], theme['line'])
        self._add_textbox(slide, data.get('center') or '闭环', cx - Inches(0.62), cy - Inches(0.18), Inches(1.24), Inches(0.36), 17, True, theme['primary'], PP_ALIGN.CENTER, self._style_for(data, 'label'), 'label')
        for i, item in enumerate(steps):
            angle = -pi / 2 + 2 * pi * i / n
            x = cx + int(r * cos(angle)) - node_w / 2
            y = cy + int(r * sin(angle)) - node_h / 2
            title, desc, _, _ = self._item_title_desc(item, i + 1)
            node = slide.shapes.add_shape(MSO_SHAPE.ROUNDED_RECTANGLE, x, y, node_w, node_h)
            self._set_shape_fill(node, colors[i % len(colors)], theme['card'])
            self._add_textbox(slide, title, x + Inches(0.08), y + Inches(0.2), node_w - Inches(0.16), Inches(0.28), 13, True, theme['card'], PP_ALIGN.CENTER, self._style_for(data, 'title', item), 'title')
            # Tangential arrow between nodes; use non-zero arrow shape to avoid corrupt connector XML.
            next_angle = -pi / 2 + 2 * pi * ((i + 0.5) % n) / n
            ax = cx + int((r + Inches(0.02)) * cos(next_angle)) - Inches(0.28)
            ay = cy + int((r + Inches(0.02)) * sin(next_angle)) - Inches(0.12)
            arrow = slide.shapes.add_shape(MSO_SHAPE.RIGHT_ARROW, ax, ay, Inches(0.56), Inches(0.24))
            self._set_shape_fill(arrow, theme['line'])
            arrow.rotation = int((next_angle + pi / 2) * 180 / pi)

    def _ppt_add_funnel(self, prs, data):
        from pptx.enum.shapes import MSO_SHAPE
        from pptx.enum.text import PP_ALIGN
        from pptx.util import Inches
        theme = self._theme()
        slide = self._blank_slide(prs, data.get('title'), data)
        stages = _as_list(data.get('stages') or data.get('steps') or data.get('items') or data.get('bullets'))[:5]
        if not stages:
            stages = ['识别场景', '评估价值', '接入工具', '形成闭环']
        n = len(stages)
        center_x = Inches(6.65)
        top0 = Inches(1.45)
        h = Inches(0.82)
        gap = Inches(0.12)
        max_w = Inches(8.6)
        min_w = Inches(3.25)
        colors = [theme['primary'], theme['purple'], theme['accent'], theme['green'], theme['primary_dark']]
        for i, item in enumerate(stages):
            width = max_w - int((max_w - min_w) * i / max(1, n - 1))
            left = center_x - width / 2
            top = top0 + i * (h + gap)
            shape = slide.shapes.add_shape(MSO_SHAPE.TRAPEZOID, left, top, width, h)
            shape.rotation = 180
            self._set_shape_fill(shape, colors[i % len(colors)], theme['card'])
            title, desc, _, _ = self._item_title_desc(item, i + 1)
            text = title if not desc else f'{title}：{desc}'
            self._add_textbox(slide, text, left + Inches(0.25), top + Inches(0.2), width - Inches(0.5), Inches(0.34), 15, True, theme['card'], PP_ALIGN.CENTER, self._style_for(data, 'title', item), 'title')

    def _ppt_add_quote(self, prs, data):
        from pptx.enum.shapes import MSO_SHAPE
        from pptx.enum.text import PP_ALIGN
        from pptx.util import Inches
        theme = self._theme()
        slide = self._blank_slide(prs, data.get('title') or '核心观点', data)
        quote = data.get('quote') or data.get('content') or data.get('text') or '在这里填写核心观点。'
        author = data.get('author') or data.get('source') or ''
        mark = slide.shapes.add_shape(MSO_SHAPE.OVAL, Inches(1.05), Inches(1.45), Inches(1.0), Inches(1.0))
        self._set_shape_fill(mark, theme['soft_blue'], theme['primary'])
        self._add_textbox(slide, '“', Inches(1.07), Inches(1.43), Inches(0.96), Inches(0.68), 42, True, theme['primary'], PP_ALIGN.CENTER, self._style_for(data, 'icon'), 'icon')
        card = slide.shapes.add_shape(MSO_SHAPE.ROUNDED_RECTANGLE, Inches(1.65), Inches(2.0), Inches(10.35), Inches(3.45))
        self._set_shape_fill(card, theme['card'], theme['line'])
        self._add_textbox(slide, quote, Inches(2.2), Inches(2.45), Inches(9.25), Inches(1.65), 26, True, theme['text'], PP_ALIGN.CENTER, self._style_for(data, 'quote'), 'quote')
        if author:
            self._add_textbox(slide, f'—— {author}', Inches(6.65), Inches(4.5), Inches(4.6), Inches(0.42), 15, False, theme['muted'], PP_ALIGN.RIGHT, self._style_for(data, 'author'), 'author')

    def _ppt_add_architecture(self, prs, data):
        from pptx.enum.shapes import MSO_SHAPE
        from pptx.enum.text import PP_ALIGN
        from pptx.util import Inches, Pt
        theme = self._theme()
        slide = self._blank_slide(prs, data.get('title') or '项目架构', data)
        layers = _as_list(data.get('layers') or data.get('modules') or data.get('items'))[:5]
        if not layers:
            layers = [
                {'title': '数据层', 'items': ['数据采集', '数据标注']},
                {'title': '处理层', 'items': ['清洗去噪', '特征构建']},
                {'title': '模型层', 'items': ['模型训练', '参数优化']},
                {'title': '应用层', 'items': ['结果分析', '系统展示']},
            ]
        n = len(layers)
        left0, top, area_w = Inches(0.75), Inches(1.45), Inches(11.85)
        gap = Inches(0.22)
        box_w = int((area_w - gap * (n - 1)) / n)
        box_h = Inches(4.85)
        colors = [theme['soft_blue'], theme['soft_green'], theme['soft_orange'], theme['soft_purple'], theme['card']]
        accents = [theme['primary'], theme['green'], theme['accent'], theme['purple'], theme['primary_dark']]
        for i, layer in enumerate(layers):
            title, desc, _, _ = self._item_title_desc(layer, i + 1)
            items = _as_list(layer.get('items') or layer.get('modules') or layer.get('components') if isinstance(layer, dict) else [])[:5]
            left = left0 + i * (box_w + gap)
            rect = slide.shapes.add_shape(MSO_SHAPE.ROUNDED_RECTANGLE, left, top, box_w, box_h)
            self._set_shape_fill(rect, colors[i % len(colors)], theme['line'])
            self._add_textbox(slide, title, left + Inches(0.12), top + Inches(0.25), box_w - Inches(0.24), Inches(0.48), 18, True, accents[i % len(accents)], PP_ALIGN.CENTER, self._style_for(data, 'title', layer), 'title')
            if desc:
                self._add_textbox(slide, desc, left + Inches(0.2), top + Inches(0.78), box_w - Inches(0.4), Inches(0.38), 11, False, theme['muted'], PP_ALIGN.CENTER, self._style_for(data, 'body', layer), 'body')
            tf_box = slide.shapes.add_textbox(left + Inches(0.22), top + Inches(1.3), box_w - Inches(0.44), Inches(2.85))
            tf = tf_box.text_frame
            self._configure_text_frame(tf)
            tf.clear()
            for j, item in enumerate(items or ['模块说明']):
                p = tf.paragraphs[0] if j == 0 else tf.add_paragraph()
                p.text = '• ' + _item_text(item)
                self._apply_text_style(p, self._style_for(data, 'body', item if isinstance(item, dict) else layer), default_size=12, default_color=theme['text'])
            if i > 0:
                arrow = slide.shapes.add_shape(MSO_SHAPE.RIGHT_ARROW, left - gap + Inches(0.03), top + Inches(2.22), gap - Inches(0.06), Inches(0.32))
                self._set_shape_fill(arrow, theme['line'])

    def _ppt_add_method_pipeline(self, prs, data):
        from pptx.enum.shapes import MSO_SHAPE
        from pptx.enum.text import PP_ALIGN
        from pptx.util import Inches, Pt
        theme = self._theme()
        slide = self._blank_slide(prs, data.get('title') or '研究方法流程', data)
        stages = _as_list(data.get('stages') or data.get('steps') or data.get('items'))[:5]
        if not stages:
            stages = [
                {'title': '数据采集', 'items': ['实验数据', '公开数据集']},
                {'title': '预处理', 'items': ['去噪', '归一化']},
                {'title': '模型训练', 'items': ['参数优化', '交叉验证']},
                {'title': '性能评估', 'items': ['指标对比', '消融实验']},
            ]
        n = len(stages)
        left0, top = Inches(0.8), Inches(1.55)
        area_w, gap = Inches(11.75), Inches(0.28)
        card_w = int((area_w - gap * (n - 1)) / n)
        for i, stage in enumerate(stages):
            title, desc, _, _ = self._item_title_desc(stage, i + 1)
            items = _as_list(stage.get('items') or stage.get('substeps') or stage.get('tasks') if isinstance(stage, dict) else [])[:4]
            left = left0 + i * (card_w + gap)
            head = slide.shapes.add_shape(MSO_SHAPE.CHEVRON, left, top, card_w, Inches(0.78))
            self._set_shape_fill(head, theme['primary'] if i % 2 == 0 else theme['accent'], theme['card'])
            self._add_textbox(slide, f'{i + 1}. {title}', left + Inches(0.08), top + Inches(0.22), card_w - Inches(0.18), Inches(0.28), 13, True, theme['card'], PP_ALIGN.CENTER, self._style_for(data, 'title', stage), 'title')
            body = slide.shapes.add_shape(MSO_SHAPE.ROUNDED_RECTANGLE, left + Inches(0.05), top + Inches(1.02), card_w - Inches(0.1), Inches(3.75))
            self._set_shape_fill(body, theme['card'], theme['line'])
            if desc:
                self._add_textbox(slide, desc, left + Inches(0.2), top + Inches(1.28), card_w - Inches(0.4), Inches(0.45), 11, False, theme['muted'], PP_ALIGN.CENTER, self._style_for(data, 'body', stage), 'body')
            tf_box = slide.shapes.add_textbox(left + Inches(0.22), top + Inches(1.85), card_w - Inches(0.44), Inches(2.4))
            tf = tf_box.text_frame
            self._configure_text_frame(tf)
            tf.clear()
            for j, item in enumerate(items or ['关键步骤']):
                p = tf.paragraphs[0] if j == 0 else tf.add_paragraph()
                p.text = '• ' + _item_text(item)
                self._apply_text_style(p, self._style_for(data, 'body', item if isinstance(item, dict) else stage), default_size=11, default_color=theme['text'])

    def _ppt_add_table(self, prs, data):
        from pptx.enum.text import PP_ALIGN
        from pptx.util import Inches, Pt
        theme = self._theme()
        slide = self._blank_slide(prs, data.get('title') or '数据表格', data)
        headers = _as_list(data.get('headers') or data.get('columns'))[:6]
        rows = _as_list(data.get('rows') or data.get('data'))[:8]
        if not headers:
            headers = ['方法', '指标A', '指标B', '结论']
        if rows and isinstance(rows[0], dict):
            rows = _rows_from_records(rows, headers)
        if not rows:
            rows = [['Baseline', '82.1', '79.4', '参考'], ['Ours', '88.6', '85.2', '最佳']]
        table_shape = slide.shapes.add_table(len(rows) + 1, len(headers), Inches(0.85), Inches(1.55), Inches(11.65), Inches(4.9))
        table = table_shape.table
        for c, h in enumerate(headers):
            cell = table.cell(0, c)
            cell.text = _as_text(h)
            self._configure_text_frame(cell.text_frame)
            cell.fill.solid(); cell.fill.fore_color.rgb = theme['primary']
            p = cell.text_frame.paragraphs[0]
            self._apply_text_style(p, self._style_for(data, 'header'), default_size=12, default_bold=True, default_color=theme['card'], default_align=PP_ALIGN.CENTER)
        for r, row in enumerate(rows, start=1):
            vals = _as_list(row)
            for c in range(len(headers)):
                cell = table.cell(r, c)
                cell.text = _as_text(vals[c] if c < len(vals) else '')
                self._configure_text_frame(cell.text_frame)
                cell.fill.solid(); cell.fill.fore_color.rgb = theme['soft_blue'] if r % 2 else theme['card']
                p = cell.text_frame.paragraphs[0]
                self._apply_text_style(p, self._style_for(data, 'body'), default_size=11, default_color=theme['text'], default_align=PP_ALIGN.CENTER)

    def _ppt_add_chart(self, prs, data):
        from pptx.enum.shapes import MSO_SHAPE
        from pptx.enum.shapes import MSO_CONNECTOR
        from pptx.enum.text import PP_ALIGN
        from pptx.util import Inches, Pt
        theme = self._theme()
        slide = self._blank_slide(prs, data.get('title') or '数据图表', data)
        chart_type = _as_text(data.get('chart_type') or data.get('kind') or 'bar').lower()
        categories = _as_list(data.get('categories') or data.get('labels'))[:8] or ['A', 'B', 'C', 'D']
        values = _as_list(data.get('values') or data.get('data'))[:len(categories)] or [30, 55, 42, 70]
        nums = []
        for v in values:
            try:
                nums.append(float(v.get('value') if isinstance(v, dict) else v))
            except Exception:
                nums.append(0.0)
        max_v = max(nums) if nums else 1.0
        max_v = max_v or 1.0
        if chart_type in ('pie', 'pie_chart'):
            self._ppt_add_pie_chart(slide, data, categories, nums, theme, Inches, PP_ALIGN)
            return
        if chart_type == 'line':
            self._add_textbox(slide, '折线图（节点表示趋势）', Inches(0.9), Inches(1.35), Inches(3.5), Inches(0.35), 12, False, theme['muted'], PP_ALIGN.LEFT, self._style_for(data, 'label'), 'label')
            left0, bottom = Inches(1.2), Inches(5.75)
            area_w, area_h = Inches(10.7), Inches(3.9)
            prev = None
            for i, (cat, val) in enumerate(zip(categories, nums)):
                x = left0 + int(area_w * i / max(1, len(nums) - 1))
                y = bottom - int(area_h * val / max_v)
                if prev:
                    px, py = prev
                    line = slide.shapes.add_connector(MSO_CONNECTOR.STRAIGHT, px, py, x, y)
                    line.line.color.rgb = theme['primary']
                    line.line.width = Pt(2)
                dot = slide.shapes.add_shape(MSO_SHAPE.OVAL, x - Inches(0.08), y - Inches(0.08), Inches(0.16), Inches(0.16))
                self._set_shape_fill(dot, theme['accent'], theme['card'])
                self._add_textbox(slide, _as_text(cat), x - Inches(0.35), bottom + Inches(0.15), Inches(0.7), Inches(0.25), 10, False, theme['muted'], PP_ALIGN.CENTER, self._style_for(data, 'label'), 'label')
                prev = (x, y)
        else:
            left0, bottom = Inches(1.0), Inches(5.85)
            area_w, area_h = Inches(11.2), Inches(4.0)
            gap = Inches(0.14)
            bar_w = int((area_w - gap * (len(nums) - 1)) / max(1, len(nums)))
            for i, (cat, val) in enumerate(zip(categories, nums)):
                h = max(Inches(0.08), int(area_h * val / max_v))
                left = left0 + i * (bar_w + gap)
                bar = slide.shapes.add_shape(MSO_SHAPE.ROUNDED_RECTANGLE, left, bottom - h, bar_w, h)
                self._set_shape_fill(bar, theme['primary'] if i % 2 == 0 else theme['accent'])
                self._add_textbox(slide, str(val).rstrip('0').rstrip('.'), left, bottom - h - Inches(0.35), bar_w, Inches(0.25), 10, True, theme['text'], PP_ALIGN.CENTER, self._style_for(data, 'value'), 'value')
                self._add_textbox(slide, _as_text(cat), left, bottom + Inches(0.12), bar_w, Inches(0.28), 10, False, theme['muted'], PP_ALIGN.CENTER, self._style_for(data, 'label'), 'label')

    def _ppt_add_pie_chart(self, slide, data, categories, nums, theme, Inches, PP_ALIGN):
        from pptx.enum.shapes import MSO_SHAPE
        total = sum(v for v in nums if v > 0) or 1.0
        colors = [theme['primary'], theme['accent'], theme['green'], theme['purple'], theme['soft_orange'], theme['soft_blue']]
        cx, cy = Inches(3.75), Inches(3.65)
        radius = Inches(1.45)
        start = 0.0
        # python-pptx 没有稳定的扇形 API，这里用“可编辑近似饼图”：圆环主图 + 图例 + 百分比。
        # 相比静默退化为柱状图，pie/pie_chart 现在有明确且一致的饼图表达。
        base = slide.shapes.add_shape(MSO_SHAPE.OVAL, cx - radius, cy - radius, radius * 2, radius * 2)
        self._set_shape_fill(base, colors[0], theme['card'])
        ring = slide.shapes.add_shape(MSO_SHAPE.OVAL, cx - int(radius * 0.62), cy - int(radius * 0.62), int(radius * 1.24), int(radius * 1.24))
        self._set_shape_fill(ring, theme['card'], theme['card'])
        for i, (cat, val) in enumerate(zip(categories, nums)):
            pct = max(0.0, val) / total
            color = colors[i % len(colors)]
            y = Inches(1.6) + i * Inches(0.52)
            swatch = slide.shapes.add_shape(MSO_SHAPE.ROUNDED_RECTANGLE, Inches(7.0), y, Inches(0.28), Inches(0.22))
            self._set_shape_fill(swatch, color, color)
            label = f'{_as_text(cat)}  {pct * 100:.1f}%'
            self._add_textbox(slide, label, Inches(7.38), y - Inches(0.08), Inches(4.4), Inches(0.38), 13, False, theme['text'], PP_ALIGN.LEFT, self._style_for(data, 'label'), 'label')
            if i > 0:
                wedge = slide.shapes.add_shape(MSO_SHAPE.ARC, cx - radius, cy - radius, radius * 2, radius * 2)
                wedge.line.color.rgb = color
                wedge.line.width = Inches(0.16)
                wedge.rotation = start * 360.0
            start += pct
        self._add_textbox(slide, '饼图', cx - Inches(0.55), cy - Inches(0.18), Inches(1.1), Inches(0.36), 15, True, theme['text'], PP_ALIGN.CENTER, self._style_for(data, 'label'), 'label')

    def _ppt_add_experiment_design(self, prs, data):
        from pptx.enum.shapes import MSO_SHAPE
        from pptx.enum.text import PP_ALIGN
        from pptx.util import Inches, Pt
        theme = self._theme()
        slide = self._blank_slide(prs, data.get('title') or '实验设计', data)
        groups = _as_list(data.get('groups') or data.get('experiments') or data.get('items'))[:4]
        if not groups:
            groups = [{'name': 'Baseline', 'method': '传统方法', 'metrics': ['Accuracy', 'F1']}, {'name': 'Ours', 'method': '改进模型', 'metrics': ['Accuracy', 'F1']}]
        n = len(groups)
        left0, top, area_w, gap = Inches(0.9), Inches(1.55), Inches(11.55), Inches(0.25)
        card_w = int((area_w - gap * (n - 1)) / n)
        for i, g in enumerate(groups):
            left = left0 + i * (card_w + gap)
            card = slide.shapes.add_shape(MSO_SHAPE.ROUNDED_RECTANGLE, left, top, card_w, Inches(4.85))
            self._set_shape_fill(card, theme['soft_blue'] if i % 2 == 0 else theme['soft_green'], theme['line'])
            name = _as_text(g.get('name') or g.get('title') or f'实验组 {i + 1}' if isinstance(g, dict) else g)
            method = _as_text(g.get('method') or g.get('desc') or '' if isinstance(g, dict) else '')
            metrics = _as_list(g.get('metrics') or g.get('indicators') or [] if isinstance(g, dict) else [])[:5]
            self._add_textbox(slide, name, left + Inches(0.18), top + Inches(0.28), card_w - Inches(0.36), Inches(0.45), 18, True, theme['primary'], PP_ALIGN.CENTER, self._style_for(data, 'title', g), 'title')
            self._add_textbox(slide, method or '实验方法', left + Inches(0.25), top + Inches(0.95), card_w - Inches(0.5), Inches(0.65), 13, False, theme['text'], PP_ALIGN.CENTER, self._style_for(data, 'body', g), 'body')
            tf_box = slide.shapes.add_textbox(left + Inches(0.35), top + Inches(1.95), card_w - Inches(0.7), Inches(2.1))
            tf = tf_box.text_frame
            self._configure_text_frame(tf)
            tf.clear()
            for j, m in enumerate(metrics or ['评价指标']):
                p = tf.paragraphs[0] if j == 0 else tf.add_paragraph()
                p.text = '✓ ' + _item_text(m)
                self._apply_text_style(p, self._style_for(data, 'body', m if isinstance(m, dict) else g), default_size=12, default_color=theme['text'])

    def _ppt_add_ablation(self, prs, data):
        headers = data.get('headers') or ['实验设置', '去除模块', '关键指标', '结论']
        rows = data.get('rows') or data.get('variants') or data.get('items') or [
            ['Full Model', '-', '88.6', '最佳'],
            ['w/o Module A', '模块A', '84.2', '下降明显'],
            ['w/o Module B', '模块B', '85.1', '有一定影响'],
        ]
        table_data = dict(data)
        table_data.update({'title': data.get('title') or '消融实验', 'headers': headers, 'rows': rows})
        self._ppt_add_table(prs, table_data)
